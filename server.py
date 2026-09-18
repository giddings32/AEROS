# AEROS server.py V1
#!/usr/bin/env python3
from http.server import ThreadingHTTPServer, SimpleHTTPRequestHandler
from pathlib import Path, PurePosixPath
from urllib.parse import unquote, urlsplit
from io import BytesIO
from copy import deepcopy
import xml.etree.ElementTree as ET
import functools, json, os, re, datetime, shutil, uuid, mimetypes, sys, zipfile, hashlib, stat, unicodedata, threading
from email.parser import BytesParser
from email.policy import default as email_default
from engagement_store import EngagementStore, EngagementAlreadyExistsError, contained, exploitation_directory
from engagement_model import (
    APPLICATION_VERSION,
    RevisionConflictError,
    SCHEMA_VERSION,
)
from artifact_identity import (
    classify_artifact, infer_artifact_type, merge_observation,
    normalize_artifact, normalize_logical_path, sha256_bytes,
)

PROJECT_MUTATION_LOCK = threading.RLock()


def serialized_project_mutation(function):
    """Serialize state/filesystem mutations that must remain one logical unit."""
    @functools.wraps(function)
    def wrapped(*args, **kwargs):
        with PROJECT_MUTATION_LOCK:
            return function(*args, **kwargs)
    return wrapped


try:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.enum.section import WD_SECTION
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
except Exception as e:
    Document = None

ROOT = Path(__file__).resolve().parent
HOST = "127.0.0.1"
PORT = int(os.environ.get("AEROS_PORT", "8765"))
APPLICATION_ID = "aeros-v1"
BUILD_IDENTITY_FILES = (
    "server.py", "desktop_app.py", "app.js", "index.html",
    "app/core/theme-bootstrap.js", "app/core/application.js",
    "app/styles/themes.css", "app/styles/core.css",
    "app/styles/components.css", "app/styles/features.css", "app/styles/pages.css",
    "templates.js", "notes.js", "highlighters.js",
    "app/features/readiness/readiness.js", "review.js",
    "app/features/workflow/exam-controls.js",
    "app/features/workflow/oscp-workflow.js",
    "app/features/navigator/navigator-core.js",
    "app/features/navigator/oscp-navigator.js",
    "app/features/host-workbench/host-workbench.js",
    "profiles.js",
    "scan-intelligence.js", "app/features/imports/import-lifecycle.js",
    "app/features/imports/autorecon-live-import.js", "app/features/imports/live-import-core.js", "app/features/imports/live-import-browser.js",
    "app/features/imports/recon-import.js", "app/features/recon/recon-projection.js", "app/features/recon/recon-triage.js", "app/features/recon/initial-recon.js",
    "app/features/peas/peas-organizer.js",
    "app/features/artifacts/artifact-identity.js", "nmapimport.js",
    "kernel-exploit-helper.js", "network-context.js", "network-context-ui.js", "reference-note-matcher.js", "asset-groups.js",
    "context-foundation.js", "app/features/access-contexts/access-contexts.js",
    "linux-software-intelligence.js", "methodology-tasks.js", "methodology-ui.js",
    "stage-awareness.js", "app/features/methodology/methodology-contexts.js",
    "host-profile.js", "command-paths.js",
    "app/features/reference-notes/reference-note-renderer.js",
    "app/features/recon/recon-workspace.js", "app/features/exploitation-path/exploitation-path.js", "app/features/recon/web-app-review.js",
    "app/features/access-leads/access-leads.js",
    "app/features/exploit-attempts/exploit-attempts.js",
    "app/features/findings/findings.js",
    "app/features/persistence/persistence.js",
    "app/features/engagements/engagements.js",
    "app/features/evidence/evidence.js", "app/features/reporting/reporting.js",
    "ux-enhancements.js", "app/features/asset-groups/group-manager.js",
    "artifact_identity.py", "engagement_model.py", "engagement_store.py", "legacy_import.py",
)


def build_identity(root=ROOT):
    """Fingerprint the runnable application files used by both launch paths."""
    root = Path(root)
    digest = hashlib.sha256()
    for relative_name in BUILD_IDENTITY_FILES:
        path = root / relative_name
        digest.update(relative_name.encode("utf-8"))
        digest.update(b"\0")
        if path.is_file():
            digest.update(path.read_bytes())
        else:
            digest.update(b"<missing>")
        digest.update(b"\0")
    return digest.hexdigest()


def data_root_fingerprint(path):
    normalized = os.path.normcase(str(Path(path).expanduser().resolve()))
    return hashlib.sha256(normalized.encode("utf-8")).hexdigest()

def application_data_root(env=None, platform_name=None, home=None):
    """Return the per-user, cross-platform AEROS V1 data directory."""
    env = os.environ if env is None else env
    platform_name = ("win32" if os.name == "nt" else sys.platform) if platform_name is None else str(platform_name)
    home = Path.home() if home is None else Path(home)
    override = (env.get("AEROS_DATA_DIR") or env.get("OSCP_REPORT_BUILDER_DATA_DIR") or "").strip()
    if override:
        return Path(override).expanduser().resolve()
    if platform_name == "nt" or platform_name.startswith("win"):
        base = Path(env.get("LOCALAPPDATA") or (home / "AppData" / "Local"))
        return base / "AEROS" / "data"
    if platform_name == "darwin":
        return home / "Library" / "Application Support" / "AEROS"
    base = Path(env.get("XDG_DATA_HOME") or (home / ".local" / "share"))
    return base / "aeros"


def legacy_application_data_roots(env=None, platform_name=None, home=None):
    """Locations used by the pre-V1 development builds."""
    env = os.environ if env is None else env
    platform_name = ("win32" if os.name == "nt" else sys.platform) if platform_name is None else str(platform_name)
    home = Path.home() if home is None else Path(home)
    if platform_name == "nt" or platform_name.startswith("win"):
        base = Path(env.get("LOCALAPPDATA") or (home / "AppData" / "Local"))
        return [base / "OSCP_Report_Builder" / "data"]
    if platform_name == "darwin":
        return [home / "Library" / "Application Support" / "OSCP Report Builder"]
    base = Path(env.get("XDG_DATA_HOME") or (home / ".local" / "share"))
    return [base / "oscp-report-builder"]


def migrate_legacy_data_root(target, legacy_roots=None):
    """Copy an existing development data store into AEROS V1 once."""
    target = Path(target)
    if target.exists() and any(target.iterdir()):
        return None
    sources = legacy_application_data_roots() if legacy_roots is None else legacy_roots
    for source in sources:
        source = Path(source)
        if not source.exists() or source.resolve() == target.resolve():
            continue
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copytree(source, target, dirs_exist_ok=True)
        return source
    return None


if os.name != "nt":
    # Project data includes credentials and evidence; new files must be private
    # to the account running the local application.
    os.umask(0o077)


DATA_ROOT = application_data_root()
DATA_ROOT_OVERRIDE = bool((os.environ.get("AEROS_DATA_DIR") or os.environ.get("OSCP_REPORT_BUILDER_DATA_DIR") or "").strip())
BUILD_IDENTITY = build_identity()
DATA_ROOT_FINGERPRINT = data_root_fingerprint(DATA_ROOT)
LEGACY_DATA_MIGRATED_FROM = None if DATA_ROOT_OVERRIDE else migrate_legacy_data_root(DATA_ROOT)
DATA_ROOT.mkdir(parents=True, exist_ok=True)
if os.name != "nt":
    os.chmod(DATA_ROOT, 0o700)
LABS_ROOT = DATA_ROOT / "labs"
ASSETS_ROOT = DATA_ROOT / "assets"
LAB_BACKUPS_ROOT = DATA_ROOT / "backups" / "labs"
ARCHIVES_ROOT = DATA_ROOT / "archives"
IMPORT_INTAKE_STAGING_ROOT = DATA_ROOT / "import-intake-staging"
PROFILES_PATH = DATA_ROOT / "profiles.json"
REPORT_TEMPLATES_ROOT = DATA_ROOT / "report_templates"
REPORT_TEMPLATES_REGISTRY = DATA_ROOT / "report_templates.json"
ENGAGEMENT_DOCUMENTS_ROOT = DATA_ROOT / "engagements"
COMMAND_NOTES_ROOT = ROOT / "command-notes"
COMMAND_NOTES_MANIFEST_PATH = COMMAND_NOTES_ROOT / "index.json"
CUSTOM_COMMAND_NOTES_ROOT = DATA_ROOT / "command-notes"
CUSTOM_REFERENCE_NOTES_ROOT = DATA_ROOT / "reference-notes"
MAX_COMMAND_NOTE_BYTES = 2 * 1024 * 1024
MAX_LAB_BACKUPS = 24
MIN_LAB_BACKUP_INTERVAL_SECONDS = 5 * 60
REFERENCE_PACKS_ROOT = DATA_ROOT / "reference-packs"
REFERENCE_PACK_MANIFEST = "index.json"
CANONICAL_REFERENCE_PACK_ID = "aeros-pentesting-notes-v1"
CANONICAL_REFERENCE_PACK_FILENAME = "AEROS-Pentesting-Notes-V1.aeros-notes"
MAX_REFERENCE_PACK_BYTES = 50 * 1024 * 1024
MAX_REFERENCE_PACK_MEMBERS = 2000
MAX_REFERENCE_PACK_TOTAL_BYTES = 100 * 1024 * 1024
MAX_REFERENCE_PACK_COMPRESSION_RATIO = 250
MAX_REFERENCE_NOTE_BYTES = 4 * 1024 * 1024
REFERENCE_NOTE_STAGES = {"recon", "enumeration", "exploitation", "foothold", "privesc", "lateral", "looting", "reporting"}
REFERENCE_NOTE_MODULES = {"general", "oscp", "web", "active-directory"}
REFERENCE_NOTE_SCOPES = {"engagement", "asset-group", "host", "service", "access", "reporting", "library-only"}
REFERENCE_NOTE_TRIGGER_MODES = {"always", "asset-group-type", "service", "port", "host-os", "access", "credential", "domain", "event", "manual"}
REFERENCE_NOTE_ACTIVATION_EVENTS = {
    "host-added", "service-discovered", "service-accessed", "credential-discovered",
    "foothold-acquired", "privilege-changed", "new-access-context", "domain-discovered",
}
REFERENCE_NOTE_OS_VALUES = {"both", "linux", "windows"}
REFERENCE_NOTE_TASK_ROLES = {"method", "conditional-method", "task-guide", "reference", "workflow-guide", "exclude", "duplicate"}
REFERENCE_NOTE_TASK_CLASSES = {"core", "conditional", "conditional-warning", "reference-only", "excluded"}
REFERENCE_NOTE_COMPLETION_RULES = {"any-method", "manual-confirm", "conditional-manual", "reference-only", "not-applicable"}
REFERENCE_NOTE_EVIDENCE_TRIGGERS = {"linux-sudo-binary", "linux-suid-binary", "linux-capability-binary"}
REFERENCE_ASSET_GROUP_TYPES = {"active-directory", "web-application", "network-segment", "custom"}
REFERENCE_NOTE_TOKEN_RE = re.compile(r"^[a-z0-9][a-z0-9._-]*$")
STORAGE_PATH = DATA_ROOT / "engagements.json"
STORE = EngagementStore(DATA_ROOT)


def read_profile_store(path=PROFILES_PATH):
    path = Path(path)
    profiles = json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}
    if not isinstance(profiles, dict):
        raise ValueError("Invalid profiles data")
    return profiles


def save_profile_store(profiles, path=PROFILES_PATH):
    if not isinstance(profiles, dict):
        raise ValueError("profiles must be an object")
    path = Path(path)
    if not profiles and read_profile_store(path):
        raise ValueError("Refusing to replace a non-empty profile store with an empty profile object")
    atomic_write_json(path, profiles)


# --- Local security controls -------------------------------------------------
# A fresh token is minted every launch and injected into index.html at serve
# time. The frontend attaches it to every /api/ request. Any other page the
# browser has open cannot read index.html cross-origin, so it cannot obtain the
# token, which neutralises cross-site requests to the local API.
import secrets
APP_TOKEN = secrets.token_urlsafe(32)
CSP_NONCE = secrets.token_urlsafe(24)
SESSION_COOKIE_NAME = "aeros_session"
NATIVE_SHUTDOWN_TOKEN = os.environ.get("AEROS_NATIVE_SHUTDOWN_TOKEN", "").strip()
NATIVE_OWNER_PID = os.environ.get("AEROS_NATIVE_OWNER_PID", "").strip()
ALLOWED_ORIGINS = {f"http://{HOST}:{PORT}", f"http://localhost:{PORT}"}
# Host-header allowlist. A DNS-rebinding page connects to 127.0.0.1 but carries
# its own hostname in the Host header, so pinning Host to the loopback names we
# actually serve on rejects rebinding without affecting real same-origin traffic
# (browser tabs, <img> loads, and the fetch shim all send a matching Host).
ALLOWED_HOSTS = {f"{HOST}:{PORT}", f"localhost:{PORT}", f"127.0.0.1:{PORT}"}
MAX_BODY_BYTES = 50 * 1024 * 1024  # hard cap on request bodies (screenshots included)
MAX_DOCX_MEMBERS = 2000
MAX_DOCX_EXPANDED_BYTES = 100 * 1024 * 1024
MAX_DOCX_COMPRESSION_RATIO = 250
MAX_DOCX_XML_MEMBER_BYTES = 10 * 1024 * 1024
MAX_EXTRACTED_DOCUMENT_CHARS = 2 * 1024 * 1024
MAX_EXTRACTED_PDF_PAGES = 250
MAX_DOCX_RELATIONSHIP_BYTES = 2 * 1024 * 1024
MAX_IMAGE_DIMENSION = 20000
MAX_IMAGE_PIXELS = 50000000

# Structured AEROS collector import limits. ZIP is a transport container only:
# every member must still be an independently valid AEROS recon JSON payload.
SUPPORTED_RECON_SCHEMAS = {"aeros-recon", "aeros-collector", "product-recon"}
RECON_SCHEMA_MAJOR = "1"
MAX_RECON_ARCHIVE_MEMBERS = 32
MAX_RECON_PAYLOADS = 16
MAX_RECON_JSON_BYTES = 25 * 1024 * 1024
MAX_RECON_TOTAL_JSON_BYTES = 50 * 1024 * 1024
MAX_RECON_COMPRESSION_RATIO = 250

# Unified Import Results intake limits. Archive members are staged only for
# operator preview/commit and are never executed.
MAX_IMPORT_ARCHIVE_BYTES = 48 * 1024 * 1024
MAX_IMPORT_EXPANDED_BYTES = 100 * 1024 * 1024
MAX_IMPORT_MEMBER_BYTES = 25 * 1024 * 1024
MAX_IMPORT_ARCHIVE_MEMBERS = 500
MAX_IMPORT_COMPRESSION_RATIO = 200
MAX_IMPORT_NESTED_ARCHIVE_DEPTH = 0
MAX_IMPORT_PREVIEW_BYTES = 256 * 1024
IMPORT_STAGING_TTL_SECONDS = 4 * 60 * 60
MAX_ENGAGEMENT_ARCHIVE_BYTES = 1024 * 1024 * 1024
MAX_ENGAGEMENT_ARCHIVE_MEMBERS = 20000
MAX_ENGAGEMENT_ARCHIVE_EXPANDED_BYTES = 2 * 1024 * 1024 * 1024
MAX_ENGAGEMENT_ARCHIVE_MEMBER_BYTES = 512 * 1024 * 1024
MAX_ENGAGEMENT_ARCHIVE_COMPRESSION_RATIO = 500
MAX_ENGAGEMENT_MANIFEST_BYTES = 2 * 1024 * 1024
# -----------------------------------------------------------------------------

# Report spacing controls
GAP_TIGHT = 0      # tight heading stack, like the OffSec sample
GAP_FIELD = 8      # visible blank line between Vulnerability Explanation / Fix / Severity / Steps
GAP_BLOCK = 12     # gap after a complete evidence item or scan block
GAP_MAJOR = 16     # gap between major host report sections

BODY_FONT_NAME = "Calibri"
BODY_FONT_SIZE = 11

SEVERITY_COLORS = {
    "critical": RGBColor(255, 0, 0),
    "high": RGBColor(255, 102, 0),
    "medium": RGBColor(191, 144, 0),
    "low": RGBColor(0, 112, 192),
    "informational": RGBColor(89, 89, 89),
    "info": RGBColor(89, 89, 89),
}
SAFE_NAME_RE = re.compile(r'[<>:"/\\|?*\x00-\x1F]+')
WINDOWS_RESERVED_NAMES = {
    "CON", "PRN", "AUX", "NUL",
    *(f"COM{index}" for index in range(1, 10)),
    *(f"LPT{index}" for index in range(1, 10)),
}

def safe_name(value):
    value = str(value or "").strip()
    value = SAFE_NAME_RE.sub("_", value).strip(" .")
    return value or "Unknown"


def storage_component(value, max_slug_length=72):
    """Return a bounded collision-resistant filesystem component."""
    raw = unicodedata.normalize("NFC", str(value or "").strip())
    slug = safe_name(raw)
    stem = slug.split(".", 1)[0].upper()
    unchanged = slug.casefold() == raw.casefold()
    portable_ascii = raw.isascii() and re.fullmatch(r"[A-Za-z0-9 _.-]+", raw) is not None
    if unchanged and portable_ascii and len(slug) <= max_slug_length and stem not in WINDOWS_RESERVED_NAMES:
        return slug
    digest = hashlib.sha256(raw.encode("utf-8", errors="replace")).hexdigest()[:12]
    prefix = slug[:max_slug_length].rstrip(" ._-") or "Unknown"
    return f"{prefix}--{digest}"


def _legacy_state_belongs_to(lab_name):
    legacy = LABS_ROOT / f"{safe_name(lab_name)}.json"
    if not legacy.is_file():
        return False
    try:
        payload = json.loads(legacy.read_text(encoding="utf-8"))
    except Exception:
        return False
    owner = str(payload.get("projectName") or payload.get("labName") or "").strip()
    return bool(owner) and owner.casefold() == str(lab_name or "").strip().casefold()


def engagement_storage_key(lab_name):
    """Preserve an owned legacy path, otherwise use a collision-resistant key."""
    legacy_key = safe_name(lab_name)
    current_key = storage_component(lab_name)
    if legacy_key != current_key and _legacy_state_belongs_to(lab_name):
        return legacy_key
    return current_key









def parse_known_artifacts(value):
    if isinstance(value, list):
        return [row for row in value if isinstance(row, dict)]
    try:
        parsed = json.loads(str(value or "[]"))
    except (TypeError, ValueError, json.JSONDecodeError):
        return []
    return [row for row in parsed if isinstance(row, dict)] if isinstance(parsed, list) else []


def store_versioned_artifact(data, *, root, original_filename, logical_path, source, objective, artifact_type, kind, known_artifacts=None, host_key=""):
    if not isinstance(data, (bytes, bytearray)) or not data:
        raise ValueError("Artifact must not be empty")
    root = Path(root).resolve(); root.mkdir(parents=True, exist_ok=True)
    digest = sha256_bytes(bytes(data))
    known = parse_known_artifacts(known_artifacts or [])

    def resolve_known_hash(stored_filename):
        target = (root / safe_name(stored_filename)).resolve()
        if not path_is_inside(target, root) or not target.is_file():
            raise FileNotFoundError("Stored artifact revision not found")
        return file_sha256(target)

    verified_known = [normalize_artifact(row, hash_resolver=resolve_known_hash) for row in known]
    outcome, existing = classify_artifact(
        verified_known, sha256=digest, logical_path=logical_path, source=source,
        objective=objective, artifact_type=artifact_type,
    )
    stored_filename = ""
    if outcome == "unchanged" and existing:
        normalized = normalize_artifact(existing)
        matching = next((row for row in normalized.get("revisions", []) if row.get("sha256") == digest), None)
        stored_filename = str((matching or {}).get("storedFilename") or normalized.get("storedFilename") or "")
        if stored_filename and not (root / safe_name(stored_filename)).is_file():
            stored_filename = ""
            outcome = "updated"
    target = None
    if not stored_filename:
        suffix = Path(original_filename).suffix.lower()
        stored_filename = f"{storage_component(Path(original_filename).stem)}--{uuid.uuid4().hex}{suffix}"
        target = (root / stored_filename).resolve()
        if not path_is_inside(target, root):
            raise ValueError("Invalid artifact filename")
    merged = merge_observation(
        existing, outcome=outcome, sha256=digest, stored_filename=stored_filename,
        original_filename=Path(original_filename).name, logical_path=logical_path,
        size=len(data), mime_type=mimetypes.guess_type(original_filename)[0] or "application/octet-stream",
        source=source, objective=objective, artifact_type=artifact_type,
        host_key=host_key, observed_at=datetime.datetime.now(datetime.timezone.utc).isoformat(),
    )
    merged["kind"] = kind
    merged["importOutcome"] = outcome
    if target is not None:
        atomic_write_bytes(target, bytes(data))
    return merged


def rollback_uncommitted_artifact(root, item, extra_filenames=()):
    if not isinstance(item, dict) or item.get("importOutcome") not in {"new", "updated"}:
        return
    root = Path(root).resolve()
    names = [item.get("storedFilename"), *extra_filenames]
    for filename in dict.fromkeys(str(value or "") for value in names):
        if not filename:
            continue
        target = (root / safe_name(filename)).resolve()
        if path_is_inside(target, root) and target.is_file():
            target.unlink()


def engagement_documents_root(lab_name):
    return contained(STORE.project_root(lab_name), "documents")


def roe_documents_root(lab_name):
    return engagement_documents_root(lab_name) / "roe"


def _bounded_extracted_text(parts, separator="\n"):
    output = []
    total = 0
    truncated = False
    for part in parts:
        text = str(part or "")
        if not text:
            continue
        separator_cost = len(separator) if output else 0
        remaining = MAX_EXTRACTED_DOCUMENT_CHARS - total - separator_cost
        if remaining <= 0:
            truncated = True
            break
        if len(text) > remaining:
            output.append(text[:remaining])
            total += remaining
            truncated = True
            break
        output.append(text)
        total += separator_cost + len(text)
    return separator.join(output), truncated


def extract_engagement_document_text(data, suffix):
    suffix = str(suffix or "").lower()
    if suffix in {".txt", ".md"}:
        text = data.decode("utf-8", errors="replace")
        if len(text) > MAX_EXTRACTED_DOCUMENT_CHARS:
            return text[:MAX_EXTRACTED_DOCUMENT_CHARS], "Extracted text was truncated at the 2 MB safety limit."
        return text, ""
    if suffix == ".docx":
        if Document is None:
            return "", "python-docx is unavailable; the original document was stored without extracted text."
        doc = Document(BytesIO(data))
        def docx_lines():
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    yield paragraph.text
            for table in doc.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        yield " | ".join(cells)
        text, truncated = _bounded_extracted_text(docx_lines())
        return text, "Extracted text was truncated at the 2 MB safety limit." if truncated else ""
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
            reader = PdfReader(BytesIO(data))
            page_limit = min(len(reader.pages), MAX_EXTRACTED_PDF_PAGES)
            text, truncated = _bounded_extracted_text(
                ((reader.pages[index].extract_text() or "").strip() for index in range(page_limit)),
                separator="\n\n",
            )
            text = text.strip()
            warnings = []
            if len(reader.pages) > page_limit:
                warnings.append(f"Text extraction stopped after {MAX_EXTRACTED_PDF_PAGES} pages.")
            if truncated:
                warnings.append("Extracted text was truncated at the 2 MB safety limit.")
            if not text:
                warnings.append("The PDF was stored, but no readable text could be extracted.")
            warning = " ".join(warnings)
            return text, warning
        except ImportError:
            return "", "Install pypdf to extract PDF text. The original PDF was stored successfully."
        except Exception as exc:
            return "", f"The PDF was stored, but text extraction failed: {exc}"
    return "", ""


def upload_roe_document(lab_name, file_item, display_name="", document_kind="roe", known_documents="[]"):
    lab_name = str(lab_name or "").strip()
    if not lab_name:
        raise ValueError("Select an engagement before uploading a document")
    if not project_state_path(lab_name).exists():
        raise FileNotFoundError("The selected engagement has not been saved yet")
    if not file_item or not file_item.filename:
        raise ValueError("Choose a PDF, DOCX, TXT, or Markdown file")
    suffix = Path(file_item.filename).suffix.lower()
    if suffix not in {".pdf", ".docx", ".txt", ".md"}:
        raise ValueError("Supported document types are PDF, DOCX, TXT, and Markdown")
    data = file_item.file.read()
    if not data or len(data) > 25 * 1024 * 1024:
        raise ValueError("Document must be non-empty and smaller than 25 MB")
    if suffix == ".docx":
        validate_docx_bytes(data)
    original = Path(file_item.filename).name
    kind = safe_name(document_kind or "roe").lower()
    root = roe_documents_root(lab_name)
    item = store_versioned_artifact(
        data, root=root, original_filename=original,
        logical_path=f"{kind}/{original}", source="authorization-document", objective=kind,
        artifact_type=f"document-{suffix.lstrip('.') or 'file'}", kind="scope-authorization",
        known_artifacts=known_documents,
    )
    extracted_name = ""
    try:
        extracted_text, warning = extract_engagement_document_text(data, suffix)
        if extracted_text:
            extracted_name = f"{Path(item['storedFilename']).stem}.txt"
            atomic_write_text(root / extracted_name, extracted_text)
        item.update({
            "category": "scope-authorization", "kind": kind,
            "displayName": str(display_name or Path(original).stem).strip() or "Rules of Engagement",
            "extractedFilename": extracted_name, "extractedText": extracted_text,
            "extractionWarning": warning,
        })
        return item
    except Exception:
        rollback_uncommitted_artifact(root, item, [extracted_name])
        raise


def roe_document_path(lab_name, stored_filename):
    target_root = roe_documents_root(lab_name).resolve()
    target = (target_root / safe_name(stored_filename)).resolve()
    if not path_is_inside(target, target_root) or not target.exists() or not target.is_file():
        raise FileNotFoundError("Engagement document not found")
    return target


def delete_roe_document(lab_name, stored_filename, extracted_filename="", stored_filenames=None):
    deleted = False
    names = [stored_filename, extracted_filename, *(stored_filenames or [])]
    for stored in list(names):
        if stored and Path(str(stored)).suffix.lower() != ".txt":
            names.append(f"{Path(str(stored)).stem}.txt")
    for name in dict.fromkeys(str(value or "") for value in names):
        if not name:
            continue
        try:
            path = roe_document_path(lab_name, name)
        except FileNotFoundError:
            continue
        path.unlink(); deleted = True
    return deleted



def general_documents_root(lab_name):
    return engagement_documents_root(lab_name) / "library"


def upload_engagement_document(lab_name, file_item, display_name="", category="supporting", notes="", known_documents="[]"):
    lab_name = str(lab_name or "").strip()
    if not lab_name:
        raise ValueError("Select an engagement before uploading a document")
    if not project_state_path(lab_name).exists():
        raise FileNotFoundError("The selected engagement has not been saved yet")
    if not file_item or not file_item.filename:
        raise ValueError("Choose a document to upload")
    data = file_item.file.read()
    if not data or len(data) > 50 * 1024 * 1024:
        raise ValueError("Document must be non-empty and smaller than 50 MB")
    original = Path(file_item.filename).name
    suffix = Path(original).suffix.lower()
    if suffix == ".docx":
        validate_docx_bytes(data)
    category_key = safe_name(category or "supporting").lower()
    root = general_documents_root(lab_name)
    item = store_versioned_artifact(
        data, root=root, original_filename=original,
        logical_path=f"{category_key}/{original}", source="engagement-document", objective=category_key,
        artifact_type=f"document-{suffix.lstrip('.') or 'file'}", kind="engagement-document",
        known_artifacts=known_documents,
    )
    extracted_name = ""
    try:
        extracted_text, warning = extract_engagement_document_text(data, suffix) if suffix in {".pdf", ".docx", ".txt", ".md"} else ("", "")
        if extracted_text:
            extracted_name = f"{Path(item['storedFilename']).stem}.txt"
            atomic_write_text(root / extracted_name, extracted_text)
        item.update({
            "category": category_key, "kind": "engagement-document",
            "displayName": str(display_name or Path(original).stem).strip() or "Engagement Document",
            "extractedFilename": extracted_name, "notes": str(notes or "").strip(),
            "extractedText": extracted_text, "extractionWarning": warning,
        })
        return item
    except Exception:
        rollback_uncommitted_artifact(root, item, [extracted_name])
        raise


def engagement_document_path(lab_name, stored_filename):
    root = general_documents_root(lab_name).resolve()
    target = (root / safe_name(stored_filename)).resolve()
    if not path_is_inside(target, root) or not target.exists() or not target.is_file():
        raise FileNotFoundError("Engagement document not found")
    return target


def delete_engagement_document(lab_name, stored_filename, extracted_filename="", stored_filenames=None):
    deleted = False
    names = [stored_filename, extracted_filename, *(stored_filenames or [])]
    for stored in list(names):
        if stored and Path(str(stored)).suffix.lower() != ".txt":
            names.append(f"{Path(str(stored)).stem}.txt")
    for name in dict.fromkeys(str(value or "") for value in names):
        if not name: continue
        try: path = engagement_document_path(lab_name, name)
        except FileNotFoundError: continue
        path.unlink(); deleted = True
    return deleted


def scan_artifacts_root(lab_name, host_key):
    if not str(lab_name or "").strip():
        raise ValueError("Select an engagement before storing scan artifacts")
    state = STORE.load_project(lab_name) or {}
    if host_key == "__engagement__" and state:
        return contained(STORE.project_root(lab_name), "scans")
    hosts = state.get("hosts", {})
    host = hosts.get(host_key) or next((h for h in hosts.values() if h.get("ip") == host_key or h.get("id") == host_key), None)
    if not host:
        raise ValueError("Save this host before storing scan artifacts")
    return contained(STORE.host_root(lab_name, host), "scans")


WEB_DISCOVERY_JSON_OBJECTIVES = frozenset({
    "web-content-discovery",
    "web-targeted-file-discovery",
    "web-common-backup-discovery",
    "web-custom-backup-discovery",
    "web-deeper-path-discovery",
    "api-endpoint-discovery",
})
WHATWEB_JSON_OBJECTIVE = "web-technology-fingerprinting"
MAX_WHATWEB_JSON_RECORDS = 5000
MAX_WHATWEB_JSON_PLUGINS = 250
MAX_WHATWEB_JSON_DEPTH = 16
MAX_WHATWEB_JSON_NODES = 250000
MAX_WHATWEB_JSON_STRING = 1024 * 1024
WHATWEB_VERBOSE_JSON_ERROR = (
    "WhatWeb verbose JSON is not supported. Generate standard output with --log-json."
)


def validate_web_discovery_json(data):
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise ValueError("Web-discovery JSON must be valid UTF-8 JSON") from exc
    if not text.strip():
        raise ValueError("Web-discovery JSON must be valid UTF-8 JSON")
    try:
        payload = json.loads(text)
    except json.JSONDecodeError as exc:
        lines = [line.strip() for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n") if line.strip()]
        if not lines or len(lines) > 250000:
            raise ValueError("Web-discovery upload must contain valid UTF-8 JSON, FFUF JSON, or Feroxbuster JSON lines") from exc
        try:
            payload = [json.loads(line) for line in lines]
        except json.JSONDecodeError as line_error:
            raise ValueError("Web-discovery upload must contain valid UTF-8 JSON, FFUF JSON, or Feroxbuster JSON lines") from line_error

    if isinstance(payload, dict) and isinstance(payload.get("results"), list):
        return payload

    rows = payload if isinstance(payload, list) else [payload]
    if not rows or len(rows) > 250000 or not all(isinstance(row, dict) for row in rows):
        raise ValueError("Web-discovery JSON must contain an ffuf results array or Feroxbuster records")
    ferox_types = {"response", "statistics", "stats", "configuration", "config", "summary", "error"}
    recognized = False
    for row in rows:
        record_type = str(row.get("type") or "").strip().lower()
        if record_type in ferox_types:
            recognized = True
            if record_type == "response" and not str(row.get("url") or row.get("target_url") or "").strip():
                raise ValueError("Feroxbuster response records must contain an exact URL")
        elif any(key in row for key in ("total_requests", "requests", "scan_complete", "completed")):
            recognized = True
    if not recognized:
        raise ValueError("Web-discovery JSON must contain an ffuf results array or Feroxbuster records")
    return payload


def exact_whatweb_target(value):
    if not isinstance(value, str):
        return False
    target = value.strip()
    if not target or len(target) > 8192 or re.search(r"[\x00-\x20\x7f]", target):
        return False
    try:
        parsed = urlsplit(target)
        host = parsed.hostname or ""
        port = parsed.port
    except ValueError:
        return False
    if parsed.scheme.lower() not in {"http", "https"} or not parsed.netloc:
        return False
    if parsed.username is not None or parsed.password is not None or not host or len(host) > 253:
        return False
    if port is not None and not 1 <= port <= 65535:
        return False
    if ":" in host:
        return bool(re.fullmatch(r"[0-9a-f:.]+", host, re.IGNORECASE) and ":" in host)
    labels = host.rstrip(".").split(".")
    return bool(
        labels
        and all(
            label
            and len(label) <= 63
            and re.fullmatch(r"[a-z0-9](?:[a-z0-9-]*[a-z0-9])?", label, re.IGNORECASE)
            for label in labels
        )
    )


def whatweb_verbose_positional_record(value):
    if not isinstance(value, list) or len(value) < 3:
        return False
    if not exact_whatweb_target(value[0]) or isinstance(value[1], bool):
        return False
    try:
        status = int(value[1])
    except (TypeError, ValueError):
        return False
    return 100 <= status <= 599 and isinstance(value[2], list)


def validate_whatweb_json_bounds(payload):
    nodes = 0
    stack = [(payload, 0)]
    while stack:
        value, depth = stack.pop()
        if depth > MAX_WHATWEB_JSON_DEPTH:
            raise ValueError("WhatWeb JSON exceeds safe nested material limits.")
        nodes += 1
        if nodes > MAX_WHATWEB_JSON_NODES:
            raise ValueError("WhatWeb JSON exceeds safe nested material limits.")
        if isinstance(value, dict):
            stack.extend((item, depth + 1) for item in value.values())
        elif isinstance(value, list):
            stack.extend((item, depth + 1) for item in value)
        elif isinstance(value, str) and len(value) > MAX_WHATWEB_JSON_STRING:
            raise ValueError("WhatWeb JSON exceeds safe nested material limits.")


def parse_standard_whatweb_json(data):
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise ValueError("The WhatWeb JSON is invalid or malformed.") from exc
    if not text.strip():
        raise ValueError("The WhatWeb JSON is invalid or malformed.")

    try:
        payload = json.loads(text)
    except json.JSONDecodeError as exc:
        lines = [line.strip() for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n") if line.strip()]
        parsed_lines = []
        if len(lines) > MAX_WHATWEB_JSON_RECORDS:
            raise ValueError(
                f"WhatWeb JSON exceeds the {MAX_WHATWEB_JSON_RECORDS}-record limit."
            ) from exc
        if len(lines) >= 2:
            try:
                parsed_lines = [json.loads(line) for line in lines]
            except json.JSONDecodeError:
                parsed_lines = []
        if parsed_lines and all(whatweb_verbose_positional_record(row) for row in parsed_lines):
            raise ValueError(WHATWEB_VERBOSE_JSON_ERROR) from exc
        if not parsed_lines or not all(isinstance(row, dict) for row in parsed_lines):
            raise ValueError("The WhatWeb JSON is invalid or malformed.") from exc
        payload = parsed_lines

    if whatweb_verbose_positional_record(payload) or (
        isinstance(payload, list)
        and payload
        and all(whatweb_verbose_positional_record(row) for row in payload)
    ):
        raise ValueError(WHATWEB_VERBOSE_JSON_ERROR)

    validate_whatweb_json_bounds(payload)
    if isinstance(payload, list):
        rows = payload
    elif isinstance(payload, dict) and isinstance(payload.get("results"), list):
        identity = " ".join(
            str(payload.get(key) or "")
            for key in ("tool", "scanner", "generator", "command", "commandline")
        )
        if not re.search(r"whatweb", identity, re.IGNORECASE):
            raise ValueError("The WhatWeb JSON is invalid or malformed.")
        rows = payload["results"]
    elif isinstance(payload, dict) and "target" in payload:
        rows = [payload]
    else:
        raise ValueError("The WhatWeb JSON is invalid or malformed.")

    if not rows:
        raise ValueError("WhatWeb JSON contains no exact HTTP(S) target record.")
    if len(rows) > MAX_WHATWEB_JSON_RECORDS:
        raise ValueError(f"WhatWeb JSON exceeds the {MAX_WHATWEB_JSON_RECORDS}-record limit.")
    if not all(isinstance(row, dict) for row in rows):
        raise ValueError("The WhatWeb JSON is invalid or malformed.")
    if not any(exact_whatweb_target(row.get("target") or row.get("url")) for row in rows):
        raise ValueError("WhatWeb JSON contains no exact HTTP(S) target record.")

    for index, row in enumerate(rows, 1):
        if not exact_whatweb_target(row.get("target") or row.get("url")):
            raise ValueError(
                f"WhatWeb JSON result {index} has no exact HTTP(S) target record."
            )
        plugins = row.get("plugins")
        if not isinstance(plugins, dict):
            raise ValueError(f"WhatWeb JSON result {index} has no plugin object.")
        if len(plugins) > MAX_WHATWEB_JSON_PLUGINS:
            raise ValueError(
                f"WhatWeb JSON result {index} exceeds the {MAX_WHATWEB_JSON_PLUGINS}-plugin limit."
            )
        for request_key in ("request_config", "requestConfig", "request"):
            if request_key in row and not isinstance(row[request_key], dict):
                raise ValueError(f"WhatWeb JSON result {index} has invalid request metadata.")
        status_marker = object()
        status = row.get("http_status", row.get("status", status_marker))
        if status is not status_marker:
            valid_status = False
            if not isinstance(status, bool):
                try:
                    numeric_status = float(status)
                    valid_status = numeric_status.is_integer() and 100 <= int(numeric_status) <= 599
                except (TypeError, ValueError, OverflowError):
                    valid_status = False
            if not valid_status:
                raise ValueError(f"WhatWeb JSON result {index} has an invalid HTTP status.")
    return rows


def upload_scan_artifact(lab_name, host_key, file_item, objective="", source="", logical_path="", known_artifacts=None, physical_source_path="", run_id=""):
    """Store one scan artifact using exact-byte identity and logical revision history.

    Existing metadata is supplied by the engagement state. The server verifies
    legacy stored files itself before using their hashes, so filename/size
    similarity never counts as duplicate proof.
    """
    if not file_item or not file_item.filename:
        raise ValueError("Choose a scan artifact to upload")
    original = Path(file_item.filename).name
    suffix = Path(original).suffix.lower()
    normalized_objective = str(objective or "").strip()
    auto_recon_raw = normalized_objective == "autorecon-raw-evidence"
    allowed_suffixes = {".xml", ".nmap", ".gnmap", ".txt", ".log"}
    if normalized_objective in WEB_DISCOVERY_JSON_OBJECTIVES | {WHATWEB_JSON_OBJECTIVE}:
        allowed_suffixes.add(".json")
    if suffix == ".json" and not auto_recon_raw and normalized_objective not in WEB_DISCOVERY_JSON_OBJECTIVES | {WHATWEB_JSON_OBJECTIVE}:
        raise ValueError("This JSON file is not supported for the selected import objective.")
    if not auto_recon_raw and suffix not in allowed_suffixes:
        if normalized_objective == WHATWEB_JSON_OBJECTIVE:
            raise ValueError("WhatWeb Technology artifacts must use JSON, text, or command-log files.")
        raise ValueError("Supported scan artifacts are XML, NMAP, GNMAP, text, command-log, and web-discovery JSON files")
    data = file_item.file.read()
    if not data and suffix == ".json" and normalized_objective == WHATWEB_JSON_OBJECTIVE:
        raise ValueError("The WhatWeb JSON is invalid or malformed.")
    if not data or len(data) > 25 * 1024 * 1024:
        raise ValueError("Scan artifact must be non-empty and no larger than 25 MiB")
    if suffix == ".json" and not auto_recon_raw:
        if normalized_objective == WHATWEB_JSON_OBJECTIVE:
            parse_standard_whatweb_json(data)
        else:
            validate_web_discovery_json(data)

    host_part = safe_name(host_key or "host") or "host"
    root = scan_artifacts_root(lab_name, host_key)
    root.mkdir(parents=True, exist_ok=True)
    digest = sha256_bytes(data)
    normalized_path = normalize_logical_path(logical_path or original, original)
    normalized_source = str(source or "Imported scan").strip() or "Imported scan"
    artifact_type = infer_artifact_type(original, normalized_objective)

    if isinstance(known_artifacts, str):
        try:
            known_artifacts = json.loads(known_artifacts or "[]")
        except json.JSONDecodeError as exc:
            raise ValueError("Existing artifact metadata is not valid JSON") from exc
    if not isinstance(known_artifacts, list):
        known_artifacts = []

    def resolve_existing_hash(stored_filename):
        path = scan_artifact_path(lab_name, host_key, stored_filename)
        return file_sha256(path)

    verified_known = [
        normalize_artifact(item, hash_resolver=resolve_existing_hash)
        for item in known_artifacts
        if isinstance(item, dict)
    ]
    outcome, existing = classify_artifact(
        verified_known,
        sha256=digest,
        logical_path=normalized_path,
        source=normalized_source,
        objective=normalized_objective,
        artifact_type=artifact_type,
    )

    stored_name = ""
    stored_new_file = outcome in {"new", "updated"}
    if stored_new_file:
        storage_suffix = suffix if re.fullmatch(r"\.[a-z0-9]{1,16}", suffix) else ".bin"
        stored_name = f"{storage_component(Path(original).stem)}--{uuid.uuid4().hex}{storage_suffix}"
        target = (root / stored_name).resolve()
        if not path_is_inside(target, root):
            raise ValueError("Invalid scan artifact filename")
        target.write_bytes(data)
    elif existing:
        matching = next((row for row in existing.get("revisions", []) if row.get("sha256") == digest), None)
        stored_name = str((matching or {}).get("storedFilename") or existing.get("storedFilename") or "")

    observed_at = datetime.datetime.now(datetime.timezone.utc).isoformat()
    artifact = merge_observation(
        existing,
        outcome=outcome,
        sha256=digest,
        stored_filename=stored_name,
        original_filename=original,
        logical_path=normalized_path,
        size=len(data),
        mime_type="application/octet-stream" if auto_recon_raw else (mimetypes.guess_type(original)[0] or "text/plain"),
        source=normalized_source,
        objective=normalized_objective,
        artifact_type=artifact_type,
        host_key=host_part,
        observed_at=observed_at,
        physical_source_path=str(physical_source_path or "")[:4096],
        run_id=str(run_id or "")[:200],
    )
    return {
        "artifact": artifact,
        "result": {
            "outcome": outcome,
            "revision": artifact.get("currentRevision", 1),
            "storedNewFile": stored_new_file,
            "reusedStoredFile": outcome == "unchanged",
            "sha256": digest,
        },
    }


def scan_artifact_path(lab_name, host_key, stored_filename):
    root = scan_artifacts_root(lab_name, host_key).resolve()
    target = (root / safe_name(stored_filename)).resolve()
    if not path_is_inside(target, root) or not target.exists() or not target.is_file():
        raise FileNotFoundError("Scan artifact not found")
    return target

def load_report_template_registry():
    REPORT_TEMPLATES_ROOT.mkdir(parents=True, exist_ok=True)
    registry = {"schemaVersion": 1, "defaultTemplateId": "", "templates": []}
    if REPORT_TEMPLATES_REGISTRY.exists():
        loaded = json.loads(REPORT_TEMPLATES_REGISTRY.read_text(encoding="utf-8"))
        if isinstance(loaded, dict):
            registry.update(loaded)
    if not isinstance(registry.get("templates"), list):
        registry["templates"] = []
    # One-time migration of the bundled OSCP template into app-owned storage.
    legacy = ROOT / "report_templates" / "OSCP-Exam-Report.docx"
    if legacy.exists() and not registry["templates"]:
        template_id = "built-in-oscp"
        stored = REPORT_TEMPLATES_ROOT / f"{template_id}.docx"
        if not stored.exists():
            shutil.copy2(legacy, stored)
        registry["templates"].append({
            "id": template_id, "displayName": "OSCP Exam Report", "filename": stored.name,
            "originalFilename": legacy.name, "createdAt": datetime.datetime.now(datetime.timezone.utc).isoformat(),
            "size": stored.stat().st_size, "builtIn": True
        })
        registry["defaultTemplateId"] = template_id
        atomic_write_json(REPORT_TEMPLATES_REGISTRY, registry)
    return registry


def save_report_template_registry(registry):
    atomic_write_json(REPORT_TEMPLATES_REGISTRY, registry)


def report_template_path(template_id):
    template_id = safe_name(template_id)
    registry = load_report_template_registry()
    item = next((x for x in registry["templates"] if x.get("id") == template_id), None)
    if not item:
        raise FileNotFoundError("Report template not found")
    path = (REPORT_TEMPLATES_ROOT / item["filename"]).resolve()
    if not path_is_inside(path, REPORT_TEMPLATES_ROOT) or not path.exists():
        raise FileNotFoundError("Stored report template file not found")
    return path, item


def validate_docx_bytes(data):
    if not data or len(data) > 25 * 1024 * 1024:
        raise ValueError("DOCX must be non-empty and smaller than 25 MB")
    try:
        with zipfile.ZipFile(BytesIO(data), "r") as zf:
            members = zf.infolist()
            if len(members) > MAX_DOCX_MEMBERS:
                raise ValueError("DOCX contains too many archive members")
            expanded = 0
            seen_names = set()
            for info in members:
                normalized = safe_import_zip_member(info)
                name_key = normalized.casefold()
                if name_key in seen_names:
                    raise ValueError(f"DOCX contains a duplicate member path: {normalized}")
                seen_names.add(name_key)
                expanded += info.file_size
                if expanded > MAX_DOCX_EXPANDED_BYTES:
                    raise ValueError("Expanded DOCX exceeds the safety limit")
                if info.file_size and info.file_size / max(info.compress_size, 1) > MAX_DOCX_COMPRESSION_RATIO:
                    raise ValueError("DOCX contains a suspiciously compressed member")
                lower_name = normalized.casefold()
                if (
                    lower_name == "word/vbaproject.bin"
                    or lower_name.startswith("word/activex/")
                    or lower_name.startswith("word/embeddings/")
                    or lower_name.startswith("customui/")
                ):
                    raise ValueError("DOCX templates may not contain macros, ActiveX controls, or embedded objects")
                if lower_name.endswith(".rels") and not info.is_dir():
                    if info.file_size > MAX_DOCX_RELATIONSHIP_BYTES:
                        raise ValueError("DOCX relationship metadata is too large")
                    try:
                        relationships = ET.fromstring(zf.read(info))
                    except ET.ParseError as exc:
                        raise ValueError("DOCX contains invalid relationship metadata") from exc
                    for relationship in relationships:
                        attributes = relationship.attrib
                        if str(attributes.get("TargetMode") or "").casefold() != "external":
                            continue
                        relationship_type = str(attributes.get("Type") or "").casefold()
                        if not relationship_type.endswith("/hyperlink"):
                            raise ValueError("DOCX contains a prohibited external relationship")
                        target = str(attributes.get("Target") or "").strip()
                        scheme = urlsplit(target).scheme.casefold()
                        if scheme not in {"http", "https", "mailto"}:
                            raise ValueError("DOCX contains an unsafe external hyperlink")
                elif lower_name.startswith("word/") and lower_name.endswith(".xml") and not info.is_dir():
                    if info.file_size > MAX_DOCX_XML_MEMBER_BYTES:
                        raise ValueError("DOCX XML metadata exceeds the safety limit")
                    try:
                        xml_root = ET.fromstring(zf.read(info))
                    except ET.ParseError as exc:
                        raise ValueError("DOCX contains invalid WordprocessingML") from exc
                    instructions = []
                    for element in xml_root.iter():
                        local_name = str(element.tag).rsplit("}", 1)[-1].casefold()
                        if local_name == "instrtext" and element.text:
                            instructions.append(element.text)
                        for attribute_name, attribute_value in element.attrib.items():
                            if str(attribute_name).rsplit("}", 1)[-1].casefold() == "instr":
                                instructions.append(str(attribute_value))
                    instruction_text = "".join(instructions)
                    if re.search(r"\b(?:DDEAUTO?|INCLUDETEXT|INCLUDEPICTURE|LINK|DATABASE)\b", instruction_text, re.I):
                        raise ValueError("DOCX contains a prohibited active field instruction")
                    for hyperlink in re.finditer(
                        r'\bHYPERLINK\s+(?:"([^"]+)"|([^\s]+))',
                        instruction_text,
                        re.I,
                    ):
                        target = str(hyperlink.group(1) or hyperlink.group(2) or "").strip()
                        if target.startswith("\\l"):
                            continue
                        if urlsplit(target).scheme.casefold() not in {"http", "https", "mailto"}:
                            raise ValueError("DOCX contains an unsafe hyperlink field")
            if "word/document.xml" not in zf.namelist() or zf.testzip() is not None:
                raise ValueError("Invalid DOCX template")
    except (RuntimeError, zipfile.BadZipFile):
        raise ValueError("Invalid DOCX template")


def upload_report_template(file_item, display_name=""):
    if not file_item or not file_item.filename:
        raise ValueError("Choose a DOCX template")
    if Path(file_item.filename).suffix.lower() != ".docx":
        raise ValueError("Only DOCX report templates are supported")
    data = file_item.file.read(); validate_docx_bytes(data)
    registry = load_report_template_registry()
    original = Path(file_item.filename).name
    label = str(display_name or Path(original).stem).strip() or "Report Template"
    item = store_versioned_artifact(
        data, root=REPORT_TEMPLATES_ROOT, original_filename=original,
        logical_path=f"templates/{original}", source="report-template", objective="global",
        artifact_type="report-template-docx", kind="report-template",
        known_artifacts=[{**row, "storedFilename": row.get("storedFilename") or row.get("filename", "")} for row in registry.get("templates", []) if not row.get("builtIn")],
    )
    existing = next((row for row in registry.get("templates", []) if row.get("id") == item.get("id")), None)
    item.update({"displayName": label, "filename": item.get("storedFilename", ""), "createdAt": item.get("firstSeenAt") or item.get("uploadedAt"), "builtIn": False})
    if existing:
        registry["templates"][registry["templates"].index(existing)] = item
    else:
        registry["templates"].append(item)
    if not registry.get("defaultTemplateId"):
        registry["defaultTemplateId"] = item["id"]
    try:
        save_report_template_registry(registry)
        return item
    except Exception:
        rollback_uncommitted_artifact(REPORT_TEMPLATES_ROOT, item)
        raise


def delete_report_template(template_id):
    registry = load_report_template_registry()
    item = next((x for x in registry["templates"] if x.get("id") == template_id), None)
    if not item:
        return False
    if item.get("builtIn"):
        raise ValueError("The built-in template cannot be deleted")
    filenames = [item.get("filename"), item.get("storedFilename"), *[row.get("storedFilename") for row in item.get("revisions", []) if isinstance(row, dict)]]
    registry["templates"] = [x for x in registry["templates"] if x.get("id") != template_id]
    if registry.get("defaultTemplateId") == template_id:
        registry["defaultTemplateId"] = registry["templates"][0]["id"] if registry["templates"] else ""
    save_report_template_registry(registry)
    for filename in dict.fromkeys(str(value or "") for value in filenames):
        if not filename: continue
        path = (REPORT_TEMPLATES_ROOT / safe_name(filename)).resolve()
        if path_is_inside(path, REPORT_TEMPLATES_ROOT) and path.exists() and path.is_file(): path.unlink()
    return True

def project_state_path(lab):
    return STORE.state_path(lab)


def initialize_persistence():
    """Use files exclusively; copy legacy data once without changing originals."""
    STORE.initialize()
    from legacy_import import import_legacy_engagements
    return import_legacy_engagements(STORE, DATA_ROOT)


def atomic_write_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_name(f".{path.name}.{uuid.uuid4().hex}.tmp")
    try:
        temp.write_text(json.dumps(value, indent=2), encoding="utf-8")
        os.replace(temp, path)
    finally:
        if temp.exists():
            temp.unlink()


def atomic_write_text(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_name(f".{path.name}.{uuid.uuid4().hex}.tmp")
    try:
        temp.write_text(str(value), encoding="utf-8")
        os.replace(temp, path)
    finally:
        if temp.exists():
            temp.unlink()


def atomic_write_bytes(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_name(f".{path.name}.{uuid.uuid4().hex}.tmp")
    try:
        temp.write_bytes(bytes(value))
        os.replace(temp, path)
    finally:
        if temp.exists():
            temp.unlink()


def load_command_notes_manifest():
    if not COMMAND_NOTES_MANIFEST_PATH.exists():
        raise FileNotFoundError("The bundled command note library is missing.")
    data = json.loads(COMMAND_NOTES_MANIFEST_PATH.read_text(encoding="utf-8"))
    notes = data.get("notes") if isinstance(data, dict) else None
    if not isinstance(notes, dict):
        raise ValueError("The command note manifest is invalid.")
    return data


def _safe_reference_relative_path(raw_path):
    relative = PurePosixPath(str(raw_path or "").replace("\\", "/"))
    if relative.is_absolute() or not relative.parts or any(part in {"", ".", ".."} for part in relative.parts):
        raise ValueError("The reference note path is invalid")
    if relative.suffix.lower() != ".md":
        raise ValueError("Reference notes must be Markdown files")
    return relative


def _reference_note_string_list(entry, field, *, allowed=None, required=False):
    value = entry.get(field)
    if not isinstance(value, list):
        raise ValueError(f"Reference note {entry.get('id') or '<unknown>'} field {field} must be an array")
    normalized = []
    for raw in value:
        token = str(raw or "").strip().lower()
        if not token or not REFERENCE_NOTE_TOKEN_RE.fullmatch(token):
            raise ValueError(f"Reference note {entry.get('id') or '<unknown>'} contains an invalid {field} value")
        if allowed is not None and token not in allowed:
            raise ValueError(f"Reference note {entry.get('id') or '<unknown>'} uses unsupported {field} value: {token}")
        if token not in normalized:
            normalized.append(token)
    if required and not normalized:
        raise ValueError(f"Reference note {entry.get('id') or '<unknown>'} field {field} cannot be empty")
    return normalized


def _validate_reference_note_metadata(entry):
    note_id = str(entry.get("id") or "").strip()
    if not note_id:
        raise ValueError("Reference note is missing id")
    if not str(entry.get("title") or "").strip():
        raise ValueError(f"Reference note {note_id} is missing title")

    stage = str(entry.get("stage") or "").strip().lower()
    if stage not in REFERENCE_NOTE_STAGES:
        raise ValueError(f"Reference note {note_id} uses unsupported stage: {stage or '<empty>'}")
    scope = str(entry.get("scope") or "").strip().lower()
    if scope not in REFERENCE_NOTE_SCOPES:
        raise ValueError(f"Reference note {note_id} uses unsupported scope: {scope or '<empty>'}")
    trigger_mode = str(entry.get("triggerMode") or "").strip().lower()
    if trigger_mode not in REFERENCE_NOTE_TRIGGER_MODES:
        raise ValueError(f"Reference note {note_id} uses unsupported triggerMode: {trigger_mode or '<empty>'}")
    os_value = str(entry.get("os") or "both").strip().lower()
    if os_value not in REFERENCE_NOTE_OS_VALUES:
        raise ValueError(f"Reference note {note_id} uses unsupported os value: {os_value or '<empty>'}")

    task_id = str(entry.get("taskId") or "").strip().lower()
    task_role = str(entry.get("taskRole") or "").strip().lower()
    if task_id and not REFERENCE_NOTE_TOKEN_RE.fullmatch(task_id):
        raise ValueError(f"Reference note {note_id} contains an invalid taskId")
    if task_id and task_role not in REFERENCE_NOTE_TASK_ROLES:
        raise ValueError(f"Reference note {note_id} uses unsupported taskRole: {task_role or '<empty>'}")
    if task_role and not task_id:
        raise ValueError(f"Reference note {note_id} defines taskRole without taskId")
    task_class = str(entry.get("taskClass") or "").strip().lower()
    if task_class and task_class not in REFERENCE_NOTE_TASK_CLASSES:
        raise ValueError(f"Reference note {note_id} uses unsupported taskClass: {task_class}")
    completion_rule = str(entry.get("completionRule") or "").strip().lower()
    if completion_rule and completion_rule not in REFERENCE_NOTE_COMPLETION_RULES:
        raise ValueError(f"Reference note {note_id} uses unsupported completionRule: {completion_rule}")

    modules = _reference_note_string_list(entry, "modules", allowed=REFERENCE_NOTE_MODULES, required=True)
    services = _reference_note_string_list(entry, "services")
    tags = _reference_note_string_list(entry, "tags")
    activation_events = _reference_note_string_list(entry, "activationEvents", allowed=REFERENCE_NOTE_ACTIVATION_EVENTS)
    if "assetGroupTypes" not in entry:
        entry["assetGroupTypes"] = []
    asset_group_types = _reference_note_string_list(entry, "assetGroupTypes", allowed=REFERENCE_ASSET_GROUP_TYPES)
    evidence_trigger = str(entry.get("evidenceTrigger") or "").strip().lower()
    if evidence_trigger and evidence_trigger not in REFERENCE_NOTE_EVIDENCE_TRIGGERS:
        raise ValueError(f"Reference note {note_id} uses unsupported evidenceTrigger: {evidence_trigger}")
    if evidence_trigger:
        if scope != "access":
            raise ValueError(f"Reference note {note_id} uses evidenceTrigger outside access scope")
        binary_names = _reference_note_string_list(entry, "binaryNames", required=True)
    else:
        binary_names = []

    ports = entry.get("ports")
    if not isinstance(ports, list):
        raise ValueError(f"Reference note {note_id} field ports must be an array")
    normalized_ports = []
    for raw in ports:
        if isinstance(raw, bool) or not isinstance(raw, int) or not 1 <= raw <= 65535:
            raise ValueError(f"Reference note {note_id} contains an invalid port")
        if raw not in normalized_ports:
            normalized_ports.append(raw)

    for field in ("requiresShell", "requiresCreds", "requiresDomain"):
        if not isinstance(entry.get(field), bool):
            raise ValueError(f"Reference note {note_id} field {field} must be true or false")

    if trigger_mode == "asset-group-type" and not asset_group_types:
        raise ValueError(f"Reference note {note_id} uses asset-group-type triggerMode without assetGroupTypes")
    if trigger_mode == "service" and not services:
        raise ValueError(f"Reference note {note_id} uses service triggerMode without services")
    if trigger_mode == "port" and not normalized_ports:
        raise ValueError(f"Reference note {note_id} uses port triggerMode without ports")
    if trigger_mode == "domain" and not entry.get("requiresDomain"):
        raise ValueError(f"Reference note {note_id} uses domain triggerMode without requiresDomain")
    if scope == "asset-group" and trigger_mode not in {"asset-group-type", "domain", "credential", "event", "manual"}:
        raise ValueError(f"Reference note {note_id} has incompatible asset-group scope and {trigger_mode} triggerMode")
    if scope == "service" and trigger_mode not in {"service", "port", "credential", "event", "manual"}:
        raise ValueError(f"Reference note {note_id} has incompatible service scope and {trigger_mode} triggerMode")
    if scope == "access" and trigger_mode not in {"access", "credential", "event", "manual"}:
        raise ValueError(f"Reference note {note_id} has incompatible access scope and {trigger_mode} triggerMode")
    if scope == "library-only" and trigger_mode != "manual":
        raise ValueError(f"Reference note {note_id} must use manual triggerMode for library-only scope")
    if scope == "reporting" and stage != "reporting":
        raise ValueError(f"Reference note {note_id} uses reporting scope outside reporting stage")

    entry["stage"] = stage
    entry["scope"] = scope
    entry["triggerMode"] = trigger_mode
    entry["os"] = os_value
    entry["modules"] = modules
    entry["services"] = services
    entry["tags"] = tags
    entry["ports"] = sorted(normalized_ports)
    entry["activationEvents"] = activation_events
    entry["assetGroupTypes"] = asset_group_types
    if evidence_trigger:
        entry["evidenceTrigger"] = evidence_trigger
        entry["binaryNames"] = binary_names
    else:
        entry.pop("evidenceTrigger", None)
        entry.pop("binaryNames", None)
    if task_id:
        entry["taskId"] = task_id
        entry["taskRole"] = task_role
    else:
        entry.pop("taskId", None)
        entry.pop("taskRole", None)
    if task_class:
        entry["taskClass"] = task_class
    if completion_rule:
        entry["completionRule"] = completion_rule
    return entry


def _reference_pack_manifest(pack_root):
    manifest_path = Path(pack_root) / REFERENCE_PACK_MANIFEST
    if not manifest_path.exists() or not manifest_path.is_file():
        raise ValueError("Reference pack is missing index.json")
    if manifest_path.stat().st_size > 8 * 1024 * 1024:
        raise ValueError("Reference pack manifest is too large")
    try:
        data = json.loads(manifest_path.read_text(encoding="utf-8"))
    except Exception as exc:
        raise ValueError(f"Reference pack manifest is invalid JSON: {exc}") from exc
    if not isinstance(data, dict):
        raise ValueError("Reference pack manifest must be an object")
    if str(data.get("packType") or "") != "aeros-reference-notes":
        raise ValueError("Unsupported reference pack type")
    if int(data.get("schemaVersion") or 0) != 1:
        raise ValueError("Unsupported reference pack schema version")
    raw_pack_id = str(data.get("packId") or "").strip()
    if not raw_pack_id:
        raise ValueError("Reference pack is missing packId")
    pack_id = safe_name(raw_pack_id).lower()
    notes = data.get("notes")
    if not isinstance(notes, dict):
        raise ValueError("Reference pack manifest is missing notes")
    normalized = {}
    for note_id, raw in notes.items():
        if not isinstance(raw, dict):
            raise ValueError(f"Reference note {note_id} must be an object")
        entry = dict(raw)
        runtime_id = str(entry.get("id") or note_id).strip()
        if not runtime_id:
            raise ValueError(f"Reference note {note_id} is missing id")
        if runtime_id in normalized:
            raise ValueError(f"Reference pack contains duplicate runtime note ID: {runtime_id}")
        entry["id"] = runtime_id
        _validate_reference_note_metadata(entry)
        relative = _safe_reference_relative_path(entry.get("path") or entry.get("notePath"))
        note_path = (Path(pack_root) / Path(*relative.parts)).resolve()
        if not path_is_inside(note_path, Path(pack_root).resolve()):
            raise ValueError("Reference note path escaped the installed pack")
        if not note_path.exists() or not note_path.is_file():
            raise ValueError(f"Reference note is missing: {relative.as_posix()}")
        if note_path.stat().st_size > MAX_REFERENCE_NOTE_BYTES:
            raise ValueError(f"Reference note is too large: {relative.as_posix()}")
        entry["id"] = runtime_id
        entry["notePath"] = relative.as_posix()
        entry["path"] = relative.as_posix()
        entry["source"] = "aeros-reference-pack"
        entry["packId"] = pack_id
        entry["packName"] = str(data.get("displayName") or data.get("library") or pack_id)
        normalized[runtime_id] = entry
    return {
        **data,
        "packId": pack_id,
        "displayName": str(data.get("displayName") or data.get("library") or pack_id),
        "version": str(data.get("version") or "V1"),
        "count": len(normalized),
        "notes": normalized,
    }


def list_reference_packs(include_notes=False):
    REFERENCE_PACKS_ROOT.mkdir(parents=True, exist_ok=True)
    packs = []
    for pack_root in sorted((p for p in REFERENCE_PACKS_ROOT.iterdir() if p.is_dir()), key=lambda p: p.name.lower()):
        try:
            manifest = _reference_pack_manifest(pack_root)
            row = {
                "packId": manifest["packId"],
                "displayName": manifest["displayName"],
                "version": manifest["version"],
                "contentRevision": str(manifest.get("contentRevision") or ""),
                "count": manifest["count"],
                "installedPath": str(pack_root),
            }
            if include_notes:
                row["notes"] = manifest["notes"]
            packs.append(row)
        except Exception as exc:
            packs.append({
                "packId": pack_root.name,
                "displayName": pack_root.name,
                "version": "",
                "count": 0,
                "installedPath": str(pack_root),
                "invalid": True,
                "error": str(exc),
                **({"notes": {}} if include_notes else {}),
            })
    return packs


def load_reference_notes_manifest():
    packs = list_reference_packs(include_notes=True)
    merged = {}
    valid_packs = []
    invalid_packs = []
    for pack in packs:
        if pack.get("invalid"):
            invalid_packs.append({k: v for k, v in pack.items() if k != "notes"})
            continue
        collisions = sorted(set(merged).intersection(pack.get("notes", {})))
        if collisions:
            invalid_packs.append({
                "packId": pack["packId"],
                "displayName": pack["displayName"],
                "error": f"Duplicate note IDs conflict with another installed pack: {', '.join(collisions[:5])}",
            })
            continue
        merged.update(pack.get("notes", {}))
        valid_packs.append({k: v for k, v in pack.items() if k != "notes"})
    library = " + ".join(pack["displayName"] for pack in valid_packs) or "No optional reference pack installed"
    return {
        "schemaVersion": 1,
        "library": library,
        "count": len(merged),
        "notes": merged,
        "packs": valid_packs,
        "invalidPacks": invalid_packs,
    }


def reference_note_entry(note_id):
    note_id = str(note_id or "").strip()
    if not note_id:
        raise ValueError("noteId is required")
    entry = load_reference_notes_manifest().get("notes", {}).get(note_id)
    if not isinstance(entry, dict):
        raise ValueError("Unknown reference note")
    relative = _safe_reference_relative_path(entry.get("path") or entry.get("notePath"))
    pack_id = safe_name(str(entry.get("packId") or "")).lower()
    if not pack_id:
        raise ValueError("Reference note is missing pack ownership")
    pack_root = (REFERENCE_PACKS_ROOT / pack_id).resolve()
    if not path_is_inside(pack_root, REFERENCE_PACKS_ROOT.resolve()):
        raise ValueError("Reference pack path escaped the library")
    return entry, relative, pack_root


def reference_note_custom_path(profile_name, entry, relative):
    profile_root = (CUSTOM_REFERENCE_NOTES_ROOT / command_note_profile_key(profile_name)).resolve()
    pack_id = safe_name(str(entry.get("packId") or "")).lower()
    if not pack_id:
        raise ValueError("Reference note is missing pack ownership")
    path = (profile_root / pack_id / Path(*relative.parts)).resolve()
    if not path_is_inside(path, profile_root):
        raise ValueError("Custom reference note path escaped the profile library")
    return path


def read_reference_note(note_id, profile_name=""):
    entry, relative, pack_root = reference_note_entry(note_id)
    path = (pack_root / Path(*relative.parts)).resolve()
    if not path_is_inside(path, pack_root):
        raise ValueError("Reference note path escaped the installed pack")
    if not path.exists() or not path.is_file():
        raise FileNotFoundError("The installed reference note was not found")
    if path.stat().st_size > MAX_REFERENCE_NOTE_BYTES:
        raise ValueError("The reference note is too large")
    default_content = path.read_text(encoding="utf-8", errors="replace")
    custom_content = None
    if str(profile_name or "").strip():
        custom_path = reference_note_custom_path(profile_name, entry, relative)
        custom_content = custom_path.read_text(encoding="utf-8", errors="replace") if custom_path.exists() else None
    return {
        "entry": entry,
        "title": str(entry.get("title") or path.stem),
        "path": str(entry.get("notePath") or entry.get("path") or relative.as_posix()),
        "defaultContent": default_content,
        "customContent": custom_content,
        "customized": custom_content is not None,
        "content": custom_content if custom_content is not None else default_content,
    }


def save_custom_reference_note(note_id, profile_name, content):
    if not isinstance(content, str):
        raise ValueError("content must be text")
    encoded = content.encode("utf-8")
    if len(encoded) > MAX_REFERENCE_NOTE_BYTES:
        raise ValueError("The reference note is too large")
    entry, relative, _ = reference_note_entry(note_id)
    path = reference_note_custom_path(profile_name, entry, relative)
    atomic_write_text(path, content)
    return {"entry": entry, "path": str(path), "bytes": len(encoded)}


def reset_custom_reference_note(note_id, profile_name):
    entry, relative, _ = reference_note_entry(note_id)
    path = reference_note_custom_path(profile_name, entry, relative)
    existed = path.exists()
    if existed:
        path.unlink()
        profile_root = (CUSTOM_REFERENCE_NOTES_ROOT / command_note_profile_key(profile_name)).resolve()
        parent = path.parent
        while parent != profile_root and path_is_inside(parent, profile_root):
            try:
                parent.rmdir()
            except OSError:
                break
            parent = parent.parent
    return {"entry": entry, "removed": existed}


def import_reference_pack(file_item):
    if not file_item or not file_item.filename:
        raise ValueError("Choose an AEROS reference pack")
    data = file_item.file.read()
    if not data or len(data) > MAX_REFERENCE_PACK_BYTES:
        raise ValueError(f"Reference pack must be non-empty and smaller than {MAX_REFERENCE_PACK_BYTES // (1024 * 1024)} MB")
    if data[:4] != b"PK\x03\x04":
        raise ValueError("Reference pack is not a valid ZIP-based AEROS pack")
    try:
        archive = zipfile.ZipFile(BytesIO(data), "r")
    except zipfile.BadZipFile as exc:
        raise ValueError("Reference pack is not a valid ZIP archive") from exc
    staging_root = DATA_ROOT / "reference-pack-staging" / uuid.uuid4().hex
    staging_root.mkdir(parents=True, exist_ok=False)
    try:
        with archive:
            members = archive.infolist()
            if len(members) > MAX_REFERENCE_PACK_MEMBERS:
                raise ValueError(f"Reference pack contains too many files (maximum {MAX_REFERENCE_PACK_MEMBERS})")
            seen = set()
            total = 0
            for info in members:
                normalized = safe_recon_zip_member(info)
                if normalized in seen:
                    raise ValueError(f"Reference pack contains a duplicate path: {normalized}")
                seen.add(normalized)
                if info.is_dir():
                    continue
                suffix = PurePosixPath(normalized).suffix.lower()
                if normalized != REFERENCE_PACK_MANIFEST and suffix != ".md":
                    raise ValueError(f"Reference pack contains an unsupported file: {normalized}")
                if suffix in {".zip", ".7z", ".rar"}:
                    raise ValueError(f"Nested archives are not supported: {normalized}")
                if info.file_size <= 0:
                    raise ValueError(f"Reference pack contains an empty file: {normalized}")
                if info.file_size > MAX_REFERENCE_NOTE_BYTES and normalized != REFERENCE_PACK_MANIFEST:
                    raise ValueError(f"Reference note exceeds the size limit: {normalized}")
                total += info.file_size
                if total > MAX_REFERENCE_PACK_TOTAL_BYTES:
                    raise ValueError("Expanded reference pack exceeds the size limit")
                if info.compress_size == 0 or info.file_size / max(info.compress_size, 1) > MAX_REFERENCE_PACK_COMPRESSION_RATIO:
                    raise ValueError(f"Suspicious compression ratio in reference pack: {normalized}")
                target = (staging_root / Path(*PurePosixPath(normalized).parts)).resolve()
                if not path_is_inside(target, staging_root.resolve()):
                    raise ValueError("Reference pack extraction escaped staging")
                target.parent.mkdir(parents=True, exist_ok=True)
                with archive.open(info, "r") as source, target.open("wb") as destination:
                    shutil.copyfileobj(source, destination)
            bad_member = archive.testzip()
            if bad_member:
                raise ValueError(f"Reference pack CRC check failed: {bad_member}")
        manifest = _reference_pack_manifest(staging_root)
        pack_id = manifest["packId"]
        final_root = (REFERENCE_PACKS_ROOT / pack_id).resolve()
        if not path_is_inside(final_root, REFERENCE_PACKS_ROOT.resolve()):
            raise ValueError("Invalid reference pack ID")
        REFERENCE_PACKS_ROOT.mkdir(parents=True, exist_ok=True)
        backup_root = None
        if final_root.exists():
            backup_root = final_root.with_name(f".{final_root.name}.backup-{uuid.uuid4().hex}")
            os.replace(final_root, backup_root)
        try:
            os.replace(staging_root, final_root)
        except Exception:
            if backup_root and backup_root.exists() and not final_root.exists():
                os.replace(backup_root, final_root)
            raise
        if backup_root and backup_root.exists():
            shutil.rmtree(backup_root, ignore_errors=True)
        return {
            "packId": pack_id,
            "displayName": manifest["displayName"],
            "version": manifest["version"],
            "contentRevision": str(manifest.get("contentRevision") or ""),
            "count": manifest["count"],
            "installedPath": str(final_root),
            "replaced": backup_root is not None,
        }
    finally:
        if staging_root.exists():
            shutil.rmtree(staging_root, ignore_errors=True)


def delete_reference_pack(pack_id):
    raw_pack_id = str(pack_id or "").strip()
    if not raw_pack_id:
        raise ValueError("packId is required")
    pack_id = safe_name(raw_pack_id).lower()
    target = (REFERENCE_PACKS_ROOT / pack_id).resolve()
    if not path_is_inside(target, REFERENCE_PACKS_ROOT.resolve()):
        raise ValueError("Invalid reference pack ID")
    if not target.exists():
        return False
    shutil.rmtree(target)
    return True


def command_note_entry(note_id):
    note_id = str(note_id or "").strip()
    if not note_id:
        raise ValueError("noteId is required")
    entry = load_command_notes_manifest().get("notes", {}).get(note_id)
    if not isinstance(entry, dict):
        raise ValueError("Unknown command note")
    relative = PurePosixPath(str(entry.get("path") or ""))
    if relative.is_absolute() or not relative.parts or any(part in {"", ".", ".."} for part in relative.parts):
        raise ValueError("The command note path is invalid")
    if relative.suffix.lower() != ".md":
        raise ValueError("Command notes must be Markdown files")
    return entry, relative


def command_note_default_path(relative):
    path = (COMMAND_NOTES_ROOT / Path(*relative.parts)).resolve()
    if not path_is_inside(path, COMMAND_NOTES_ROOT.resolve()):
        raise ValueError("Command note path escaped the library")
    return path


def command_note_profile_key(profile_name):
    name = str(profile_name or "").strip()
    if not name:
        raise ValueError("profile is required")
    visible = safe_name(name)[:72] or "profile"
    digest = hashlib.sha256(name.encode("utf-8")).hexdigest()[:12]
    return f"{visible}-{digest}"


def command_note_custom_path(profile_name, relative):
    profile_root = (CUSTOM_COMMAND_NOTES_ROOT / command_note_profile_key(profile_name)).resolve()
    path = (profile_root / Path(*relative.parts)).resolve()
    if not path_is_inside(path, profile_root):
        raise ValueError("Custom command note path escaped the profile library")
    return path


def read_command_note(note_id, profile_name):
    entry, relative = command_note_entry(note_id)
    default_path = command_note_default_path(relative)
    if not default_path.exists() or not default_path.is_file():
        raise FileNotFoundError("The bundled command note was not found")
    default_content = default_path.read_text(encoding="utf-8")
    custom_path = command_note_custom_path(profile_name, relative)
    custom_content = custom_path.read_text(encoding="utf-8") if custom_path.exists() else None
    return {
        "entry": entry,
        "defaultContent": default_content,
        "customContent": custom_content,
        "customized": custom_content is not None,
    }


def save_custom_command_note(note_id, profile_name, content):
    if not isinstance(content, str):
        raise ValueError("content must be text")
    encoded = content.encode("utf-8")
    if len(encoded) > MAX_COMMAND_NOTE_BYTES:
        raise ValueError("The command note is too large")
    entry, relative = command_note_entry(note_id)
    path = command_note_custom_path(profile_name, relative)
    atomic_write_text(path, content)
    return {"entry": entry, "path": str(path), "bytes": len(encoded)}


def reset_custom_command_note(note_id, profile_name):
    entry, relative = command_note_entry(note_id)
    path = command_note_custom_path(profile_name, relative)
    existed = path.exists()
    if existed:
        path.unlink()
        parent = path.parent
        profile_root = (CUSTOM_COMMAND_NOTES_ROOT / command_note_profile_key(profile_name)).resolve()
        while parent != profile_root and path_is_inside(parent, profile_root):
            try:
                parent.rmdir()
            except OSError:
                break
            parent = parent.parent
    return {"entry": entry, "removed": existed}


def backup_lab_state(path, lab_name):
    """Create a rate-limited, bounded copy before replacing a lab mirror."""
    if not path.exists():
        return None
    backup_dir = contained(STORE.project_root(lab_name), ".aeros", "revisions")
    backup_dir.mkdir(parents=True, exist_ok=True)
    existing = sorted(backup_dir.glob("*.json"), key=lambda item: item.stat().st_mtime)
    if existing:
        age = datetime.datetime.now().timestamp() - existing[-1].stat().st_mtime
        if age < MIN_LAB_BACKUP_INTERVAL_SECONDS:
            return None
    stamp = datetime.datetime.now().strftime("%Y%m%d-%H%M%S-%f")
    backup_path = backup_dir / f"{stamp}.json"
    shutil.copy2(path, backup_path)
    existing.append(backup_path)
    for stale in existing[:-MAX_LAB_BACKUPS]:
        stale.unlink()
    return backup_path


def internal_host_dir(lab_name, host):
    if not STORE.project_exists(lab_name):
        raise ValueError("Save the engagement before storing host outputs")
    return STORE.host_root(lab_name, host)


def exploit_attempt_artifacts_root(lab_name, host, attempt_id):
    if not str(lab_name or "").strip():
        raise ValueError("labName is required")
    if not isinstance(host, dict) or not (host.get("ip") or host.get("hostname") or host.get("id")):
        raise ValueError("host identity is required")
    attempt_key = storage_component(attempt_id or "attempt")
    root = (internal_host_dir(lab_name, host) / "exploit-attempts" / attempt_key / "artifacts").resolve()
    if not path_is_inside(root, STORE.project_root(lab_name)):
        raise ValueError("Invalid exploit artifact path")
    return root


def _artifact_storage_view(item, root, relative_base):
    root = Path(root).resolve()
    relative_base = Path(relative_base)
    result = dict(item or {})
    revisions = []
    for row in result.get("revisions", []) if isinstance(result.get("revisions"), list) else []:
        revision = dict(row)
        stored = safe_name(revision.get("storedFilename")) if revision.get("storedFilename") else ""
        if stored:
            target = (root / stored).resolve()
            if path_is_inside(target, root):
                revision["storedPath"] = str(target)
                revision["relativePath"] = str(relative_base / stored).replace("\\", "/")
        revisions.append(revision)
    result["revisions"] = revisions
    stored = safe_name(result.get("storedFilename")) if result.get("storedFilename") else ""
    if stored:
        target = (root / stored).resolve()
        if path_is_inside(target, root):
            result["storedPath"] = str(target)
            result["relativePath"] = str(relative_base / stored).replace("\\", "/")
    return result


def peas_artifacts_root(lab_name, host):
    if not str(lab_name or "").strip():
        raise ValueError("labName is required")
    if not isinstance(host, dict) or not (host.get("ip") or host.get("hostname") or host.get("id")):
        raise ValueError("host identity is required")
    root = (internal_host_dir(lab_name, host) / "peas-output").resolve()
    if not path_is_inside(root, STORE.project_root(lab_name)):
        raise ValueError("Invalid PEAS artifact path")
    return root


def upload_peas_artifact(lab_name, host, access_context_id, tool, file_item, known_artifacts=None):
    if file_item is None or not getattr(file_item, "filename", ""):
        raise ValueError("Choose a LinPEAS or WinPEAS text output file")
    data = file_item.file.read()
    if not data:
        raise ValueError("PEAS output must not be empty")
    if len(data) > 25 * 1024 * 1024:
        raise ValueError("PEAS output must be smaller than 25 MB")
    original = Path(file_item.filename).name
    suffix = Path(original).suffix.lower()
    if suffix not in {"", ".txt", ".log", ".out", ".ansi"}:
        raise ValueError("PEAS output must be a text, log, out, or ANSI file")
    try:
        data.decode("utf-8")
    except UnicodeDecodeError:
        try:
            data.decode("cp1252")
        except UnicodeDecodeError as exc:
            raise ValueError("PEAS output is not readable text") from exc
    resolved_tool = str(tool or "unknown").strip().lower()
    if resolved_tool not in {"linpeas", "winpeas", "unknown"}:
        raise ValueError("Unsupported PEAS tool type")
    context_key = safe_name(access_context_id or "unscoped")
    root = peas_artifacts_root(lab_name, host)
    item = store_versioned_artifact(
        data, root=root, original_filename=original,
        logical_path=f"peas/{context_key}/{resolved_tool}", source="peas-import",
        objective=f"{resolved_tool}:{context_key}", artifact_type="peas-output-text", kind="peas-output",
        known_artifacts=parse_known_artifacts(known_artifacts or []),
        host_key=str(host.get("id") or host.get("ip") or host.get("hostname") or ""),
    )
    return _artifact_storage_view(item, root, Path("peas-output"))


def peas_artifact_from_stored_path(raw_path):
    path = Path(str(raw_path or "")).expanduser().resolve()
    owner = STORE.asset_owner(path)
    relative = path.relative_to(owner)
    if "peas-output" not in relative.parts:
        raise ValueError("Path is not a PEAS output artifact")
    if not path.exists() or not path.is_file():
        raise FileNotFoundError("PEAS output artifact not found")
    return path


def delete_peas_artifact_by_path(raw_path):
    path = peas_artifact_from_stored_path(raw_path)
    path.unlink()
    current = path.parent
    stop = STORE.asset_owner(path)
    while current != stop and path_is_inside(current, stop):
        try:
            current.rmdir()
        except OSError:
            break
        current = current.parent
    return True


def upload_exploit_attempt_artifact(lab_name, host, attempt_id, file_item, known_artifact=None):
    if file_item is None or not getattr(file_item, "filename", ""):
        raise ValueError("Choose the exact artifact used for the attempt")
    data = file_item.file.read()
    if not data:
        raise ValueError("Exploit artifact must not be empty")
    if len(data) > 50 * 1024 * 1024:
        raise ValueError("Exploit artifact must be smaller than 50 MB")
    original = Path(file_item.filename).name
    root = exploit_attempt_artifacts_root(lab_name, host, attempt_id)
    relative_base = Path("exploit-attempts") / safe_name(attempt_id or "attempt") / "artifacts"
    item = store_versioned_artifact(
        data, root=root, original_filename=original,
        logical_path="attempted-artifact", source="exploit-attempt", objective=safe_name(attempt_id or "attempt"),
        artifact_type="exploit-attempt-artifact", kind="exploit-attempt-artifact",
        known_artifacts=[known_artifact] if isinstance(known_artifact, dict) and known_artifact else [],
        host_key=str(host.get("id") or host.get("ip") or host.get("hostname") or ""),
    )
    return _artifact_storage_view(item, root, relative_base)


def exploit_attempt_artifact_from_stored_path(raw_path):
    path = Path(str(raw_path or "")).expanduser().resolve()
    owner = STORE.asset_owner(path)
    relative = path.relative_to(owner)
    if "exploit-attempts" not in relative.parts:
        raise ValueError("Path is not an exploit-attempt artifact")
    if not path.exists() or not path.is_file():
        raise FileNotFoundError("Exploit artifact not found")
    return path


def delete_exploit_attempt_artifact_by_path(raw_path):
    path = exploit_attempt_artifact_from_stored_path(raw_path)
    path.unlink()
    current = path.parent
    stop = STORE.asset_owner(path)
    while current != stop and path_is_inside(current, stop):
        try:
            current.rmdir()
        except OSError:
            break
        current = current.parent
    return True


def internal_screenshot_from_stored_path(raw_path):
    path = Path(str(raw_path or "")).expanduser().resolve()
    owner = STORE.asset_owner(path)
    relative = path.relative_to(owner)
    if "screenshots" not in relative.parts:
        raise ValueError("Path is not an internally stored screenshot")
    if path.suffix.lower() not in IMAGE_EXTS:
        raise ValueError("Path is not a supported screenshot")
    if not path.exists() or not path.is_file():
        raise FileNotFoundError("Screenshot not found")
    return path


def delete_internal_screenshot_by_path(raw_path):
    path = internal_screenshot_from_stored_path(raw_path)
    path.unlink()
    try:
        path.parent.rmdir()
    except OSError:
        pass
    return True





@serialized_project_mutation
def delete_project(lab_name, expected_revision=None):
    """Remove the catalog entry; leave engagement files available on disk."""
    return STORE.delete_project(str(lab_name), expected_revision=expected_revision)



def file_sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def archive_path_for_id(archive_id):
    archive_id = str(archive_id or "").strip()
    if (
        not archive_id
        or len(archive_id) > 200
        or archive_id in {".", ".."}
        or Path(archive_id).name != archive_id
        or safe_name(archive_id) != archive_id
    ):
        raise ValueError("Invalid archive identifier")
    path = (ARCHIVES_ROOT / f"{archive_id}.zip").resolve()
    if not path_is_inside(path, ARCHIVES_ROOT):
        raise ValueError("Invalid archive identifier")
    return path


@serialized_project_mutation
def archive_engagement(lab_name, remove_working=False):
    """Create a verified portable ZIP for one engagement."""
    lab_name = str(lab_name or "").strip()
    if not lab_name:
        raise ValueError("labName is required")
    state_path = project_state_path(lab_name)
    record = STORE.load_project_record(lab_name)
    if record is None:
        raise FileNotFoundError("Engagement not found")
    state = record["state"]
    snapshot_revision = int(record["revision"])
    stamp = datetime.datetime.now().strftime("%Y%m%d-%H%M%S")
    archive_id = f"{safe_name(lab_name)}_{stamp}_{uuid.uuid4().hex[:8]}"
    ARCHIVES_ROOT.mkdir(parents=True, exist_ok=True)
    final_path = archive_path_for_id(archive_id)
    temp_path = final_path.with_suffix(".zip.tmp")

    # Archives contain a transport snapshot assembled from the current field files.
    # The working directory itself only needs the hidden structural index.
    state_bytes = json.dumps(state, ensure_ascii=False, indent=2).encode("utf-8")
    sources = []
    asset_dir = STORE.project_root(lab_name)
    if asset_dir.exists():
        for item in sorted(asset_dir.rglob("*")):
            relative = item.relative_to(asset_dir)
            if item.is_file() and not item.is_symlink() and relative.parts[0] not in {"engagement.json", ".aeros", "documents"}:
                contained(asset_dir, relative)
                sources.append((item, f"assets/{relative.as_posix()}"))
    backup_dir = contained(STORE.project_root(lab_name), ".aeros", "revisions")
    if backup_dir.exists():
        for item in sorted(backup_dir.iterdir()):
            if item.is_file() and item.suffix in {".json", ".zip"}:
                sources.append((item, f"backups/{item.name}"))
    document_dir = engagement_documents_root(lab_name)
    if document_dir.exists():
        for item in sorted(document_dir.rglob("*")):
            if item.is_file():
                sources.append((item, f"documents/{item.relative_to(document_dir).as_posix()}"))

    inventory = [{"path": "engagement.json", "size": len(state_bytes), "sha256": hashlib.sha256(state_bytes).hexdigest()}]
    for source, archive_name in sources:
        inventory.append({
            "path": archive_name,
            "size": source.stat().st_size,
            "sha256": file_sha256(source),
        })
    config = state.get("engagementConfig") if isinstance(state, dict) else {}
    identity = config.get("identity", {}) if isinstance(config, dict) else {}
    manifest = {
        "format": "aeros-engagement-archive",
        "schemaVersion": SCHEMA_VERSION,
        "applicationVersion": APPLICATION_VERSION,
        "archiveId": archive_id,
        "engagementName": lab_name,
        "storageKey": engagement_storage_key(lab_name),
        "clientName": identity.get("clientName", ""),
        "engagementStatus": identity.get("status", ""),
        "archivedAt": datetime.datetime.now(datetime.timezone.utc).isoformat(),
        "removeWorkingCopy": bool(remove_working),
        "files": inventory,
    }

    try:
        with zipfile.ZipFile(temp_path, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as zf:
            zf.writestr("manifest.json", json.dumps(manifest, indent=2))
            zf.writestr("engagement.json", state_bytes)
            for source, archive_name in sources:
                zf.write(source, archive_name)
        saved_manifest = read_archive_manifest(temp_path, verify_files=True)
        if saved_manifest.get("archiveId") != archive_id:
            raise ValueError("Archive manifest verification failed")
        os.replace(temp_path, final_path)
    finally:
        if temp_path.exists():
            temp_path.unlink()

    if remove_working:
        delete_project(lab_name, expected_revision=snapshot_revision)
    else:
        config = state.setdefault("engagementConfig", {})
        config.pop("storage", None)
        config.pop("commandPaths", None)
        identity = config.setdefault("identity", {})
        identity["status"] = "archived"
        state["archivedAt"] = manifest["archivedAt"]
        db_result = STORE.save_project(
            state,
            expected_name=lab_name,
            expected_revision=snapshot_revision,
            allow_empty_hosts=True,
            reason="archive-status",
        )


    return {"archiveId": archive_id, "path": str(final_path), "manifest": manifest, "size": final_path.stat().st_size}


def _archive_member_digest(zf, info):
    digest = hashlib.sha256()
    actual_size = 0
    with zf.open(info, "r") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            actual_size += len(chunk)
            if actual_size > info.file_size:
                raise ValueError(f"Archive member changed size while reading: {info.filename}")
            digest.update(chunk)
    if actual_size != info.file_size:
        raise ValueError(f"Archive member size does not match metadata: {info.filename}")
    return digest.hexdigest()


def validate_engagement_archive(path, verify_files=False):
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError("Archive not found")
    if path.stat().st_size > MAX_ENGAGEMENT_ARCHIVE_BYTES:
        raise ValueError("Engagement archive exceeds the 1 GB safety limit")
    try:
        archive = zipfile.ZipFile(path, "r")
    except zipfile.BadZipFile as exc:
        raise ValueError("Invalid engagement archive") from exc
    with archive as zf:
        infos = zf.infolist()
        if len(infos) > MAX_ENGAGEMENT_ARCHIVE_MEMBERS:
            raise ValueError("Engagement archive contains too many members")
        files = {}
        seen_windows_names = set()
        expanded = 0
        for info in infos:
            normalized = safe_import_zip_member(info)
            windows_key = normalized.casefold()
            if windows_key in seen_windows_names:
                raise ValueError(f"Engagement archive contains a duplicate path: {normalized}")
            seen_windows_names.add(windows_key)
            parts = PurePosixPath(normalized).parts
            top_level = parts[0]
            exact_root_file = normalized in {"manifest.json", "engagement.json"}
            tree_member = top_level in {"assets", "documents", "backups"} and (
                len(parts) > 1 or info.is_dir()
            )
            if not exact_root_file and not tree_member:
                raise ValueError(f"Unsupported engagement archive member: {normalized}")
            if info.is_dir():
                continue
            if info.file_size > MAX_ENGAGEMENT_ARCHIVE_MEMBER_BYTES:
                raise ValueError(f"Engagement archive member exceeds the safety limit: {normalized}")
            expanded += info.file_size
            if expanded > MAX_ENGAGEMENT_ARCHIVE_EXPANDED_BYTES:
                raise ValueError("Expanded engagement archive exceeds the 2 GB safety limit")
            ratio = info.file_size / max(info.compress_size, 1)
            if info.file_size and ratio > MAX_ENGAGEMENT_ARCHIVE_COMPRESSION_RATIO:
                raise ValueError(f"Suspicious compression ratio in engagement archive: {normalized}")
            files[normalized] = info
        if "manifest.json" not in files or "engagement.json" not in files:
            raise ValueError("Engagement archive is missing manifest.json or engagement.json")
        manifest_info = files["manifest.json"]
        if manifest_info.file_size > MAX_ENGAGEMENT_MANIFEST_BYTES:
            raise ValueError("Engagement archive manifest is too large")
        try:
            manifest = json.loads(zf.read(manifest_info).decode("utf-8"))
        except (UnicodeDecodeError, json.JSONDecodeError) as exc:
            raise ValueError("Engagement archive manifest is invalid") from exc
        if not isinstance(manifest, dict):
            raise ValueError("Engagement archive manifest must be an object")
        inventory_rows = manifest.get("files")
        if not isinstance(inventory_rows, list):
            raise ValueError("Engagement archive manifest inventory is missing")
        inventory = {}
        for row in inventory_rows:
            if not isinstance(row, dict):
                raise ValueError("Engagement archive manifest contains an invalid inventory row")
            member_name = str(row.get("path") or "")
            if member_name in inventory or member_name == "manifest.json":
                raise ValueError(f"Engagement archive manifest contains a duplicate path: {member_name}")
            if member_name not in files:
                raise ValueError(f"Engagement archive manifest references a missing file: {member_name}")
            try:
                expected_size = int(row.get("size"))
            except (TypeError, ValueError) as exc:
                raise ValueError(f"Engagement archive manifest size is invalid: {member_name}") from exc
            expected_hash = str(row.get("sha256") or "").lower()
            if expected_size != files[member_name].file_size or not re.fullmatch(r"[a-f0-9]{64}", expected_hash):
                raise ValueError(f"Engagement archive manifest metadata is invalid: {member_name}")
            inventory[member_name] = (expected_size, expected_hash)
        actual_inventory = set(files) - {"manifest.json"}
        if set(inventory) != actual_inventory:
            raise ValueError("Engagement archive manifest does not exactly match archive contents")
        if verify_files:
            for member_name, (_, expected_hash) in inventory.items():
                if _archive_member_digest(zf, files[member_name]) != expected_hash:
                    raise ValueError(f"Engagement archive hash verification failed: {member_name}")
        return manifest


def read_archive_manifest(path, verify_files=False):
    manifest = validate_engagement_archive(path, verify_files=verify_files)
    if manifest.get("format") not in {"aeros-engagement-archive", "oscp-report-builder-engagement-archive"}:
        raise ValueError("Unsupported archive format")
    return manifest


def list_engagement_archives():
    if not ARCHIVES_ROOT.exists():
        return []
    rows = []
    for path in sorted(ARCHIVES_ROOT.glob("*.zip"), key=lambda p: p.stat().st_mtime, reverse=True):
        try:
            manifest = read_archive_manifest(path)
            rows.append({
                "archiveId": path.stem,
                "engagementName": manifest.get("engagementName", path.stem),
                "clientName": manifest.get("clientName", ""),
                "archivedAt": manifest.get("archivedAt", ""),
                "schemaVersion": manifest.get("schemaVersion", 1),
                "size": path.stat().st_size,
            })
        except Exception:
            rows.append({"archiveId": path.stem, "engagementName": path.stem, "clientName": "", "archivedAt": "", "schemaVersion": 0, "size": path.stat().st_size, "invalid": True})
    return rows


RESTORED_INTERNAL_PATH_FIELDS = {
    "screenshotAbsPath", "storedPath", "absolutePath", "artifactPath",
}


def _rewrite_restored_internal_paths(value, asset_root, document_root, asset_members, document_members, field_name=""):
    if isinstance(value, list):
        return [
            _rewrite_restored_internal_paths(
                item, asset_root, document_root, asset_members, document_members, field_name
            )
            for item in value
        ]
    if isinstance(value, dict):
        return {
            key: _rewrite_restored_internal_paths(
                item, asset_root, document_root, asset_members, document_members, key
            )
            for key, item in value.items()
        }
    if not isinstance(value, str) or field_name not in RESTORED_INTERNAL_PATH_FIELDS:
        return value
    normalized = value.replace("\\", "/").rstrip("/")
    for relative in sorted(asset_members, key=len, reverse=True):
        if normalized.casefold().endswith(f"/{relative}".casefold()):
            return str((asset_root / Path(*PurePosixPath(relative).parts)).resolve())
    for relative in sorted(document_members, key=len, reverse=True):
        if normalized.casefold().endswith(f"/{relative}".casefold()):
            return str((document_root / Path(*PurePosixPath(relative).parts)).resolve())
    return value


def _extract_archive_tree(zf, prefix, destination):
    extracted = False
    for info in zf.infolist():
        if info.is_dir() or not info.filename.startswith(f"{prefix}/"):
            continue
        relative = PurePosixPath(info.filename).relative_to(prefix)
        target = (destination / Path(*relative.parts)).resolve()
        if not path_is_inside(target, destination):
            raise ValueError("Unsafe path in engagement archive")
        target.parent.mkdir(parents=True, exist_ok=True)
        with zf.open(info, "r") as source, target.open("wb") as output:
            shutil.copyfileobj(source, output, length=1024 * 1024)
        extracted = True
    return extracted


def _remove_empty_parent_chain(path, stop):
    current = Path(path).parent
    stop = Path(stop).resolve()
    while current != stop and path_is_inside(current, stop):
        try:
            current.rmdir()
        except OSError:
            break
        current = current.parent


@serialized_project_mutation
def restore_engagement_archive(archive_id, new_name=""):
    path = archive_path_for_id(archive_id)
    if not path.exists():
        raise FileNotFoundError("Archive not found")
    manifest = read_archive_manifest(path, verify_files=True)
    requested_name = str(new_name or manifest.get("engagementName") or "").strip()
    if not requested_name:
        raise ValueError("The archive does not contain an engagement name")
    destination = project_state_path(requested_name)
    asset_root = STORE.project_root(requested_name).resolve()
    document_root = engagement_documents_root(requested_name).resolve()
    backup_root = contained(asset_root, ".aeros", "revisions")
    if (
        destination.exists()
        or STORE.project_exists(requested_name)
        or asset_root.exists()
        or document_root.parent.exists()
        or backup_root.exists()
    ):
        raise FileExistsError("An engagement with that name already exists")

    staging_root = (DATA_ROOT / ".restore-staging" / uuid.uuid4().hex).resolve()
    staged_assets = staging_root / "assets"
    staged_documents = staging_root / "documents"
    staged_backups = staging_root / "backups"
    moved_destinations = []
    try:
        with zipfile.ZipFile(path, "r") as zf:
            state_info = zf.getinfo("engagement.json")
            if state_info.file_size > MAX_BODY_BYTES:
                raise ValueError("Archived engagement state exceeds the 50 MB safety limit")
            try:
                state = json.loads(zf.read(state_info).decode("utf-8"))
            except (UnicodeDecodeError, json.JSONDecodeError) as exc:
                raise ValueError("Archived engagement state is invalid") from exc
            if not isinstance(state, dict):
                raise ValueError("Archived engagement state must be an object")
            has_assets = _extract_archive_tree(zf, "assets", staged_assets)
            has_documents = _extract_archive_tree(zf, "documents", staged_documents)
            has_backups = _extract_archive_tree(zf, "backups", staged_backups)
            asset_members = {
                info.filename.removeprefix("assets/")
                for info in zf.infolist()
                if not info.is_dir() and info.filename.startswith("assets/")
            }
            document_members = {
                info.filename.removeprefix("documents/")
                for info in zf.infolist()
                if not info.is_dir() and info.filename.startswith("documents/")
            }
        state["projectName"] = requested_name
        state.pop("_fileStorage", None)
        config = state.setdefault("engagementConfig", {})
        if not isinstance(config, dict):
            raise ValueError("Archived engagement configuration is invalid")
        config.pop("storage", None)
        config.pop("commandPaths", None)
        identity = config.setdefault("identity", {})
        if not isinstance(identity, dict):
            raise ValueError("Archived engagement identity is invalid")
        if identity.get("status") == "archived":
            identity["status"] = "active"
        state.pop("archivedAt", None)
        state = _rewrite_restored_internal_paths(
            state, asset_root, document_root, asset_members, document_members
        )
        for staged, final, present in (
            (staged_assets, asset_root, has_assets),
            (staged_documents, document_root, has_documents),
            (staged_backups, backup_root, has_backups),
        ):
            if not present:
                continue
            final.parent.mkdir(parents=True, exist_ok=True)
            shutil.copytree(staged, final, dirs_exist_ok=True)
            moved_destinations.append(final)
        db_result = STORE.save_project(
            state,
            expected_name=requested_name,
            expected_revision=0,
            allow_empty_hosts=True,
            reason="archive-restore",
        )
        state = db_result["state"]
        mirror_warning = ""
        return {
            "labName": requested_name,
            "path": str(destination),
            "storagePath": str(STORE.project_root(requested_name)),
            "mirrorWarning": mirror_warning,
        }
    except Exception:
        if not STORE.project_exists(requested_name):
            for destination_path in reversed(moved_destinations):
                if destination_path.is_dir():
                    shutil.rmtree(destination_path, ignore_errors=True)
                elif destination_path.exists():
                    destination_path.unlink()
            for destination_path, owner_root in (
                (asset_root, ASSETS_ROOT),
                (document_root, ENGAGEMENT_DOCUMENTS_ROOT),
                (backup_root, LAB_BACKUPS_ROOT),
            ):
                _remove_empty_parent_chain(destination_path, owner_root)
        raise
    finally:
        shutil.rmtree(staging_root, ignore_errors=True)


@serialized_project_mutation
def delete_engagement_archive(archive_id):
    path = archive_path_for_id(archive_id)
    if not path.exists():
        return False
    path.unlink()
    return True





















def service_label(key):
    return {
        "ftp":"FTP","ssh":"SSH","smtp":"SMTP","dns":"DNS","http":"HTTP","https":"HTTPS","smb":"SMB","rdp":"RDP","winrm":"WinRM","snmp":"SNMP","ldap":"LDAP","kerberos":"Kerberos","mssql":"MSSQL","mysql":"MySQL","postgresql":"PostgreSQL","vnc":"VNC"
    }.get(key, str(key).upper())

def selected_services(host):
    return [(k,v) for k,v in (host.get("services") or {}).items() if v and v.get("checked")]

def split_ports(value):
    return [p for p in re.split(r"[\s,;/]+", str(value or "")) if p]











RESERVED_LAB_CHILD_DIRS = {
    "templates", "template",
    "credentials",
    "loot",
    "lateral movement", "lateral-movement", "lateral_movement",
    "screenshots",
    "report", "reports",
    "report_templates", "report-templates",
    "notes",
    "exports",
    "wordlists", "wordlist",
    "evidence",
}

def normalize_reserved_name(value):
    return re.sub(r"\s+", " ", str(value or "").strip().lower().replace("_", " ").replace("-", " "))

def is_reserved_lab_child_dir(name):
    n = normalize_reserved_name(name)
    return n in {normalize_reserved_name(x) for x in RESERVED_LAB_CHILD_DIRS}


def clean_reserved_hosts_from_state(state):
    hosts = state.get("hosts") or {}
    for hid, host in list(hosts.items()):
        if is_reserved_lab_child_dir(host.get("ip") or hid) or is_reserved_lab_child_dir(host.get("hostname") or ""):
            del hosts[hid]
    if state.get("activeHost") not in hosts:
        state["activeHost"] = next(iter(hosts), None)
    return state






def parse_open_tcp_ports_from_text(text):
    ports = set()
    for m in re.finditer(r"(\d{1,5})/tcp\s+open\s+", str(text or ""), flags=re.I):
        try:
            port = int(m.group(1))
            if 0 < port <= 65535:
                ports.add(port)
        except Exception:
            pass
    return sorted(ports)


def selected_tcp_ports(host):
    ports = set(parse_open_tcp_ports_from_text((host.get("scans") or {}).get("port", "")))
    ports.update(parse_open_tcp_ports_from_text((host.get("scans") or {}).get("tcp", "")))
    for _key, val in selected_services(host):
        for p in split_ports(val.get("ports", "")):
            try:
                port = int(p)
                if 0 < port <= 65535:
                    ports.add(port)
            except Exception:
                pass
    return sorted(ports)


def scan_command(scan_type, host):
    ip = safe_name(host.get("ip") or "<rhost>")
    if scan_type == "port":
        return f"sudo nmap -n -Pn -sS -p- --min-rate 20000 --open {ip}"
    if scan_type == "udp":
        return f"sudo nmap -n -Pn -sU --top-ports 100 --open {ip}"
    ports = ",".join(str(p) for p in selected_tcp_ports(host)) or "<ports>"
    return f"sudo nmap -n -Pn -sS -sC -sV -p {ports} {ip}"







def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    value = str(text or "")
    lines = value.splitlines() or [""]
    for i, line in enumerate(lines):
        if i:
            p.add_run().add_break()
        r = p.add_run(line)
        r.bold = bold
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.name = BODY_FONT_NAME
            run.font.size = Pt(10)

def add_code_block(doc, text):
    if not text:
        p = doc.add_paragraph("TBD")
        normal = get_style(doc, "Normal")
        if normal is not None:
            p.style = normal
        return
    # Limit giant scan blocks in the Word report to keep it readable.
    text = str(text)
    if len(text) > 9000:
        text = text[:9000] + "\n\n[TRUNCATED IN REPORT - SEE RAW NOTES]"
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = "Consolas"
    r.font.size = Pt(8)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "E7E6E6")
    p._p.get_or_add_pPr().append(shading)

def add_labeled_para(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(label + ": ")
    r.bold = True
    r.font.name = BODY_FONT_NAME
    r.font.size = Pt(BODY_FONT_SIZE)
    value = str(text or "TBD")
    parts = value.splitlines()
    if not parts:
        vr = p.add_run("TBD")
        vr.font.name = BODY_FONT_NAME
        vr.font.size = Pt(BODY_FONT_SIZE)
    else:
        vr = p.add_run(parts[0])
        vr.font.name = BODY_FONT_NAME
        vr.font.size = Pt(BODY_FONT_SIZE)
        for line in parts[1:]:
            p.add_run().add_break()
            vr = p.add_run(line)
            vr.font.name = BODY_FONT_NAME
            vr.font.size = Pt(BODY_FONT_SIZE)
    return p

def add_labeled_para_with_gap(doc, label, text, gap_font_size=GAP_FIELD):
    """Add a labeled paragraph and force a visible blank line inside it.

    Separate spacer paragraphs can be collapsed by the OffSec Word template.
    This puts the blank line inside the paragraph so Word visibly preserves it.
    """
    p = add_labeled_para(doc, label, text)
    p.add_run().add_break()
    spacer = p.add_run("\u00A0")
    spacer.font.name = BODY_FONT_NAME
    spacer.font.size = Pt(gap_font_size)
    return p


def add_severity_para_with_gap(doc, severity, gap_font_size=GAP_FIELD):
    """Add OSCP-style Severity line and color-code the severity value."""
    severity = str(severity or "TBD")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    label_run = p.add_run("Severity: ")
    label_run.bold = True
    label_run.font.name = BODY_FONT_NAME
    label_run.font.size = Pt(BODY_FONT_SIZE)

    value_run = p.add_run(severity)
    value_run.bold = True
    value_run.font.name = BODY_FONT_NAME
    value_run.font.size = Pt(BODY_FONT_SIZE)
    color = SEVERITY_COLORS.get(severity.strip().lower())
    if color:
        value_run.font.color.rgb = color

    p.add_run().add_break()
    spacer = p.add_run("\u00A0")
    spacer.font.name = BODY_FONT_NAME
    spacer.font.size = Pt(gap_font_size)
    return p

def get_style(doc, name):
    try:
        return doc.styles[name]
    except Exception:
        return None


def ensure_paragraph_style(doc, name, size=None, bold=False):
    style = get_style(doc, name)
    if style is None:
        try:
            style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        except Exception:
            return None
    try:
        style.font.name = BODY_FONT_NAME
        if size:
            style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)
    except Exception:
        pass
    return style


def ensure_report_styles(doc):
    # Some .docx templates don't expose the standard heading styles through python-docx.
    # Creating them up front prevents: KeyError: "no style with name 'Heading 1'".
    ensure_paragraph_style(doc, "Normal", size=BODY_FONT_SIZE, bold=False)
    ensure_paragraph_style(doc, "Heading 1", size=16, bold=True)
    ensure_paragraph_style(doc, "Heading 2", size=13, bold=True)
    ensure_paragraph_style(doc, "Heading 3", size=BODY_FONT_SIZE, bold=True)
    ensure_paragraph_style(doc, "List Bullet", size=BODY_FONT_SIZE, bold=False)


def safe_table_grid(table):
    """Apply a visible Word table grid even when the source template hides borders."""
    try:
        table.style = "Table Grid"
    except Exception:
        pass
    try:
        tbl = table._tbl
        tblPr = tbl.tblPr
        if tblPr is None:
            tblPr = OxmlElement("w:tblPr")
            tbl.insert(0, tblPr)
        borders = tblPr.find(qn("w:tblBorders"))
        if borders is None:
            borders = OxmlElement("w:tblBorders")
            tblPr.append(borders)
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            tag = qn(f"w:{edge}")
            el = borders.find(tag)
            if el is None:
                el = OxmlElement(f"w:{edge}")
                borders.append(el)
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "6")
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), "808080")
    except Exception:
        pass


def replace_text_in_doc(doc, replacements):
    def replace_in_paragraph(paragraph):
        full = "".join(run.text for run in paragraph.runs)
        new = full
        for old, val in replacements.items():
            new = new.replace(old, str(val or ""))
        if new != full:
            for run in paragraph.runs:
                run.text = ""
            if paragraph.runs:
                paragraph.runs[0].text = new
            else:
                paragraph.add_run(new)
    for p in doc.paragraphs:
        replace_in_paragraph(p)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    replace_in_paragraph(p)


def keep_cover_page_only(doc):
    body = doc._body._element
    children = list(body)
    keep_until_idx = None
    for i, child in enumerate(children):
        if child.xpath('.//w:br[@w:type="page"]'):
            keep_until_idx = i
            break
    if keep_until_idx is None:
        paras = list(doc.paragraphs)
        if len(paras) >= 8:
            keep_el = paras[7]._p
            keep_until_idx = children.index(keep_el)
        else:
            keep_until_idx = len(children)-1
    sectPr = None
    for child in list(body):
        if child.tag == qn('w:sectPr'):
            sectPr = child
            break
    for i, child in enumerate(list(body)):
        if child.tag == qn('w:sectPr'):
            continue
        if i > keep_until_idx:
            body.remove(child)
    if sectPr is not None and sectPr.getparent() is None:
        body.append(sectPr)




def set_table_fixed_layout(table):
    """Prevent Word from squeezing the IP column when port lists are long."""
    try:
        table.autofit = False
    except Exception:
        pass
    try:
        tblPr = table._tbl.tblPr
        if tblPr is None:
            tblPr = OxmlElement("w:tblPr")
            table._tbl.insert(0, tblPr)
        layout = tblPr.find(qn("w:tblLayout"))
        if layout is None:
            layout = OxmlElement("w:tblLayout")
            tblPr.append(layout)
        layout.set(qn("w:type"), "fixed")
    except Exception:
        pass


def set_cell_width(cell, inches):
    """Set cell width in both python-docx and OOXML units."""
    try:
        cell.width = Inches(inches)
    except Exception:
        pass
    try:
        tcPr = cell._tc.get_or_add_tcPr()
        tcW = tcPr.find(qn("w:tcW"))
        if tcW is None:
            tcW = OxmlElement("w:tcW")
            tcPr.append(tcW)
        tcW.set(qn("w:w"), str(int(inches * 1440)))
        tcW.set(qn("w:type"), "dxa")
    except Exception:
        pass


def set_table_grid_widths(table, widths):
    """Force actual table grid column widths so Word/LibreOffice don't equalize columns."""
    try:
        tbl = table._tbl
        tblGrid = tbl.find(qn("w:tblGrid"))
        if tblGrid is None:
            tblGrid = OxmlElement("w:tblGrid")
            # Insert grid after tblPr when possible.
            tbl.insert(1 if tbl.tblPr is not None else 0, tblGrid)
        for child in list(tblGrid):
            tblGrid.remove(child)
        for width in widths:
            gridCol = OxmlElement("w:gridCol")
            gridCol.set(qn("w:w"), str(int(width * 1440)))
            tblGrid.append(gridCol)
    except Exception:
        pass


def apply_port_summary_widths(table):
    # Letter page with normal-ish margins: keep IP wide enough for IPv4/hostnames,
    # keep protocol small, give the port list the remaining width.
    widths = [1.65, 0.75, 4.60]
    set_table_grid_widths(table, widths)
    for row in table.rows:
        cells = row.cells
        if len(cells) >= 3:
            set_cell_width(cells[0], widths[0])
            set_cell_width(cells[1], widths[1])
            set_cell_width(cells[2], widths[2])


def wrap_port_list(ports, max_chars=75):
    """Wrap comma-separated port lists at comma boundaries, up to max_chars per line."""
    clean = [str(p).strip() for p in ports if str(p).strip()]
    if not clean:
        return "TBD"

    lines = []
    current = ""

    for port in clean:
        candidate = port if not current else f"{current},{port}"
        if current and len(candidate) > max_chars:
            lines.append(current)
            current = port
        else:
            current = candidate

    if current:
        lines.append(current)

    return "\n".join(lines)

def add_toc_line(doc, section, title, page="", bold=False):
    """Static TOC line without a Word table. Uses tabs like the OffSec-style sample."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    try:
        p.paragraph_format.tab_stops.add_tab_stop(Inches(0.80))
        p.paragraph_format.tab_stops.add_tab_stop(Inches(6.45), WD_TAB_ALIGNMENT.RIGHT)
    except Exception:
        pass
    r = p.add_run(f"{section}\t{title}\t{page}")
    r.bold = bold
    r.font.name = BODY_FONT_NAME
    r.font.size = Pt(10 if not bold else 11)
    return p

def first_name(full_name):
    full_name = str(full_name or "").strip()
    if not full_name:
        return "The student"
    return full_name.split()[0]


def evidence_has_type(host, proof_type):
    """Return True when a host has posted local/proof evidence, even if Proof Value was left blank.

    This lets screenshot-based evidence count correctly in the high-level summary.
    The report still prints the proof value when the user typed one, but a posted
    local.txt/proof.txt screenshot should count as submitted evidence.
    """
    proof_file = "local.txt" if proof_type == "local" else "proof.txt"
    for ev in host.get("evidence") or []:
        ev_type = str(ev.get("type") or "").strip().lower()
        ev_file = str(ev.get("proofFile") or "").strip().lower()
        title = str(ev.get("title") or "").strip().lower()
        if ev_type == proof_type:
            return True
        if ev_file == proof_file:
            return True
        if proof_file in title:
            return True
    return False


def has_local_or_proof(host):
    return bool(
        str(host.get("localProof", "")).strip()
        or str(host.get("proofTxt", "")).strip()
        or evidence_has_type(host, "local")
        or evidence_has_type(host, "proof")
    )


def has_system_proof(host):
    return bool(
        str(host.get("proofTxt", "")).strip()
        or evidence_has_type(host, "proof")
    )


def proof_counts(hosts):
    total = len(hosts)
    gained = sum(1 for h in hosts if has_local_or_proof(h))
    elevated = sum(1 for h in hosts if has_system_proof(h))
    return total, gained, elevated


def ip_list(hosts):
    ips = [str(h.get("ip", "")).strip() for h in hosts if str(h.get("ip", "")).strip()]
    return ", ".join(ips) if ips else "TBD"


def host_name_ip(host):
    hostname = str(host.get("hostname", "")).strip()
    ip = str(host.get("ip", "")).strip() or "TBD"
    return f"{hostname} - {ip}" if hostname else ip


def host_os_report_facts(host):
    resolution = host.get("osResolution") if isinstance(host.get("osResolution"), dict) else {}
    status = str(resolution.get("status") or "").strip().lower()
    family = str(resolution.get("family") or host.get("os") or "unknown").strip().lower()
    label = "Conflict" if status == "conflict" else ("Unknown" if family == "unknown" else family.title())
    confidence = str(resolution.get("confidence") or ("Unknown" if family == "unknown" else "Recorded")).strip()
    sources = [str(value).strip() for value in resolution.get("sourceLabels", []) if str(value).strip()] if isinstance(resolution.get("sourceLabels"), list) else []
    return {"label": label, "confidence": confidence, "sources": sources}


def target_title(idx, host):
    return f"Target #{idx} - {str(host.get('ip','TBD')).strip() or 'TBD'}"



def add_report_heading(doc, text, level=1):
    """Add an OSCP-style static report heading."""
    p = doc.add_paragraph()
    try:
        normal = get_style(doc, "Normal")
        if normal is not None:
            p.style = normal
    except Exception:
        pass
    r = p.add_run(str(text or ""))
    r.bold = True
    r.font.name = BODY_FONT_NAME
    if level == 1:
        r.font.size = Pt(16)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
    elif level == 2:
        r.font.size = Pt(13)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(3)
    else:
        r.font.size = Pt(BODY_FONT_SIZE)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(10)
    return p

def add_visible_blank_line(doc, after=8):
    """Insert a visible spacer paragraph that Word templates should not collapse."""
    if after is None or after <= 0:
        return None
    p = doc.add_paragraph()
    r = p.add_run("\u00A0")
    r.font.name = BODY_FONT_NAME
    r.font.size = Pt(10)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1
    return p

def canonical_toc_title(state, host, stage, legacy_title):
    asset_id = str(host.get("id") or "").strip()
    if not canonical_stage_managed(state, asset_id, stage):
        return legacy_title
    primary = next(
        (row for row in ordered_report_findings(state, asset_id, stage) if row["placement"]["role"] == "primary"),
        None,
    )
    return str(primary["finding"].get("title") or "Finding").strip() if primary else ""


def add_static_toc_page(doc, independent, ad_hosts, state=None):
    state = state if isinstance(state, dict) else {}
    ensure_report_styles(doc)
    add_report_heading(doc, "Table of Contents", level=1)
    toc_items = [
        ("1", "OffSec Certified Professional Exam Report"),
        ("1.1", "Introduction"),
        ("1.2", "Objective"),
        ("1.3", "Requirements"),
        ("2", "High-Level Summary"),
        ("2.1", "Recommendations"),
        ("3", "Methodologies"),
        ("3.1", "Information Gathering"),
        ("3.2", "Service Enumeration"),
        ("3.3", "Penetration"),
        ("3.4", "Maintaining Access"),
        ("3.5", "House Cleaning"),
        ("4", "Independent Challenges"),
    ]
    for i, h in enumerate(independent, 1):
        base = f"4.{i}"
        initial_title = canonical_toc_title(state, h, "initial-access", h.get("findingTitle") or "Finding TBD")
        privesc_title = canonical_toc_title(state, h, "privilege-escalation", legacy_privesc_title(h))
        toc_items.extend([
            (base, target_title(i, h)),
            (f"{base}.1", f"Initial Access{' - ' + initial_title if initial_title else ''}"),
            (f"{base}.2", "Service Enumeration"),
            (f"{base}.3", "Initial Access Walkthrough"),
            (f"{base}.4", f"Privilege Escalation{' - ' + privesc_title if privesc_title else ''}"),
            (f"{base}.5", "Post Exploitation"),
        ])
    toc_items.append(("5", "Active Directory Set"))
    for i, h in enumerate(ad_hosts, 1):
        base = f"5.{i}"
        initial_title = canonical_toc_title(state, h, "initial-access", h.get("findingTitle") or "TBD")
        privesc_title = canonical_toc_title(state, h, "privilege-escalation", legacy_privesc_title(h))
        toc_items.extend([
            (base, host_name_ip(h)),
            (f"{base}.1", f"Initial Access{' - ' + initial_title if initial_title else ''}"),
            (f"{base}.2", f"Privilege Escalation{' - ' + privesc_title if privesc_title else ''}"),
            (f"{base}.3", "Post-Exploitation"),
        ])

    add_toc_line(doc, "Section", "Title", "Page", bold=True)
    for number, title in toc_items:
        add_toc_line(doc, number, title, "")
    note = doc.add_paragraph("Update page numbers manually or insert an automatic Word table of contents after final edits.")
    note.paragraph_format.space_before = Pt(8)
    doc.add_page_break()

def scan_port_rows(host, proto="tcp"):
    """Return de-duplicated port rows for Word tables.

    Supports TCP open rows and UDP open/open|filtered rows. The report table
    keeps one row per port and combines service notes from port scan, full Nmap,
    UDP output, and checked services.
    """
    proto = proto.lower()
    scans = host.get("scans") or {}
    scan_text = "\n".join([scans.get("port", ""), scans.get("tcp", "")]) if proto == "tcp" else scans.get("udp", "")
    by_port = {}

    # Examples:
    # 80/tcp open http
    # 123/udp open|filtered ntp
    pattern = re.compile(r"(\d{1,5})/" + re.escape(proto) + r"\s+(open(?:\|filtered)?)\s+([^\s]+)", flags=re.I)
    for m in pattern.finditer(scan_text):
        port = str(m.group(1))
        state = str(m.group(2)).lower()
        svc = str(m.group(3)).strip()
        entry = by_port.setdefault(port, {"state": state, "services": []})
        if entry.get("state") != "open" and state == "open":
            entry["state"] = "open"
        if svc and svc.lower() not in [x.lower() for x in entry["services"]]:
            entry["services"].append(svc)

    if proto == "tcp":
        for k, v in selected_services(host):
            svc = service_label(k)
            ports = split_ports(v.get("ports", ""))
            if not ports:
                key = f"svc-{svc}"
                entry = by_port.setdefault(key, {"state": "open", "services": []})
                if svc.lower() not in [x.lower() for x in entry["services"]]:
                    entry["services"].append(svc)
                continue
            for port in ports:
                entry = by_port.setdefault(str(port), {"state": "open", "services": []})
                if svc.lower() not in [x.lower() for x in entry["services"]]:
                    entry["services"].append(svc)

    def sort_key(item):
        port = item[0]
        try:
            return (0, int(port))
        except Exception:
            return (1, port)

    rows = []
    for port, entry in sorted(by_port.items(), key=sort_key):
        display_port = port.replace("svc-", "") if port.startswith("svc-") else port
        rows.append((proto.upper(), display_port, entry.get("state") or "open", "; ".join(entry.get("services") or ["TBD"])))
    return rows


def port_numbers_for_proto(host, proto="tcp"):
    """Return de-duplicated numeric port strings for the OSCP-style port summary table."""
    ports = []
    for row in scan_port_rows(host, proto):
        port = str(row[1] or "").strip()
        if not port or not port.isdigit():
            continue
        if port not in ports:
            ports.append(port)
    return ports


def add_port_summary_table(doc, hosts):
    """OSCP-style summary table: IP on the left, TCP ports to the right, UDP below."""
    if not isinstance(hosts, (list, tuple)):
        hosts = [hosts]
    p = doc.add_paragraph("Port Scan Results")
    if p.runs:
        p.runs[0].bold = True
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(5)

    table = doc.add_table(rows=1, cols=3)
    safe_table_grid(table)
    set_table_fixed_layout(table)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    set_cell_text(hdr[0], "IP Address", True)
    set_cell_text(hdr[1], "Protocol", True)
    set_cell_text(hdr[2], "Ports Open", True)
    for cell in hdr:
        set_cell_shading(cell, "D9EAF7")
    apply_port_summary_widths(table)

    for host in hosts:
        ip = str(host.get("ip", "TBD") or "TBD")
        tcp_ports = port_numbers_for_proto(host, "tcp")
        udp_ports = port_numbers_for_proto(host, "udp")

        row = table.add_row().cells
        set_cell_text(row[0], ip)
        set_cell_text(row[1], "TCP")
        set_cell_text(row[2], wrap_port_list(tcp_ports))
        apply_port_summary_widths(table)

        if udp_ports:
            row = table.add_row().cells
            set_cell_text(row[0], "")
            set_cell_text(row[1], "UDP")
            set_cell_text(row[2], wrap_port_list(udp_ports))
            apply_port_summary_widths(table)

    add_visible_blank_line(doc, 12)
    return True


def add_report_spacer(doc, after=8):
    return add_visible_blank_line(doc, after)


CANONICAL_REPORT_STAGES = ("initial-access", "privilege-escalation", "lateral-movement", "other")
CANONICAL_REPORT_ROLES = ("primary", "additional", "supporting")


def report_token(value):
    return "-".join(str(value or "").strip().lower().replace("_", "-").split())


def report_elevated(value):
    return report_token(value) in {"root-admin", "root", "administrator", "admin", "system"}


def active_report_contexts(host):
    return [row for row in (host.get("accessContexts") or []) if isinstance(row, dict) and row.get("active") is not False]


def confirmed_report_attempts(host):
    rows = []
    for attempt in host.get("exploitAttempts") or []:
        if not isinstance(attempt, dict):
            continue
        success = attempt.get("success") if isinstance(attempt.get("success"), dict) else {}
        if success.get("confirmed") is True and not success.get("revokedAt") and not attempt.get("revokedAt"):
            rows.append(attempt)
    return rows


def host_requires_report_stage(host, stage):
    host = host if isinstance(host, dict) else {}
    stage = report_token(stage)
    contexts = active_report_contexts(host)
    attempts = confirmed_report_attempts(host)
    explicit_attempt = any(report_token((row.get("success") or {}).get("report", {}).get("stage")) == stage for row in attempts)
    explicit_step = any(
        isinstance(row, dict) and row.get("active") is not False and not row.get("deletedAt")
        and report_token(row.get("stage") or row.get("reportStage")) == stage
        for row in host.get("attackPathSteps") or []
    )
    expected_evidence = {"initial-access": "initial-access", "privilege-escalation": "privilege-escalation", "lateral-movement": "lateral-movement"}.get(stage)
    explicit_evidence = bool(expected_evidence) and any(
        isinstance(row, dict) and row.get("active") is not False and not row.get("deletedAt")
        and report_token(row.get("type")) == expected_evidence
        for row in host.get("evidence") or []
    )
    if explicit_attempt or explicit_step or explicit_evidence:
        return True
    if stage == "initial-access":
        access_result = any(
            report_token((row.get("success") or {}).get("resultType")) in {"interactive-shell", "command-execution", "authenticated-session"}
            and (row.get("success") or {}).get("accessConfirmed") is True
            for row in attempts
        )
        return bool(contexts or access_result or host.get("hasShell") is True or str(host.get("localProof") or "").strip())
    if stage == "privilege-escalation":
        if str(host.get("privescTitle") or "").strip():
            return True
        elevated = any(report_elevated(row.get("privilege")) for row in contexts) or any(
            (row.get("success") or {}).get("accessConfirmed") is True
            and report_elevated((row.get("success") or {}).get("privilege"))
            for row in attempts
        )
        lower = any(
            bool(report_token(row.get("privilege")))
            and report_token(row.get("privilege")) != "unknown"
            and not report_elevated(row.get("privilege"))
            for row in contexts
        ) or any(
            (row.get("success") or {}).get("accessConfirmed") is True
            and bool(report_token((row.get("success") or {}).get("privilege")))
            and report_token((row.get("success") or {}).get("privilege")) != "unknown"
            and not report_elevated((row.get("success") or {}).get("privilege"))
            for row in attempts
        )
        explicit_direct_elevated = elevated and not lower and (
            any(report_elevated(row.get("privilege")) for row in contexts)
            or any(
                report_token((row.get("success") or {}).get("report", {}).get("stage")) == "initial-access"
                and (row.get("success") or {}).get("accessConfirmed") is True
                and report_elevated((row.get("success") or {}).get("privilege"))
                for row in attempts
            )
        )
        if str(host.get("proofTxt") or "").strip() and not explicit_direct_elevated:
            return True
        return bool(elevated and lower)
    return False


def direct_elevated_initial_access(host):
    host = host if isinstance(host, dict) else {}
    if host_requires_report_stage(host, "privilege-escalation"):
        return False
    contexts = active_report_contexts(host)
    attempts = confirmed_report_attempts(host)
    elevated_context = any(report_elevated(row.get("privilege")) for row in contexts)
    elevated_initial_attempt = any(
        report_token((row.get("success") or {}).get("report", {}).get("stage")) == "initial-access"
        and (row.get("success") or {}).get("accessConfirmed") is True
        and report_elevated((row.get("success") or {}).get("privilege"))
        for row in attempts
    )
    lower_context = any(
        bool(report_token(row.get("privilege")))
        and report_token(row.get("privilege")) != "unknown"
        and not report_elevated(row.get("privilege"))
        for row in contexts
    )
    lower_attempt = any(
        (row.get("success") or {}).get("accessConfirmed") is True
        and bool(report_token((row.get("success") or {}).get("privilege")))
        and report_token((row.get("success") or {}).get("privilege")) != "unknown"
        and not report_elevated((row.get("success") or {}).get("privilege"))
        for row in attempts
    )
    return bool((elevated_context or elevated_initial_attempt) and not (lower_context or lower_attempt))


def legacy_privesc_title(host):
    title = str((host or {}).get("privescTitle") or "").strip()
    if title:
        return title
    return "Not Applicable" if direct_elevated_initial_access(host) else "TBD"


def required_primary_placements(state):
    state = state if isinstance(state, dict) else {}
    rows = report_placement_rows(state)
    managed_assets = {
        str(row.get("placement", {}).get("assetId") or row.get("placement", {}).get("hostId") or "").strip()
        for row in rows
    }
    required = []
    seen = set()
    hosts = state.get("hosts") if isinstance(state.get("hosts"), dict) else {}
    for key, host in hosts.items():
        if not isinstance(host, dict):
            continue
        asset_id = str(host.get("id") or key or "").strip()
        if not asset_id or asset_id not in managed_assets:
            continue
        for stage in CANONICAL_REPORT_STAGES:
            if stage == "other" or not host_requires_report_stage(host, stage):
                continue
            marker = (asset_id, stage)
            if marker not in seen:
                seen.add(marker)
                required.append({"assetId": asset_id, "stage": stage})
    return required


def finding_deleted(finding):
    return bool(str(finding.get("deletedAt") or "").strip()) or str(finding.get("status") or "").strip().lower() == "deleted"


def finding_included(finding):
    return not finding_deleted(finding) and str(finding.get("status") or "").strip().lower() != "not-reported"


def report_asset_ids(state):
    ids = set()
    hosts = state.get("hosts") if isinstance(state.get("hosts"), dict) else {}
    for key, host in hosts.items():
        if str(key).strip():
            ids.add(str(key).strip())
        if isinstance(host, dict) and str(host.get("id") or "").strip():
            ids.add(str(host.get("id")).strip())
    return ids


def report_placement_rows(state):
    rows = []
    findings = state.get("findings") if isinstance(state.get("findings"), list) else []
    for finding_index, finding in enumerate(findings):
        if not isinstance(finding, dict) or finding_deleted(finding):
            continue
        placements = finding.get("reportPlacements") if isinstance(finding.get("reportPlacements"), list) else []
        for index, placement in enumerate(placements):
            if not isinstance(placement, dict):
                placement = {}
            finding_id = str(finding.get("id") or "").strip()
            rows.append({
                "finding": finding,
                "findingIndex": finding_index,
                "placement": placement,
                "placementId": str(placement.get("id") or f"placement-{finding_id or 'finding'}-{index + 1}").strip(),
                "included": finding_included(finding),
            })
    return rows


def report_placement_issue(code, message, row=None, finding_ids=None, placement_ids=None):
    row = row or {}
    placement = row.get("placement") if isinstance(row.get("placement"), dict) else {}
    finding = row.get("finding") if isinstance(row.get("finding"), dict) else {}
    return {
        "code": code,
        "message": message,
        "assetId": str(placement.get("assetId") or placement.get("hostId") or "").strip(),
        "stage": str(placement.get("stage") or "").strip().lower(),
        "findingIds": sorted(set(str(item).strip() for item in (finding_ids or [finding.get("id")]) if str(item or "").strip())),
        "placementIds": sorted(set(str(item).strip() for item in (placement_ids or [row.get("placementId")]) if str(item or "").strip())),
    }


def validate_report_placements(state, options=None):
    state = state if isinstance(state, dict) else {}
    options = options if isinstance(options, dict) else {}
    assets = report_asset_ids(state)
    rows = report_placement_rows(state)
    issues = []
    valid_rows = []
    seen_finding_stage = {}
    required_primary = set()
    required_items = options.get("requiredPrimary") if isinstance(options.get("requiredPrimary"), list) else required_primary_placements(state)
    for item in required_items:
        if not isinstance(item, dict):
            continue
        asset_id = str(item.get("assetId") or "").strip()
        stage = str(item.get("stage") or "").strip().lower()
        if asset_id and stage:
            required_primary.add((asset_id, stage))
    rows_to_validate = rows if options.get("includeNotReported") is True else [row for row in rows if row.get("included")]
    for row in rows_to_validate:
        placement = row["placement"]
        asset_id = str(placement.get("assetId") or placement.get("hostId") or "").strip()
        stage = str(placement.get("stage") or "").strip().lower()
        role = str(placement.get("role") or "").strip().lower()
        order_value = placement.get("order")
        valid_order = isinstance(order_value, int) and not isinstance(order_value, bool) and order_value >= 0
        structurally_valid = True
        if stage not in CANONICAL_REPORT_STAGES:
            issues.append(report_placement_issue("invalid-stage", "Report placement has an invalid stage.", row))
            structurally_valid = False
        if role not in CANONICAL_REPORT_ROLES:
            issues.append(report_placement_issue("invalid-role", "Report placement has an invalid role.", row))
            structurally_valid = False
        if not asset_id or asset_id not in assets:
            issues.append(report_placement_issue("missing-asset", "Report placement references a missing asset.", row))
            structurally_valid = False
        if not valid_order:
            issues.append(report_placement_issue("invalid-order", "Report placement order must be a non-negative integer.", row))
            structurally_valid = False
        duplicate_key = (str(row["finding"].get("id") or "").strip(), asset_id, stage)
        if duplicate_key in seen_finding_stage:
            prior = seen_finding_stage[duplicate_key]
            issues.append(report_placement_issue(
                "duplicate-placement",
                "A Finding has more than one placement for the same asset and stage.",
                row,
                finding_ids=[row["finding"].get("id")],
                placement_ids=[prior["placementId"], row["placementId"]],
            ))
            structurally_valid = False
        else:
            seen_finding_stage[duplicate_key] = row
        if structurally_valid:
            valid_rows.append({**row, "assetId": asset_id, "stage": stage, "role": role, "order": order_value})

    groups = {}
    for row in valid_rows:
        groups.setdefault((row["assetId"], row["stage"]), []).append(row)
    for group in groups.values():
        included = [row for row in group if row["included"]]
        primaries = [row for row in included if row["role"] == "primary"]
        if len(primaries) > 1:
            issues.append(report_placement_issue(
                "multiple-primary",
                "More than one active primary Finding exists for the same asset and stage.",
                primaries[0],
                finding_ids=[row["finding"].get("id") for row in primaries],
                placement_ids=[row["placementId"] for row in primaries],
            ))
        included_primary = [row for row in included if row["role"] == "primary"]
        if any(row["role"] == "additional" for row in included) and not included_primary:
            issues.append(report_placement_issue(
                "missing-primary",
                "Additional Findings require one active primary Finding for the asset and stage.",
                included[0],
                finding_ids=[row["finding"].get("id") for row in included],
                placement_ids=[row["placementId"] for row in included],
            ))
        group_key = (group[0]["assetId"], group[0]["stage"])
        if (
            group_key in required_primary
            and not included_primary
            and not any(
                item["code"] == "missing-primary"
                and (item["assetId"], item["stage"]) == group_key
                for item in issues
            )
        ):
            issues.append(report_placement_issue(
                "missing-primary",
                "Report readiness requires one included primary Finding for the asset and stage.",
                group[0],
                finding_ids=[row["finding"].get("id") for row in group],
                placement_ids=[row["placementId"] for row in group],
            ))
        order_groups = {}
        for row in included:
            if row["role"] != "primary":
                order_groups.setdefault((row["role"], row["order"]), []).append(row)
        for rows_at_order in order_groups.values():
            if len(rows_at_order) < 2:
                continue
            issues.append(report_placement_issue(
                "unstable-order",
                "Findings in the same role cannot share a display order.",
                rows_at_order[0],
                finding_ids=[row["finding"].get("id") for row in rows_at_order],
                placement_ids=[row["placementId"] for row in rows_at_order],
            ))
    for asset_id, stage in required_primary:
        if (asset_id, stage) in groups:
            continue
        if asset_id not in assets or stage not in CANONICAL_REPORT_STAGES:
            continue
        issues.append(report_placement_issue(
            "missing-primary",
            "Report readiness requires one included primary Finding for the asset and stage.",
            {"finding": {}, "placement": {"assetId": asset_id, "stage": stage}, "placementId": ""},
            finding_ids=[],
            placement_ids=[],
        ))
    issues.sort(key=lambda item: (
        item["assetId"], item["stage"], item["code"],
        ",".join(item["findingIds"]), ",".join(item["placementIds"])
    ))
    return {
        "valid": not issues,
        "blocking": bool(issues),
        "legacyOnly": not rows,
        "canonical": bool(rows),
        "issues": issues,
        "placementCount": len(rows),
    }


def report_host_for_asset(state, asset_id):
    wanted = str(asset_id or "").strip()
    hosts = state.get("hosts") if isinstance(state.get("hosts"), dict) else {}
    for key, host in hosts.items():
        if not isinstance(host, dict):
            continue
        if str(key).strip() == wanted or str(host.get("id") or "").strip() == wanted:
            return host
    return None


def report_attempt_for_ref(state, ref):
    ref = ref if isinstance(ref, dict) else {}
    host = report_host_for_asset(state, ref.get("assetId") or ref.get("hostId"))
    attempt_id = str(ref.get("attemptId") or ref.get("investigationId") or "").strip()
    attempt = next((
        item for item in (host.get("exploitAttempts") or [])
        if isinstance(item, dict) and str(item.get("id") or "").strip() == attempt_id
    ), None) if host else None
    return host, attempt


def canonical_reproduction_rows(state, finding):
    refs = finding.get("reproductionRefs") if isinstance(finding.get("reproductionRefs"), list) else []
    rows = []
    for index, ref in enumerate(refs):
        if not isinstance(ref, dict) or ref.get("included") is False:
            continue
        host, attempt = report_attempt_for_ref(state, ref)
        run_id = str(ref.get("runId") or ref.get("activityId") or "").strip()
        activity = next((
            item for item in (attempt.get("runs") or [])
            if isinstance(item, dict) and str(item.get("id") or "").strip() == run_id
        ), None) if attempt else None
        if activity:
            body = str(activity.get("body") or activity.get("requestOrCommand") or "").strip()
            if not body:
                pieces = [
                    activity.get("requestOrCommand"),
                    f"Changed values / payload:\n{activity.get('changedValues')}" if activity.get("changedValues") else "",
                    f"Output / response:\n{activity.get('output')}" if activity.get("output") else "",
                    activity.get("resultSummary"),
                ]
                body = "\n\n".join(str(item).strip() for item in pieces if str(item or "").strip())
            rows.append({
                "ref": ref, "host": host, "attempt": attempt, "activity": activity,
                "body": body, "order": ref.get("order") if isinstance(ref.get("order"), int) else index,
                "reportAddendum": str(ref.get("reportAddendum") or "").strip(),
            })
    rows.sort(key=lambda row: (row["order"], str(row["ref"].get("runId") or "")))
    if not rows and not refs:
        legacy = str(finding.get("reproduction") or finding.get("steps") or "").strip()
        if legacy:
            rows.append({"ref": None, "host": None, "attempt": None, "activity": None, "body": legacy, "order": 0, "reportAddendum": "", "legacy": True})
    return rows


def canonical_evidence_refs(finding):
    finding = finding if isinstance(finding, dict) else {}
    refs = [dict(ref) for ref in finding.get("evidenceRefs") or [] if isinstance(ref, dict) and str(ref.get("evidenceId") or "").strip()]
    legacy_ids = list(dict.fromkeys(str(item).strip() for item in (finding.get("evidenceIds") or []) if str(item or "").strip()))
    has_evidence_selection_version = "evidenceSelectionVersion" in finding
    supplied_evidence_refs = isinstance(finding.get("evidenceRefs"), list)
    try:
        evidence_selection_version = max(0, int(finding.get("evidenceSelectionVersion") or 0))
    except (TypeError, ValueError):
        evidence_selection_version = 0
    scoped_evidence_selection = (
        evidence_selection_version >= 1
        or (not has_evidence_selection_version and supplied_evidence_refs)
    )
    candidates = [
        str(ref.get("assetId") or ref.get("hostId") or "").strip()
        for ref in refs if str(ref.get("assetId") or ref.get("hostId") or "").strip()
    ]
    for value in finding.get("assetIds") or finding.get("affectedAssetIds") or []:
        if str(value or "").strip():
            candidates.append(str(value).strip())
    for placement in finding.get("reportPlacements") or []:
        if isinstance(placement, dict) and str(placement.get("assetId") or placement.get("hostId") or "").strip():
            candidates.append(str(placement.get("assetId") or placement.get("hostId")).strip())
    candidates = list(dict.fromkeys(candidates))
    if not scoped_evidence_selection and len(candidates) == 1:
        represented = {str(ref.get("evidenceId") or "").strip() for ref in refs}
        refs.extend({
            "assetId": candidates[0],
            "attemptId": "",
            "evidenceId": evidence_id,
        } for evidence_id in legacy_ids if evidence_id not in represented)
    seen = set()
    resolved = []
    for ref in refs:
        key = (
            str(ref.get("assetId") or ref.get("hostId") or "").strip(),
            str(ref.get("attemptId") or ref.get("investigationId") or "").strip(),
            str(ref.get("evidenceId") or "").strip(),
        )
        if not key[2] or key in seen:
            continue
        seen.add(key)
        resolved.append(ref)
    return resolved


def canonical_evidence_rows(state, finding):
    rows = []
    for ref in canonical_evidence_refs(finding):
        host = report_host_for_asset(state, ref.get("assetId") or ref.get("hostId"))
        evidence_id = str(ref.get("evidenceId") or "").strip()
        evidence = next((
            item for item in (host.get("evidence") or [])
            if isinstance(item, dict) and str(item.get("id") or "").strip() == evidence_id
        ), None) if host else None
        if evidence and evidence.get("active") is not False and not evidence.get("deletedAt"):
            rows.append({"ref": ref, "host": host, "evidence": evidence})
    return rows


def finding_reference_issue(code, message, finding, ref=None, kind="finding"):
    ref = ref if isinstance(ref, dict) else {}
    return {
        "code": code, "message": message, "kind": kind,
        "assetId": str(ref.get("assetId") or ref.get("hostId") or "").strip(),
        "attemptId": str(ref.get("attemptId") or ref.get("investigationId") or "").strip(),
        "findingIds": [str(finding.get("id") or "").strip()] if str(finding.get("id") or "").strip() else [],
        "reference": dict(ref),
    }


def validate_finding_references(state, options=None):
    state = state if isinstance(state, dict) else {}
    options = options if isinstance(options, dict) else {}
    issues = []
    require_report_content = options.get("requireReportContent") is not False
    findings = state.get("findings") if isinstance(state.get("findings"), list) else []
    for finding in findings:
        if not isinstance(finding, dict) or finding_deleted(finding):
            continue
        if not finding_included(finding) and options.get("includeNotReported") is not True:
            continue
        placement = next((item for item in (finding.get("reportPlacements") or []) if isinstance(item, dict)), {})
        fallback_ref = {"assetId": placement.get("assetId") or placement.get("hostId")}
        if require_report_content and not str(finding.get("title") or "").strip():
            issues.append(finding_reference_issue("missing-finding-title", "Canonical Finding title is required.", finding, fallback_ref))
        legacy_snapshot = finding.get("legacySnapshot") if isinstance(finding.get("legacySnapshot"), dict) else {}
        if require_report_content and any(str(value or "").strip() for value in legacy_snapshot.values()) and finding.get("legacyReviewConfirmed") is not True:
            issues.append(finding_reference_issue("legacy-review-required", "Legacy report values must be reviewed before canonical ownership is reportable.", finding, fallback_ref))
        if require_report_content and not any(row.get("body") for row in canonical_reproduction_rows(state, finding)):
            issues.append(finding_reference_issue("missing-reproduction", "Canonical Finding must include at least one valid reproduction activity.", finding, fallback_ref, "activity"))

        for ref in finding.get("sourceInvestigationRefs") if isinstance(finding.get("sourceInvestigationRefs"), list) else []:
            if not isinstance(ref, dict):
                continue
            host, attempt = report_attempt_for_ref(state, ref)
            if not host:
                issues.append(finding_reference_issue("missing-source-asset", "Finding source investigation references a missing asset.", finding, ref, "investigation"))
            elif not attempt:
                issues.append(finding_reference_issue("missing-source-investigation", "Finding source investigation no longer exists on the referenced asset.", finding, ref, "investigation"))
        for ref in finding.get("sourceResultRefs") if isinstance(finding.get("sourceResultRefs"), list) else []:
            if not isinstance(ref, dict):
                continue
            host, attempt = report_attempt_for_ref(state, ref)
            success = attempt.get("success") if attempt and isinstance(attempt.get("success"), dict) else None
            actual_id = str((success or {}).get("id") or (f"result-{attempt.get('id')}" if attempt else "")).strip()
            if not host:
                issues.append(finding_reference_issue("missing-source-asset", "Finding source result references a missing asset.", finding, ref, "result"))
            elif not attempt:
                issues.append(finding_reference_issue("missing-source-investigation", "Finding source result references a missing investigation.", finding, ref, "result"))
            elif not success or actual_id != str(ref.get("resultId") or "").strip():
                issues.append(finding_reference_issue("missing-source-result", "Finding source result no longer matches the referenced investigation.", finding, ref, "result"))
            elif success.get("confirmed") is not True or success.get("revokedAt") or attempt.get("revokedAt"):
                issues.append(finding_reference_issue("revoked-source-result", "Finding source result is not currently confirmed.", finding, ref, "result"))
        activity_refs = []
        for name in ("sourceActivityRefs", "reproductionRefs"):
            activity_refs.extend((name, ref) for ref in (finding.get(name) if isinstance(finding.get(name), list) else []))
        for ref_name, ref in activity_refs:
            if not isinstance(ref, dict):
                continue
            host, attempt = report_attempt_for_ref(state, ref)
            run_id = str(ref.get("runId") or ref.get("activityId") or "").strip()
            activity = next((item for item in (attempt.get("runs") or []) if isinstance(item, dict) and str(item.get("id") or "").strip() == run_id), None) if attempt else None
            if not host:
                issues.append(finding_reference_issue("missing-source-asset", "Finding activity references a missing asset.", finding, ref, "activity"))
            elif not attempt:
                issues.append(finding_reference_issue("missing-source-investigation", "Finding activity references a missing investigation.", finding, ref, "activity"))
            elif not activity:
                issues.append(finding_reference_issue("missing-source-activity", "Finding activity no longer exists in the referenced investigation.", finding, ref, "activity"))
            elif ref_name == "sourceActivityRefs" and str((attempt.get("success") or {}).get("sourceActivityId") or "").strip() != run_id:
                issues.append(finding_reference_issue("source-activity-result-mismatch", "Finding source activity no longer matches the confirmed result.", finding, ref, "activity"))
        orders = set()
        for ref in finding.get("reproductionRefs") if isinstance(finding.get("reproductionRefs"), list) else []:
            if not isinstance(ref, dict) or ref.get("included") is False:
                continue
            order = ref.get("order")
            if not isinstance(order, int) or isinstance(order, bool) or order < 0 or order in orders:
                issues.append(finding_reference_issue("duplicate-reproduction-order", "Included reproduction activities must have unique non-negative display orders.", finding, ref, "activity"))
            orders.add(order)
        for ref in canonical_evidence_refs(finding):
            if not isinstance(ref, dict):
                continue
            host = report_host_for_asset(state, ref.get("assetId") or ref.get("hostId"))
            evidence_id = str(ref.get("evidenceId") or "").strip()
            evidence = next((item for item in (host.get("evidence") or []) if isinstance(item, dict) and str(item.get("id") or "").strip() == evidence_id), None) if host else None
            if not host:
                issues.append(finding_reference_issue("missing-evidence-asset", "Finding evidence references a missing asset.", finding, ref, "evidence"))
            elif not evidence or evidence.get("active") is False or evidence.get("deletedAt"):
                issues.append(finding_reference_issue("missing-evidence", "Finding evidence is missing or inactive.", finding, ref, "evidence"))
            elif ref.get("attemptId") and str(evidence.get("sourceExploitAttemptId") or evidence.get("relatedExploitAttemptId") or evidence.get("attemptId") or "").strip() != str(ref.get("attemptId")).strip():
                attempt_ids = [str(item).strip() for item in evidence.get("attemptIds") or []]
                if str(ref.get("attemptId")).strip() not in attempt_ids:
                    issues.append(finding_reference_issue("evidence-investigation-mismatch", "Finding evidence is not associated with the referenced investigation.", finding, ref, "evidence"))
        for ref in finding.get("attackPathRefs") if isinstance(finding.get("attackPathRefs"), list) else []:
            if not isinstance(ref, dict):
                continue
            host = report_host_for_asset(state, ref.get("assetId") or ref.get("hostId"))
            step_id = str(ref.get("stepId") or ref.get("attackPathStepId") or "").strip()
            step = next((item for item in (host.get("attackPathSteps") or []) if isinstance(item, dict) and str(item.get("id") or "").strip() == step_id), None) if host else None
            if not host:
                issues.append(finding_reference_issue("missing-attack-path-asset", "Finding attack-path reference uses a missing asset.", finding, ref, "attack-path"))
            elif not step or step.get("active") is False or step.get("deletedAt"):
                issues.append(finding_reference_issue("missing-attack-path-step", "Finding attack-path step is missing or inactive.", finding, ref, "attack-path"))
            elif ref.get("attemptId") and str(step.get("sourceExploitAttemptId") or step.get("attemptId") or "").strip() != str(ref.get("attemptId")).strip():
                issues.append(finding_reference_issue("attack-path-investigation-mismatch", "Finding attack-path step belongs to a different investigation.", finding, ref, "attack-path"))
    issues.sort(key=lambda item: (item.get("assetId", ""), ",".join(item.get("findingIds") or []), item.get("kind", ""), item.get("code", "")))
    return {"valid": not issues, "blocking": bool(issues), "issues": issues}


def validate_canonical_findings(state, options=None):
    resolved_options = dict(options) if isinstance(options, dict) else {}
    if not isinstance(resolved_options.get("requiredPrimary"), list):
        resolved_options["requiredPrimary"] = required_primary_placements(state)
    placements = validate_report_placements(state, resolved_options)
    references = validate_finding_references(state, resolved_options)
    return {
        "valid": placements["valid"] and references["valid"],
        "blocking": placements["blocking"] or references["blocking"],
        "legacyOnly": placements["legacyOnly"],
        "canonical": placements["canonical"],
        "placementCount": placements["placementCount"],
        "issues": placements["issues"] + references["issues"],
        "placements": placements,
        "references": references,
    }


def canonical_stage_managed(state, asset_id, stage):
    asset_id = str(asset_id or "").strip()
    stage = str(stage or "").strip().lower()
    return any(
        str(row["placement"].get("assetId") or row["placement"].get("hostId") or "").strip() == asset_id
        and str(row["placement"].get("stage") or "").strip().lower() == stage
        for row in report_placement_rows(state)
    )


def ordered_report_findings(state, asset_id, stage):
    asset_id = str(asset_id or "").strip()
    stage = str(stage or "").strip().lower()
    role_weight = {"primary": 0, "additional": 1, "supporting": 2}
    rows = []
    for row in report_placement_rows(state):
        placement = row["placement"]
        role = str(placement.get("role") or "").strip().lower()
        if (
            str(placement.get("assetId") or placement.get("hostId") or "").strip() != asset_id
            or str(placement.get("stage") or "").strip().lower() != stage
            or role not in CANONICAL_REPORT_ROLES
            or not row["included"]
        ):
            continue
        rows.append({
            "finding": row["finding"],
            "placement": {
                **placement,
                "id": row["placementId"],
                "assetId": asset_id,
                "stage": stage,
                "role": role,
                "order": placement.get("order"),
            },
        })
    rows.sort(key=lambda row: (
        role_weight[row["placement"]["role"]],
        row["placement"]["order"],
        row["placement"]["id"],
    ))
    return rows


def safe_placement_error(validation):
    first = (validation.get("issues") or [{}])[0]
    clean_token = lambda value: re.sub(r"[^A-Za-z0-9._:-]+", "-", str(value or ""))[:80]
    asset_id = clean_token(first.get("assetId"))
    stage = clean_token(first.get("stage"))
    scope = ", ".join(part for part in (
        f"asset {asset_id}" if asset_id else "",
        f"stage {stage}" if stage else "",
    ) if part)
    return f"Canonical Finding is invalid{f' for {scope}' if scope else ''}: {first.get('code') or 'invalid-finding'}"


def first_evidence_of_type(host, ev_type):
    for ev in host.get("evidence") or []:
        if isinstance(ev, dict) and ev.get("active") is not False and ev.get("type") == ev_type:
            return ev
    return {}


def evidence_notes_for(host, ev_type):
    ev = first_evidence_of_type(host, ev_type)
    pieces = []
    if ev.get("notes"):
        pieces.append(str(ev.get("notes")))
    if ev.get("explanation"):
        pieces.append("Explanation: " + str(ev.get("explanation")))
    if ev.get("fix"):
        pieces.append("Fix: " + str(ev.get("fix")))
    return "\n\n".join(pieces)


def privilege_escalation_notes(host):
    return host.get("reportNotes") or evidence_notes_for(host, "privilege_escalation") or "TBD"



def add_scan_report_item(doc, host, scan_type, label):
    scans = host.get("scans") or {}
    output = scans.get(scan_type, "")
    img = ((host.get("scanImages") or {}).get(scan_type) or {})
    has_img = bool(img.get("screenshotAbsPath") or img.get("screenshotRelPath"))
    if not output and not has_img:
        return
    p = doc.add_paragraph(label)
    if p.runs:
        p.runs[0].bold = True
    add_labeled_para(doc, "Command", scan_command(scan_type, host))
    if img.get("screenshotAbsPath") and Path(img.get("screenshotAbsPath")).exists():
        try:
            doc.add_picture(img.get("screenshotAbsPath"), width=Inches(5.7))
            return
        except Exception as e:
            add_labeled_para(doc, "Screenshot", f"Could not embed image: {img.get('screenshotAbsPath')} ({e})")
    elif img.get("screenshotRelPath"):
        add_labeled_para(doc, "Screenshot", img.get("screenshotRelPath"))
        return
    if output:
        add_code_block(doc, output)

def add_finding_detail_fields(doc, host, ev=None):
    """Write the OSCP-style finding fields for Initial Access / Privilege Escalation sections."""
    ev = ev or {}
    explanation = ev.get("explanation") or host.get("vulnExplanation") or "TBD"
    fix = ev.get("fix") or host.get("vulnFix") or "TBD"
    severity = ev.get("severity") or host.get("severity") or "TBD"
    steps = ev.get("notes") or host_attack_path_text(host) or "TBD"

    add_labeled_para_with_gap(doc, "Vulnerability Explanation", explanation)
    add_labeled_para_with_gap(doc, "Vulnerability Fix", fix)
    add_severity_para_with_gap(doc, severity)
    add_labeled_para_with_gap(doc, "Steps to reproduce the attack", steps)


def canonical_url_path_label(url, path):
    address = str(url or "").strip()
    suffix = str(path or "").strip()
    if not address:
        return suffix
    if not suffix:
        return address
    try:
        from urllib.parse import urlsplit
        if urlsplit(address).path == suffix:
            return address
    except Exception:
        pass
    return address if address.endswith(suffix) else f"{address} {suffix}"


def add_canonical_evidence_items(doc, rows):
    for row in rows:
        ev = row.get("evidence") if isinstance(row, dict) and isinstance(row.get("evidence"), dict) else {}
        add_labeled_para_with_gap(doc, "Evidence", ev.get("title") or evidence_type_label(ev.get("type")))
        if ev.get("proofFile"):
            add_labeled_para_with_gap(doc, "Proof File", ev.get("proofFile"))
        if ev.get("proofValue"):
            add_labeled_para_with_gap(doc, "Proof Value", ev.get("proofValue"))
        if ev.get("lootPath"):
            add_labeled_para_with_gap(doc, "Credential/File Path", ev.get("lootPath"))
        if ev.get("credentials"):
            add_labeled_para_with_gap(doc, "Credentials Found", "\n".join(credential_pair_lines(ev.get("credentials"))))
        if ev.get("credentialContext"):
            add_labeled_para_with_gap(doc, "Credential Context / Reuse Notes", ev.get("credentialContext"))
        if ev.get("sourceHost") or ev.get("targetHost"):
            add_labeled_para_with_gap(doc, "Lateral Movement Path", f"{ev.get('sourceHost') or 'source TBD'} -> {ev.get('targetHost') or 'target TBD'}")
        if ev.get("method"):
            add_labeled_para_with_gap(doc, "Method / Protocol", ev.get("method"))
        if ev.get("credentialUsed"):
            add_labeled_para_with_gap(doc, "Credential Used", ev.get("credentialUsed"))
        if ev.get("notes"):
            add_labeled_para_with_gap(doc, "Evidence Notes", ev.get("notes"))
        if ev.get("identityOutput"):
            add_labeled_para_with_gap(doc, "Identity Output", ev.get("identityOutput"))
        if ev.get("networkOutput"):
            add_labeled_para_with_gap(doc, "Target Identity Output", ev.get("networkOutput"))
        img = ev.get("screenshotAbsPath")
        if img and Path(img).exists():
            try:
                doc.add_picture(img, width=Inches(5.7))
                add_report_spacer(doc, GAP_BLOCK)
            except Exception as error:
                add_labeled_para_with_gap(doc, "Screenshot", f"Could not embed image: {img} ({error})")
        elif ev.get("screenshotRelPath"):
            add_labeled_para_with_gap(doc, "Screenshot", ev.get("screenshotRelPath"))
        add_report_spacer(doc, GAP_BLOCK)


def add_canonical_finding_fields(doc, state, host, finding):
    description = finding.get("description") or finding.get("summary") or "TBD"
    root_cause = finding.get("rootCause") or ""
    impact = finding.get("impact") or ""
    remediation = finding.get("remediation") or finding.get("fix") or "TBD"
    references = finding.get("references") or ""
    endpoint = finding.get("endpoint") if isinstance(finding.get("endpoint"), dict) else {}
    affected = []
    asset = " / ".join(str(item).strip() for item in (host.get("ip"), host.get("hostname")) if str(item or "").strip()) or str(host.get("id") or "")
    if asset:
        affected.append(f"Asset: {asset}")
    port_service = " · ".join(item for item in (
        f"{endpoint.get('port')}/{str(endpoint.get('protocol') or 'tcp').upper()}" if endpoint.get("port") else "",
        str(endpoint.get("service") or "").strip(),
    ) if item)
    if port_service:
        affected.append(f"Port / Service: {port_service}")
    url_path = canonical_url_path_label(endpoint.get("url"), endpoint.get("path"))
    if url_path:
        affected.append(f"URL / Path: {url_path}")
    if endpoint.get("httpMethod"):
        affected.append(f"Method: {endpoint.get('httpMethod')}")
    if endpoint.get("parameter"):
        affected.append(f"Parameter / Input: {endpoint.get('parameter')}")
    product_version = " ".join(str(item).strip() for item in (endpoint.get("product"), endpoint.get("version")) if str(item or "").strip())
    if product_version:
        affected.append(f"Product / Version: {product_version}")
    identifiers = " · ".join(str(item).strip() for item in (finding.get("cve"), finding.get("cwe"), finding.get("identifier")) if str(item or "").strip())
    if identifiers:
        affected.append(f"Identifiers: {identifiers}")
    if affected:
        add_labeled_para_with_gap(doc, "Affected Location", "\n".join(affected))
    add_labeled_para_with_gap(doc, "Technical Description", description)
    if root_cause:
        add_labeled_para_with_gap(doc, "Root Cause", root_cause)
    if impact:
        add_labeled_para_with_gap(doc, "Impact", impact)
    add_labeled_para_with_gap(doc, "Remediation", remediation)
    add_severity_para_with_gap(doc, finding.get("severity") or "not-assessed")
    if finding.get("reportNarrative"):
        add_labeled_para_with_gap(doc, "Report Note", finding.get("reportNarrative"))
    evidence_rows = canonical_evidence_rows(state, finding)
    if evidence_rows:
        add_report_heading(doc, "Supporting Evidence", level=5)
        add_canonical_evidence_items(doc, evidence_rows)
    if references:
        add_labeled_para_with_gap(doc, "References", references)


def add_canonical_finding_stage(doc, state, host, stage):
    asset_id = str(host.get("id") or "").strip()
    if not canonical_stage_managed(state, asset_id, stage):
        return False
    rows = ordered_report_findings(state, asset_id, stage)
    main = [row for row in rows if row["placement"]["role"] != "supporting"]
    supporting = [row for row in rows if row["placement"]["role"] == "supporting"]
    for row in main:
        finding = row["finding"]
        role = row["placement"]["role"].title()
        add_report_heading(doc, f"{role} Finding - {finding.get('title') or 'Title required'}", level=4)
        add_canonical_finding_fields(doc, state, host, finding)
    if supporting:
        add_report_heading(doc, "Supporting Findings", level=4)
        for row in supporting:
            finding = row["finding"]
            add_report_heading(doc, finding.get("title") or "Title required", level=5)
            add_canonical_finding_fields(doc, state, host, finding)
    return True


def add_canonical_reproduction(doc, state, host, stage):
    asset_id = str(host.get("id") or "").strip()
    if not canonical_stage_managed(state, asset_id, stage):
        return False
    rows = ordered_report_findings(state, asset_id, stage)
    for row in rows:
        finding = row["finding"]
        add_report_heading(doc, finding.get("title") or "Finding", level=4)
        for index, activity in enumerate(canonical_reproduction_rows(state, finding), 1):
            add_labeled_para_with_gap(doc, f"Activity {index}", activity.get("body") or "TBD")
            if activity.get("reportAddendum"):
                add_labeled_para_with_gap(doc, "Transition Note", activity.get("reportAddendum"))
    return True


def add_independent_host_section(doc, idx, host, state=None):
    state = state if isinstance(state, dict) else {}
    base = f"4.{idx}"
    add_report_heading(doc, f"{base} {target_title(idx, host)}", level=2)
    add_report_spacer(doc, GAP_TIGHT)
    os_facts = host_os_report_facts(host)
    doc.add_paragraph(f"Operating System: {os_facts['label']} ({os_facts['confidence']})")
    if os_facts["sources"]:
        doc.add_paragraph(f"OS Sources: {' · '.join(os_facts['sources'])}")

    finding_title = host.get("findingTitle") or "Finding TBD"
    privesc_title = legacy_privesc_title(host)

    initial_managed = canonical_stage_managed(state, str(host.get("id") or ""), "initial-access")
    add_report_heading(doc, f"{base}.1 Initial Access{'' if initial_managed else f' - {finding_title}'}", level=3)
    add_report_spacer(doc, GAP_TIGHT)
    if add_canonical_finding_stage(doc, state, host, "initial-access"):
        pass
    elif first_evidence_of_type(host, "initial_access"):
        add_evidence_items(doc, host, {"initial_access"})
    else:
        add_finding_detail_fields(doc, host)
    add_report_spacer(doc, GAP_MAJOR)

    add_report_heading(doc, f"{base}.2 Service Enumeration", level=3)
    add_report_spacer(doc, GAP_TIGHT)
    add_port_summary_table(doc, host)
    add_evidence_items(doc, host, {"service"})
    add_report_spacer(doc, GAP_BLOCK)
    add_scan_report_item(doc, host, "port", "Port Scan")
    add_report_spacer(doc, GAP_TIGHT)
    add_scan_report_item(doc, host, "tcp", "Nmap Full Scan")
    add_report_spacer(doc, GAP_TIGHT)
    add_scan_report_item(doc, host, "udp", "UDP Scan")
    add_report_spacer(doc, GAP_MAJOR)

    add_report_heading(doc, f"{base}.3 Initial Access Walkthrough", level=3)
    add_report_spacer(doc, GAP_TIGHT)
    if not add_canonical_reproduction(doc, state, host, "initial-access"):
        doc.add_paragraph(host_attack_path_text(host) or "TBD")
    add_report_spacer(doc, GAP_MAJOR)

    privesc_managed = canonical_stage_managed(state, str(host.get("id") or ""), "privilege-escalation")
    add_report_heading(doc, f"{base}.4 Privilege Escalation{'' if privesc_managed else f' - {privesc_title}'}", level=3)
    add_report_spacer(doc, GAP_TIGHT)
    privesc_ev = first_evidence_of_type(host, "privilege_escalation")
    if add_canonical_finding_stage(doc, state, host, "privilege-escalation"):
        add_canonical_reproduction(doc, state, host, "privilege-escalation")
    elif direct_elevated_initial_access(host):
        add_labeled_para_with_gap(doc, "Privilege Escalation", "Not applicable — elevated access was obtained during initial access.")
    elif privesc_ev:
        add_evidence_items(doc, host, {"privilege_escalation"})
    else:
        add_labeled_para_with_gap(doc, "Privilege Escalation Notes", privilege_escalation_notes(host))
    add_report_spacer(doc, GAP_MAJOR)

    add_report_heading(doc, f"{base}.5 Post Exploitation", level=3)
    add_report_spacer(doc, GAP_TIGHT)
    if add_canonical_finding_stage(doc, state, host, "lateral-movement"):
        add_canonical_reproduction(doc, state, host, "lateral-movement")
    if add_canonical_finding_stage(doc, state, host, "other"):
        add_canonical_reproduction(doc, state, host, "other")
    add_proof_block(doc, host)
    doc.add_page_break()

def add_ad_host_section(doc, idx, host, state=None):
    state = state if isinstance(state, dict) else {}
    base = f"5.{idx}"
    add_report_heading(doc, f"{base} {host_name_ip(host)}", level=2)
    add_report_spacer(doc, GAP_TIGHT)
    os_facts = host_os_report_facts(host)
    doc.add_paragraph(f"Operating System: {os_facts['label']} ({os_facts['confidence']})")
    if os_facts["sources"]:
        doc.add_paragraph(f"OS Sources: {' · '.join(os_facts['sources'])}")

    finding_title = host.get("findingTitle") or "Initial Access TBD"
    privesc_title = legacy_privesc_title(host)

    initial_managed = canonical_stage_managed(state, str(host.get("id") or ""), "initial-access")
    add_report_heading(doc, f"{base}.1 Initial Access{'' if initial_managed else f' - {finding_title}'}", level=3)
    add_report_spacer(doc, GAP_TIGHT)
    if add_canonical_finding_stage(doc, state, host, "initial-access"):
        add_canonical_reproduction(doc, state, host, "initial-access")
    elif first_evidence_of_type(host, "initial_access"):
        add_evidence_items(doc, host, {"initial_access"})
    else:
        add_finding_detail_fields(doc, host)
    add_report_spacer(doc, GAP_MAJOR)

    privesc_managed = canonical_stage_managed(state, str(host.get("id") or ""), "privilege-escalation")
    add_report_heading(doc, f"{base}.2 Privilege Escalation{'' if privesc_managed else f' - {privesc_title}'}", level=3)
    add_report_spacer(doc, GAP_TIGHT)
    privesc_ev = first_evidence_of_type(host, "privilege_escalation")
    if add_canonical_finding_stage(doc, state, host, "privilege-escalation"):
        add_canonical_reproduction(doc, state, host, "privilege-escalation")
    elif direct_elevated_initial_access(host):
        add_labeled_para_with_gap(doc, "Privilege Escalation", "Not applicable — elevated access was obtained during initial access.")
    elif privesc_ev:
        add_evidence_items(doc, host, {"privilege_escalation"})
    else:
        add_labeled_para_with_gap(doc, "Privilege Escalation Notes", privilege_escalation_notes(host))
    add_report_spacer(doc, GAP_MAJOR)

    add_report_heading(doc, f"{base}.3 Post-Exploitation", level=3)
    add_report_spacer(doc, GAP_TIGHT)
    if add_canonical_finding_stage(doc, state, host, "lateral-movement"):
        add_canonical_reproduction(doc, state, host, "lateral-movement")
    if add_canonical_finding_stage(doc, state, host, "other"):
        add_canonical_reproduction(doc, state, host, "other")
    add_proof_block(doc, host)
    doc.add_page_break()

def proof_value_from_evidence(host, proof_type):
    proof_file = "local.txt" if proof_type == "local" else "proof.txt"
    for ev in host.get("evidence") or []:
        if ev.get("type") == proof_type or ev.get("proofFile") == proof_file:
            value = str(ev.get("proofValue") or "").strip()
            if value:
                return value
    return ""


def add_proof_block(doc, host):
    local_proof = str(host.get("localProof", "")).strip() or proof_value_from_evidence(host, "local")
    system_proof = str(host.get("proofTxt", "")).strip() or proof_value_from_evidence(host, "proof")
    if local_proof:
        add_labeled_para_with_gap(doc, "Local Proof / local.txt", local_proof)
    if system_proof:
        add_labeled_para_with_gap(doc, "System Proof / proof.txt", system_proof)
    if not local_proof and not system_proof:
        doc.add_paragraph("Proof: TBD")
        add_report_spacer(doc, GAP_FIELD)
    if host.get("reportNotes"):
        add_labeled_para_with_gap(doc, "Post-Exploitation Notes", host.get("reportNotes"))
    add_evidence_items(doc, host, {"local", "proof", "post_exploitation", "credential_dump", "lateral_movement"})
    if not (host.get("evidence") or []):
        doc.add_paragraph("System Proof Screenshot: TBD")
    add_report_spacer(doc, GAP_BLOCK)


def resolve_report_image_path(project_name, host, image):
    if not isinstance(image, dict):
        return None
    project_root = STORE.project_root(project_name).resolve()
    filename = str(image.get("screenshotFilename") or "").strip()
    candidates = []
    if filename:
        candidates.append((internal_host_dir(project_name, host) / "screenshots" / safe_name(filename)).resolve())
    raw_path = str(image.get("screenshotAbsPath") or "").strip()
    if raw_path:
        try:
            candidates.append(Path(raw_path).expanduser().resolve())
        except (OSError, RuntimeError):
            pass
    expected_hash = str(image.get("screenshotSha256") or image.get("sha256") or "").strip().lower()
    for candidate in dict.fromkeys(candidates):
        if (
            not path_is_inside(candidate, project_root)
            or "screenshots" not in {part.casefold() for part in candidate.parts}
            or candidate.suffix.lower() not in IMAGE_EXTS
            or not candidate.is_file()
            or candidate.stat().st_size <= 0
            or candidate.stat().st_size > 25 * 1024 * 1024
        ):
            continue
        if filename and candidate.name.casefold() != safe_name(filename).casefold():
            continue
        if expected_hash and re.fullmatch(r"[a-f0-9]{64}", expected_hash) and file_sha256(candidate) != expected_hash:
            continue
        try:
            validate_image_bytes(candidate.read_bytes(), candidate.suffix)
        except (OSError, ValueError):
            continue
        return candidate
    return None


def sanitize_report_image_references(payload):
    state = deepcopy(payload) if isinstance(payload, dict) else {}
    project_name = str(state.get("projectName") or "").strip()
    hosts = state.get("hosts")
    host_rows = hosts.values() if isinstance(hosts, dict) else hosts if isinstance(hosts, list) else []
    for host in host_rows:
        if not isinstance(host, dict):
            continue
        scan_images = host.get("scanImages")
        if isinstance(scan_images, dict):
            for image in scan_images.values():
                if not isinstance(image, dict):
                    continue
                resolved = resolve_report_image_path(project_name, host, image)
                if resolved:
                    image["screenshotAbsPath"] = str(resolved)
                else:
                    image.pop("screenshotAbsPath", None)
        for evidence in host.get("evidence") or []:
            if not isinstance(evidence, dict):
                continue
            resolved = resolve_report_image_path(project_name, host, evidence)
            if resolved:
                evidence["screenshotAbsPath"] = str(resolved)
            else:
                evidence.pop("screenshotAbsPath", None)
    return state


def generate_report_docx(payload):
    if Document is None:
        raise RuntimeError("python-docx is not installed. Run: pip install python-docx")
    payload = sanitize_report_image_references(payload)
    canonical_validation = validate_canonical_findings(payload)
    if not canonical_validation["valid"]:
        raise ValueError(safe_placement_error(canonical_validation))

    project = payload.get("projectName") or "OSCP Practice Report"
    tester = payload.get("testerName") or "Student"
    first = first_name(tester)
    email = payload.get("studentEmail") or "student@example.com"
    osid = payload.get("osid") or "XXXXX"
    if isinstance(payload.get("hosts"), dict):
        hosts = []
        for key, raw_host in payload.get("hosts", {}).items():
            host = dict(raw_host) if isinstance(raw_host, dict) else {}
            if not str(host.get("id") or "").strip():
                host["id"] = str(key)
            hosts.append(host)
    else:
        hosts = payload.get("hosts", [])
    independent = [h for h in hosts if h.get("group", "independent") != "ad"]
    ad_hosts = [h for h in hosts if h.get("group") == "ad" or h.get("isDomain") or h.get("isDc")]
    total, gained, elevated = proof_counts(hosts)

    config = payload.get("engagementConfig") or {}
    reporting = config.get("reporting") or {} if isinstance(config, dict) else {}
    template_id = payload.get("reportTemplateId") or reporting.get("templateId") or ""
    template_mode = reporting.get("templateMode") or "default"
    template_path = ""
    if template_mode != "later":
        registry = load_report_template_registry()
        selected_id = template_id or registry.get("defaultTemplateId") or ""
        if selected_id:
            try: template_path = str(report_template_path(selected_id)[0])
            except Exception: template_path = ""
    using_template = bool(template_path and Path(template_path).exists())

    if using_template:
        doc = Document(template_path)
        ensure_report_styles(doc)
        replace_text_in_doc(doc, {
            "student@youremailaddress.com": email,
            "OSID: XXXXX": f"OSID: {osid}",
        })
        keep_cover_page_only(doc)
        add_static_toc_page(doc, independent, ad_hosts, payload)
    else:
        doc = Document()
        ensure_report_styles(doc)
        sec = doc.sections[0]
        sec.top_margin = Inches(0.75); sec.bottom_margin = Inches(0.65); sec.left_margin = Inches(0.75); sec.right_margin = Inches(0.75)
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("OffSec Certified Professional")
        r.bold = True; r.font.size = Pt(28)
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("Exam Report")
        r.bold = True; r.font.size = Pt(24)
        for line in ["v.2.0", email, f"OSID: {osid}", project, f"Prepared by: {tester}"]:
            p = doc.add_paragraph(line); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_page_break()
        add_static_toc_page(doc, independent, ad_hosts, payload)

    ensure_report_styles(doc)

    add_report_heading(doc, "1 OffSec Certified Professional Exam Report", level=1)
    add_report_heading(doc, "1.1 Introduction", level=2)
    doc.add_paragraph("The OffSec Certified Professional exam report contains the efforts conducted during the assessment. This report should document the items used to complete the assessment and provide enough detail to validate the methodology, technical findings, exploitation steps, and proof collection.")

    add_report_heading(doc, "1.2 Objective", level=2)
    doc.add_paragraph("The objective of this assessment is to perform an internal penetration test against the scoped lab or exam network. The tester is tasked with following a methodical approach to identifying systems, enumerating exposed services, validating vulnerabilities, gaining access where possible, and documenting all relevant evidence.")

    add_report_heading(doc, "1.3 Requirements", level=2)
    list_style = "List Bullet" if get_style(doc, "List Bullet") is not None else None
    for item in [
        "Overall High-Level Summary and Recommendations (non-technical)",
        "Methodology walkthrough and detailed outline of steps taken",
        "Each finding with included screenshots, walkthrough, sample commands, and proof.txt or local.txt where applicable",
        "Any additional items that were not included elsewhere in the report",
    ]:
        doc.add_paragraph(item, style=list_style)

    add_report_heading(doc, "2 High-Level Summary", level=1)
    doc.add_paragraph(f"{tester} was tasked with performing an internal penetration test against the scoped environment. The focus of this test was to identify live systems, enumerate exposed services, validate attack paths, obtain access where possible, and report the findings in a reproducible manner.")
    doc.add_paragraph(f"During the assessment, {first} gained access to {gained} out of {total} scoped systems based on submitted local.txt or proof.txt evidence. {first} obtained administrative or root-level access to {elevated} out of {total} scoped systems based on submitted proof.txt evidence.")

    add_report_heading(doc, "2.1 Recommendations", level=2)
    doc.add_paragraph(f"{first} recommends remediating the vulnerabilities identified during testing, removing unnecessary exposed services, enforcing strong authentication, applying vendor patches, and reviewing least-privilege configurations. Systems should remain on a regular patching and configuration review schedule to reduce exposure to newly discovered vulnerabilities.")

    add_report_heading(doc, "3 Methodologies", level=1)
    doc.add_paragraph(f"{first} used a methodical penetration testing approach to identify, enumerate, exploit, and document the scoped systems. The following sections summarize the methodology used during the assessment.")

    add_report_heading(doc, "3.1 Information Gathering", level=2)
    doc.add_paragraph("The information gathering phase focused on identifying the scope of the assessment and confirming which systems were reachable.")
    doc.add_paragraph("Exam Network:")
    doc.add_paragraph(ip_list(hosts))

    add_report_heading(doc, "3.2 Service Enumeration", level=2)
    doc.add_paragraph("The service enumeration phase focused on gathering information about exposed services on each system. This information was used to identify potential attack vectors, authentication surfaces, vulnerable applications, and misconfigurations.")

    add_report_heading(doc, "3.3 Penetration", level=2)
    doc.add_paragraph(f"The penetration phase focused on validating attack paths and obtaining initial access where possible. During this assessment, {first} gained access to {gained} out of {total} scoped systems based on provided local.txt or proof.txt evidence.")

    add_report_heading(doc, "3.4 Maintaining Access", level=2)
    doc.add_paragraph("The maintaining access phase focuses on ensuring that access obtained during testing can be reliably documented and reproduced during the assessment window. Any temporary access methods, payloads, sessions, or credentials used during testing should be documented in the relevant host section.")

    add_report_heading(doc, "3.5 House Cleaning", level=2)
    doc.add_paragraph("The house cleaning phase ensures that remnants of testing are identified and removed where applicable. Temporary files, payloads, test accounts, scripts, and services used during the assessment should be documented and cleaned up when allowed by the rules of engagement.")

    # Always start the Independent Challenges section on a new page.
    p_independent = add_report_heading(doc, "4 Independent Challenges", level=1)
    p_independent.paragraph_format.page_break_before = True
    add_report_spacer(doc, 10)
    if independent:
        for i, h in enumerate(independent, 1):
            add_independent_host_section(doc, i, h, payload)
    else:
        doc.add_paragraph("No independent challenge hosts added.")

    add_report_heading(doc, "5 Active Directory Set", level=1)
    if ad_hosts:
        add_port_summary_table(doc, ad_hosts)
        for i, h in enumerate(ad_hosts, 1):
            add_ad_host_section(doc, i, h, payload)
    else:
        doc.add_paragraph("No Active Directory hosts added.")

    bio = BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio.getvalue()



class MultipartFile:
    def __init__(self, filename, data):
        self.filename = filename
        self.file = BytesIO(data)


def parse_multipart_form(headers, body):
    content_type = headers.get("Content-Type", "")
    raw = (f"Content-Type: {content_type}\r\nMIME-Version: 1.0\r\n\r\n").encode("utf-8") + body
    msg = BytesParser(policy=email_default).parsebytes(raw)
    fields = {}
    files = {}
    if not msg.is_multipart():
        return fields, files
    for part in msg.iter_parts():
        disp = part.get("Content-Disposition", "")
        if "form-data" not in disp:
            continue
        name = part.get_param("name", header="content-disposition")
        filename = part.get_filename()
        data = part.get_payload(decode=True) or b""
        if filename:
            files[name] = MultipartFile(filename, data)
        else:
            fields[name] = data.decode(part.get_content_charset() or "utf-8", errors="replace")
    return fields, files


def validate_recon_payload(payload, filename="recon.json"):
    """Validate one independently importable AEROS Linux recon document."""
    if not isinstance(payload, dict):
        raise ValueError(f"{filename}: the JSON root must be an object")

    schema = payload.get("schema")
    if not isinstance(schema, dict):
        raise ValueError(f"{filename}: missing schema object")
    schema_name = str(schema.get("name") or "").strip()
    if schema_name not in SUPPORTED_RECON_SCHEMAS:
        supported = ", ".join(sorted(SUPPORTED_RECON_SCHEMAS))
        raise ValueError(f"{filename}: unsupported schema.name '{schema_name or 'missing'}' (supported: {supported})")
    schema_version = str(schema.get("version") or "").strip()
    if not schema_version or schema_version.split(".", 1)[0] != RECON_SCHEMA_MAJOR:
        raise ValueError(f"{filename}: unsupported {schema_name} schema version '{schema_version or 'missing'}'")

    collector = payload.get("collector")
    if not isinstance(collector, dict) or str(collector.get("platform") or "").lower() != "linux":
        raise ValueError(f"{filename}: only Linux AEROS recon payloads are supported")
    if not isinstance(payload.get("collection"), dict):
        raise ValueError(f"{filename}: missing collection object")
    if not isinstance(payload.get("target"), dict):
        raise ValueError(f"{filename}: missing target object")
    if not isinstance(payload.get("sections"), dict):
        raise ValueError(f"{filename}: missing sections object")
    return payload


def parse_recon_json_bytes(data, filename="recon.json"):
    if not data:
        raise ValueError(f"{filename}: file is empty")
    if len(data) > MAX_RECON_JSON_BYTES:
        raise ValueError(f"{filename}: JSON payload exceeds the {MAX_RECON_JSON_BYTES // (1024 * 1024)} MB limit")
    try:
        decoded = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError(f"{filename}: JSON must be UTF-8 ({exc})") from exc
    try:
        payload = json.loads(decoded)
    except json.JSONDecodeError as exc:
        raise ValueError(f"{filename}: invalid JSON at line {exc.lineno}, column {exc.colno}") from exc
    validate_recon_payload(payload, filename)
    return {"filename": Path(filename).name or "aeros-recon.json", "data": payload}


def safe_recon_zip_member(info):
    raw_name = str(info.filename or "")
    if not raw_name or "\x00" in raw_name:
        raise ValueError("ZIP contains an invalid member name")
    normalized = raw_name.replace("\\", "/")
    member_path = PurePosixPath(normalized)
    if member_path.is_absolute() or re.match(r"^[A-Za-z]:", normalized):
        raise ValueError(f"ZIP member uses an absolute path: {raw_name}")
    if any(part in {"", ".", ".."} for part in member_path.parts):
        raise ValueError(f"ZIP member uses an unsafe path: {raw_name}")
    mode = (info.external_attr >> 16) & 0xFFFF
    if mode and stat.S_ISLNK(mode):
        raise ValueError(f"ZIP member is a symbolic link: {raw_name}")
    if info.flag_bits & 0x1:
        raise ValueError(f"ZIP member is encrypted: {raw_name}")
    if info.compress_type not in {zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED}:
        raise ValueError(f"ZIP member uses an unsupported compression method: {raw_name}")
    return normalized


def recon_target_sets(payloads):
    hostnames = set()
    addresses = set()
    collection_ids = set()
    for item in payloads:
        data = item["data"]
        target = data.get("target") or {}
        hostname = str(target.get("hostname") or "").strip().lower().rstrip(".")
        address = str(target.get("primaryAddress") or "").strip().lower()
        if hostname:
            hostnames.add(hostname)
        if address:
            addresses.add(address)
        collection_id = str((data.get("collection") or {}).get("id") or "").strip()
        if collection_id:
            if collection_id in collection_ids:
                raise ValueError(f"Duplicate collection.id in upload: {collection_id}")
            collection_ids.add(collection_id)
    if len(hostnames) > 1 or len(addresses) > 1:
        raise ValueError("A recon ZIP or upload batch must describe one target host")
    return hostnames, addresses


def parse_recon_zip_bytes(data, archive_name="aeros-recon.zip"):
    try:
        archive = zipfile.ZipFile(BytesIO(data), "r")
    except zipfile.BadZipFile as exc:
        raise ValueError(f"{archive_name}: invalid ZIP archive") from exc

    with archive:
        members = archive.infolist()
        if len(members) > MAX_RECON_ARCHIVE_MEMBERS:
            raise ValueError(f"{archive_name}: ZIP contains too many members (maximum {MAX_RECON_ARCHIVE_MEMBERS})")

        json_members = []
        seen_names = set()
        total_uncompressed = 0
        for info in members:
            normalized = safe_recon_zip_member(info)
            if info.is_dir():
                continue
            lower = normalized.lower()
            if lower.startswith("__macosx/") or lower.endswith("/.ds_store") or lower == ".ds_store":
                continue
            suffix = PurePosixPath(normalized).suffix.lower()
            if suffix == ".zip":
                raise ValueError(f"{archive_name}: nested ZIP archives are not supported ({normalized})")
            if suffix != ".json":
                raise ValueError(f"{archive_name}: ZIP member is not JSON ({normalized})")
            if normalized in seen_names:
                raise ValueError(f"{archive_name}: duplicate ZIP member name ({normalized})")
            seen_names.add(normalized)
            if info.file_size <= 0:
                raise ValueError(f"{archive_name}: JSON member is empty ({normalized})")
            if info.file_size > MAX_RECON_JSON_BYTES:
                raise ValueError(f"{archive_name}: JSON member exceeds the per-file limit ({normalized})")
            total_uncompressed += info.file_size
            if total_uncompressed > MAX_RECON_TOTAL_JSON_BYTES:
                raise ValueError(f"{archive_name}: expanded JSON data exceeds the {MAX_RECON_TOTAL_JSON_BYTES // (1024 * 1024)} MB limit")
            if info.compress_size == 0 or info.file_size / max(info.compress_size, 1) > MAX_RECON_COMPRESSION_RATIO:
                raise ValueError(f"{archive_name}: suspicious compression ratio for {normalized}")
            json_members.append((info, normalized))

        if not json_members:
            raise ValueError(f"{archive_name}: ZIP contains no JSON payloads")
        if len(json_members) > MAX_RECON_PAYLOADS:
            raise ValueError(f"{archive_name}: ZIP contains too many JSON payloads (maximum {MAX_RECON_PAYLOADS})")
        bad_member = archive.testzip()
        if bad_member:
            raise ValueError(f"{archive_name}: CRC check failed for {bad_member}")

        payloads = []
        for info, normalized in json_members:
            try:
                member_data = archive.read(info)
            except (RuntimeError, zipfile.BadZipFile) as exc:
                raise ValueError(f"{archive_name}: could not read {normalized}") from exc
            payloads.append(parse_recon_json_bytes(member_data, normalized))

    recon_target_sets(payloads)
    legacy_schemas = sorted({
        str((item["data"].get("schema") or {}).get("name") or "")
        for item in payloads
        if str((item["data"].get("schema") or {}).get("name") or "") != "aeros-recon"
    })
    return {
        "sourceType": "zip",
        "sourceName": Path(archive_name).name or "aeros-recon.zip",
        "payloads": payloads,
        "legacySchemas": legacy_schemas,
    }


def parse_recon_upload(file_item):
    if not file_item or not file_item.filename:
        raise ValueError("Choose an AEROS recon ZIP archive or JSON file")
    original_name = Path(file_item.filename).name
    data = file_item.file.read()
    if not data:
        raise ValueError("The selected recon file is empty")
    suffix = Path(original_name).suffix.lower()
    looks_like_zip = data.startswith((b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08"))
    if suffix == ".zip" or looks_like_zip:
        return parse_recon_zip_bytes(data, original_name)
    if suffix != ".json":
        raise ValueError("Supported recon upload types are .zip and .json")
    payload = parse_recon_json_bytes(data, original_name)
    schema_name = str((payload["data"].get("schema") or {}).get("name") or "")
    return {
        "sourceType": "json",
        "sourceName": original_name,
        "payloads": [payload],
        "legacySchemas": [] if schema_name == "aeros-recon" else [schema_name],
    }


WINDOWS_DEVICE_NAMES = {
    "con", "prn", "aux", "nul", "clock$",
    *(f"com{number}" for number in range(1, 10)),
    *(f"lpt{number}" for number in range(1, 10)),
}
NESTED_ARCHIVE_SUFFIXES = {".zip", ".7z", ".rar", ".tar", ".gz", ".tgz", ".bz2", ".xz"}
IGNORED_IMPORT_NAMES = {".ds_store", "thumbs.db", "desktop.ini"}


def safe_import_zip_member(info):
    normalized = safe_recon_zip_member(info)
    member_path = PurePosixPath(normalized)
    for part in member_path.parts:
        trimmed = part.rstrip(" .")
        stem = trimmed.split(".", 1)[0].lower()
        if not trimmed or trimmed != part or ":" in part or stem in WINDOWS_DEVICE_NAMES:
            raise ValueError(f"ZIP member uses an unsafe Windows path: {info.filename}")
    mode = (info.external_attr >> 16) & 0xFFFF
    file_type = stat.S_IFMT(mode)
    if file_type and not (stat.S_ISREG(mode) or stat.S_ISDIR(mode)):
        raise ValueError(f"ZIP member is not a regular file or directory: {info.filename}")
    return normalized


def import_intake_limits():
    return {
        "maximumArchiveBytes": MAX_IMPORT_ARCHIVE_BYTES,
        "maximumExpandedBytes": MAX_IMPORT_EXPANDED_BYTES,
        "maximumMemberBytes": MAX_IMPORT_MEMBER_BYTES,
        "maximumContainedFiles": MAX_IMPORT_ARCHIVE_MEMBERS,
        "maximumCompressionRatio": MAX_IMPORT_COMPRESSION_RATIO,
        "maximumNestedArchiveDepth": MAX_IMPORT_NESTED_ARCHIVE_DEPTH,
        "previewBytesPerFile": MAX_IMPORT_PREVIEW_BYTES,
    }


def cleanup_import_intake_staging():
    if not IMPORT_INTAKE_STAGING_ROOT.exists():
        return
    cutoff = datetime.datetime.now(datetime.timezone.utc).timestamp() - IMPORT_STAGING_TTL_SECONDS
    for candidate in IMPORT_INTAKE_STAGING_ROOT.iterdir():
        try:
            if candidate.is_dir() and candidate.stat().st_mtime < cutoff:
                shutil.rmtree(candidate)
        except OSError:
            continue


def import_intake_stage(token):
    value = str(token or "").strip().lower()
    if not re.fullmatch(r"[a-f0-9]{32}", value):
        raise ValueError("Invalid import staging token")
    root = (IMPORT_INTAKE_STAGING_ROOT / value).resolve()
    if not path_is_inside(root, IMPORT_INTAKE_STAGING_ROOT.resolve()):
        raise ValueError("Invalid import staging token")
    return root


def import_intake_member_path(token, member):
    root = import_intake_stage(token)
    normalized = str(member or "").replace("\\", "/")
    path = PurePosixPath(normalized)
    if path.is_absolute() or any(part in {"", ".", ".."} for part in path.parts):
        raise ValueError("Invalid staged import member")
    target = (root / Path(*path.parts)).resolve()
    if not path_is_inside(target, root) or not target.exists() or not target.is_file():
        raise FileNotFoundError("Staged import member not found")
    return target


def discard_import_intake(token):
    root = import_intake_stage(token)
    if root.exists():
        shutil.rmtree(root)
        return True
    return False


def inspect_import_archive(file_item):
    if not file_item or not file_item.filename:
        raise ValueError("Choose a ZIP archive")
    archive_name = Path(file_item.filename).name or "import-results.zip"
    data = file_item.file.read(MAX_IMPORT_ARCHIVE_BYTES + 1)
    if not data:
        raise ValueError(f"{archive_name}: archive is empty")
    if len(data) > MAX_IMPORT_ARCHIVE_BYTES:
        raise ValueError(
            f"{archive_name}: archive exceeds the {MAX_IMPORT_ARCHIVE_BYTES // (1024 * 1024)} MiB limit"
        )
    if not data.startswith((b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08")):
        raise ValueError(f"{archive_name}: file is not a ZIP archive")
    try:
        archive = zipfile.ZipFile(BytesIO(data), "r")
    except zipfile.BadZipFile as exc:
        raise ValueError(f"{archive_name}: malformed ZIP archive") from exc

    cleanup_import_intake_staging()
    IMPORT_INTAKE_STAGING_ROOT.mkdir(parents=True, exist_ok=True)
    token = uuid.uuid4().hex
    staging_root = import_intake_stage(token)
    staging_root.mkdir(parents=False, exist_ok=False)
    rows = []
    try:
        with archive:
            members = archive.infolist()
            if len(members) > MAX_IMPORT_ARCHIVE_MEMBERS:
                raise ValueError(
                    f"{archive_name}: ZIP contains too many members "
                    f"(maximum {MAX_IMPORT_ARCHIVE_MEMBERS})"
                )
            seen = set()
            expanded = 0
            admitted = []
            for info in members:
                normalized = safe_import_zip_member(info)
                key = normalized.lower()
                if key in seen:
                    raise ValueError(f"{archive_name}: duplicate member path ({normalized})")
                seen.add(key)
                if info.is_dir():
                    continue
                suffix = PurePosixPath(normalized).suffix.lower()
                if suffix in NESTED_ARCHIVE_SUFFIXES:
                    raise ValueError(
                        f"{archive_name}: nested archives are not supported ({normalized}); "
                        f"maximum nested-archive depth is {MAX_IMPORT_NESTED_ARCHIVE_DEPTH}"
                    )
                if info.file_size > MAX_IMPORT_MEMBER_BYTES:
                    raise ValueError(
                        f"{archive_name}: member exceeds the {MAX_IMPORT_MEMBER_BYTES // (1024 * 1024)} MiB limit "
                        f"({normalized})"
                    )
                expanded += info.file_size
                if expanded > MAX_IMPORT_EXPANDED_BYTES:
                    raise ValueError(
                        f"{archive_name}: expanded data exceeds the "
                        f"{MAX_IMPORT_EXPANDED_BYTES // (1024 * 1024)} MiB limit"
                    )
                ratio = info.file_size / max(info.compress_size, 1)
                if info.file_size and ratio > MAX_IMPORT_COMPRESSION_RATIO:
                    raise ValueError(
                        f"{archive_name}: suspicious compression ratio for {normalized} "
                        f"(maximum {MAX_IMPORT_COMPRESSION_RATIO}:1)"
                    )
                admitted.append((info, normalized, ratio))
            bad_member = archive.testzip()
            if bad_member:
                raise ValueError(f"{archive_name}: CRC check failed for {bad_member}")
            for info, normalized, ratio in admitted:
                path = PurePosixPath(normalized)
                ignored = (
                    any(part.lower() == "__macosx" for part in path.parts)
                    or path.name.lower() in IGNORED_IMPORT_NAMES
                )
                if ignored:
                    rows.append({
                        "path": normalized, "name": path.name, "size": info.file_size,
                        "compressedSize": info.compress_size, "compressionRatio": round(ratio, 2),
                        "ignoredMetadata": True, "previewText": "", "signature": "", "sha256": "",
                    })
                    continue
                target = (staging_root / Path(*path.parts)).resolve()
                if not path_is_inside(target, staging_root):
                    raise ValueError(f"{archive_name}: extraction escaped temporary staging")
                target.parent.mkdir(parents=True, exist_ok=True)
                digest = hashlib.sha256()
                preview = bytearray()
                actual = 0
                with archive.open(info, "r") as source, target.open("xb") as destination:
                    while True:
                        chunk = source.read(1024 * 1024)
                        if not chunk:
                            break
                        actual += len(chunk)
                        if actual > MAX_IMPORT_MEMBER_BYTES:
                            raise ValueError(f"{archive_name}: expanded member exceeded its declared size ({normalized})")
                        digest.update(chunk)
                        destination.write(chunk)
                        if len(preview) < MAX_IMPORT_PREVIEW_BYTES:
                            preview.extend(chunk[:MAX_IMPORT_PREVIEW_BYTES - len(preview)])
                if actual != info.file_size:
                    raise ValueError(f"{archive_name}: member size changed during extraction ({normalized})")
                sample = bytes(preview)
                preview_text = ""
                if b"\x00" not in sample[:4096]:
                    try:
                        preview_text = sample.decode("utf-8-sig")
                    except UnicodeDecodeError:
                        preview_text = ""
                rows.append({
                    "path": normalized, "name": path.name, "size": info.file_size,
                    "compressedSize": info.compress_size, "compressionRatio": round(ratio, 2),
                    "ignoredMetadata": False, "previewText": preview_text,
                    "signature": sample[:16].hex(), "sha256": digest.hexdigest(),
                })
        if not any(not row["ignoredMetadata"] for row in rows):
            raise ValueError(f"{archive_name}: ZIP contains no importable files")
        return {
            "token": token,
            "archiveName": archive_name,
            "archiveSize": len(data),
            "members": rows,
            "limits": import_intake_limits(),
        }
    except Exception:
        shutil.rmtree(staging_root, ignore_errors=True)
        raise


IMAGE_EXTS = {".png", ".jpg", ".jpeg"}


def _jpeg_dimensions(data):
    offset = 2
    start_of_frame = {
        0xC0, 0xC1, 0xC2, 0xC3, 0xC5, 0xC6, 0xC7,
        0xC9, 0xCA, 0xCB, 0xCD, 0xCE, 0xCF,
    }
    while offset + 3 < len(data):
        if data[offset] != 0xFF:
            offset += 1
            continue
        while offset < len(data) and data[offset] == 0xFF:
            offset += 1
        if offset >= len(data):
            break
        marker = data[offset]
        offset += 1
        if marker in {0xD8, 0xD9} or 0xD0 <= marker <= 0xD7:
            continue
        if offset + 1 >= len(data):
            break
        length = int.from_bytes(data[offset:offset + 2], "big")
        if length < 2 or offset + length > len(data):
            break
        if marker in start_of_frame and length >= 7:
            height = int.from_bytes(data[offset + 3:offset + 5], "big")
            width = int.from_bytes(data[offset + 5:offset + 7], "big")
            return width, height
        offset += length
    return None


def validate_image_bytes(data, extension):
    """Reject extension-only image uploads before browsers render them."""
    extension = str(extension or "").lower()
    dimensions = None
    if extension == ".png":
        if len(data) < 24 or not data.startswith(b"\x89PNG\r\n\x1a\n") or data[12:16] != b"IHDR":
            raise ValueError("PNG screenshot has an invalid file signature")
        dimensions = (
            int.from_bytes(data[16:20], "big"),
            int.from_bytes(data[20:24], "big"),
        )
    elif extension in {".jpg", ".jpeg"}:
        if not data.startswith(b"\xff\xd8\xff"):
            raise ValueError("JPEG screenshot has an invalid file signature")
        dimensions = _jpeg_dimensions(data)
        if dimensions is None:
            raise ValueError("JPEG screenshot has invalid dimensions")
    else:
        raise ValueError("Unsupported screenshot type")
    width, height = dimensions
    if width <= 0 or height <= 0:
        raise ValueError("Screenshot dimensions are invalid")
    if width > MAX_IMAGE_DIMENSION or height > MAX_IMAGE_DIMENSION or width * height > MAX_IMAGE_PIXELS:
        raise ValueError("Screenshot dimensions exceed the safety limit")
    return dimensions

def evidence_type_label(value):
    return {
        "local": "local.txt proof",
        "proof": "proof.txt proof",
        "initial_access": "Initial Access / Vulnerability",
        "privilege_escalation": "Privilege Escalation",
        "post_exploitation": "Post-Exploitation / Loot",
        "credential_dump": "Credential / Password Dump",
        "lateral_movement": "Lateral Movement",
        "service": "Service Enumeration",
        "custom": "Custom Report Evidence",
    }.get(value, value or "Evidence")







def screenshot_artifact_record(meta, *, logical_path="screenshot", source="screenshot", objective=""):
    if not isinstance(meta, dict):
        return {}
    original = str(meta.get("originalFilename") or meta.get("screenshotFilename") or "screenshot.png")
    stored = str(meta.get("storedFilename") or meta.get("screenshotFilename") or "")
    revisions = meta.get("revisions") if isinstance(meta.get("revisions"), list) else meta.get("screenshotRevisions") if isinstance(meta.get("screenshotRevisions"), list) else []
    return {
        **meta,
        "id": str(meta.get("artifactId") or meta.get("id") or ""),
        "kind": "screenshot-artifact",
        "originalFilename": original,
        "storedFilename": stored,
        "logicalPath": str(meta.get("logicalPath") or logical_path),
        "source": str(meta.get("source") or source),
        "objective": str(meta.get("objective") or objective),
        "artifactType": str(meta.get("artifactType") or f"screenshot-{Path(original).suffix.lower().lstrip('.') or 'image'}"),
        "sha256": str(meta.get("sha256") or ""),
        "currentRevision": meta.get("currentRevision") or 1,
        "revisionCount": meta.get("revisionCount") or max(1, len(revisions)),
        "revisions": revisions,
        "importHistory": meta.get("importHistory") if isinstance(meta.get("importHistory"), list) else [],
        "firstSeenAt": str(meta.get("firstSeenAt") or meta.get("createdAt") or ""),
        "lastSeenAt": str(meta.get("lastSeenAt") or meta.get("createdAt") or ""),
    }


def hydrate_screenshot_hash(record, root):
    item = screenshot_artifact_record(record)
    if item.get("sha256"):
        return item
    stored = str(item.get("storedFilename") or "")
    if stored:
        path = (Path(root).resolve() / safe_name(stored)).resolve()
        if path_is_inside(path, Path(root).resolve()) and path.is_file():
            item["sha256"] = file_sha256(path)
    return item


def scan_type_label(value):
    return {"port": "Port Scan", "tcp": "Nmap Full Scan", "udp": "UDP Scan"}.get(value, value or "Scan")


@serialized_project_mutation
def upload_exploitation_image(lab_name, host_id, path_id, file_item):
    state = STORE.load_project(lab_name)
    host = state.get('hosts', {}).get(host_id)
    if not isinstance(host, dict):
        raise ValueError('The selected host is not in this engagement')
    record = host.get('exploitationPaths', {}).get('paths', {}).get(path_id)
    if not isinstance(record, dict):
        raise ValueError('Save this exploitation path before adding screenshots')
    if file_item is None or not getattr(file_item, 'filename', ''):
        raise ValueError('Choose a screenshot')
    original = str(file_item.filename).replace('\\', '/').rsplit('/', 1)[-1]
    extension = Path(original).suffix.lower()
    data = file_item.file.read(25 * 1024 * 1024 + 1)
    if not data or len(data) > 25 * 1024 * 1024:
        raise ValueError('Screenshot must be non-empty and smaller than 25 MB')
    validate_image_bytes(data, extension)
    digest = hashlib.sha256(data).hexdigest()
    directory = contained(STORE.host_root(lab_name, host), exploitation_directory(path_id, record) + '/images')
    STORE.asset_owner(directory)
    destination = contained(directory, digest + extension)
    directory.mkdir(parents=True, exist_ok=True)
    if destination.exists():
        if destination.read_bytes() != data:
            raise ValueError('An existing screenshot has different contents; preserve it before retrying')
    else:
        temporary = directory / ('.upload-' + uuid.uuid4().hex)
        try:
            with temporary.open('xb') as output:
                output.write(data); output.flush(); os.fsync(output.fileno())
            os.replace(temporary, destination)
        finally:
            temporary.unlink(missing_ok=True)
    return {'originalFilename': original, 'relativePath': 'images/' + destination.name,
            'storedPath': str(destination), 'sha256': digest, 'bytes': len(data)}


def upload_scan_image(lab_name, host, scan_type, command, file_item=None):
    if not lab_name:
        raise ValueError("labName is required")
    if not host or not host.get("ip"):
        raise ValueError("host IP is required")
    if scan_type not in {"port", "tcp", "udp"}:
        raise ValueError("scanType must be port, tcp, or udp")
    if file_item is None or not getattr(file_item, "filename", ""):
        raise ValueError("Screenshot is required")
    original = Path(file_item.filename).name
    ext = Path(original).suffix.lower()
    if ext not in IMAGE_EXTS:
        raise ValueError("Screenshot must be PNG, JPG, or JPEG")
    data = file_item.file.read()
    if not data or len(data) > 25 * 1024 * 1024:
        raise ValueError("Screenshot must be non-empty and smaller than 25 MB")
    validate_image_bytes(data, ext)
    host_dir = internal_host_dir(lab_name, host)
    screenshots_dir = host_dir / "screenshots"
    existing = ((host.get("scanImages") or {}).get(scan_type) if isinstance(host.get("scanImages"), dict) else None)
    known = [hydrate_screenshot_hash(screenshot_artifact_record(existing or {}, logical_path=f"scan/{scan_type}", source="scan-screenshot", objective=scan_type), screenshots_dir)] if existing else []
    artifact = store_versioned_artifact(
        data, root=screenshots_dir, original_filename=original,
        logical_path=f"scan/{scan_type}", source="scan-screenshot", objective=scan_type,
        artifact_type=f"screenshot-{ext.lstrip('.')}", kind="scan-screenshot",
        known_artifacts=known, host_key=str(host.get("id") or host.get("ip") or ""),
    )
    artifact = _artifact_storage_view(artifact, screenshots_dir, Path("screenshots"))
    image = {
        "type": scan_type, "title": scan_type_label(scan_type),
        "command": command or scan_command(scan_type, host),
        "artifactId": artifact.get("id"), "originalFilename": artifact.get("originalFilename"),
        "screenshotFilename": artifact.get("storedFilename"),
        "screenshotRelPath": f"screenshots/{artifact.get('storedFilename')}",
        "screenshotAbsPath": str((screenshots_dir / artifact.get("storedFilename")).resolve()),
        "createdAt": artifact.get("firstSeenAt") or datetime.datetime.now().isoformat(timespec="seconds"),
        "sha256": artifact.get("sha256"), "currentRevision": artifact.get("currentRevision"),
        "revisionCount": artifact.get("revisionCount"), "revisions": artifact.get("revisions", []),
        "importHistory": artifact.get("importHistory", []), "importOutcome": artifact.get("importOutcome"),
        "lastSeenAt": artifact.get("lastSeenAt"), "logicalPath": artifact.get("logicalPath"),
        "source": artifact.get("source"), "objective": artifact.get("objective"), "artifactType": artifact.get("artifactType"),
    }
    host.setdefault("scanImages", {"port": None, "tcp": None, "udp": None})
    host["scanImages"][scan_type] = image
    return image







def post_evidence(lab_name, host, evidence, file_item=None):
    if not lab_name:
        raise ValueError("labName is required")
    if not host or not host.get("ip"):
        raise ValueError("host IP is required")

    host_dir = internal_host_dir(lab_name, host); host_dir.mkdir(parents=True, exist_ok=True)
    screenshots_dir = host_dir / "screenshots"; screenshots_dir.mkdir(parents=True, exist_ok=True)

    evidence = dict(evidence or {})
    ev_type = evidence.get("type") or "custom"
    title = evidence.get("title") or evidence_type_label(ev_type)
    if ev_type in {"initial_access", "privilege_escalation", "post_exploitation", "credential_dump", "lateral_movement", "service", "custom"} and not title.strip():
        raise ValueError("Evidence name is required for this evidence type")
    evidence["title"] = title
    evidence["id"] = evidence.get("id") or str(uuid.uuid4())
    evidence["createdAt"] = evidence.get("createdAt") or datetime.datetime.now().isoformat(timespec="seconds")
    evidence["updatedAt"] = datetime.datetime.now().isoformat(timespec="seconds")

    if file_item is not None and getattr(file_item, "filename", ""):
        original = Path(file_item.filename).name
        ext = Path(original).suffix.lower()
        if ext not in IMAGE_EXTS:
            raise ValueError("Screenshot must be PNG, JPG, or JPEG")
        data = file_item.file.read()
        if not data or len(data) > 25 * 1024 * 1024:
            raise ValueError("Screenshot must be non-empty and smaller than 25 MB")
        validate_image_bytes(data, ext)
        known = []
        for row in host.get("evidence") or []:
            if not isinstance(row, dict) or not row.get("screenshotFilename"):
                continue
            known.append(hydrate_screenshot_hash(screenshot_artifact_record(row, logical_path=f"evidence/{row.get('id') or 'item'}", source="evidence-screenshot", objective=row.get("type") or "evidence"), screenshots_dir))
        artifact = store_versioned_artifact(
            data, root=screenshots_dir, original_filename=original,
            logical_path=f"evidence/{evidence['id']}", source="evidence-screenshot", objective=ev_type,
            artifact_type=f"screenshot-{ext.lstrip('.')}", kind="evidence-screenshot",
            known_artifacts=known, host_key=str(host.get("id") or host.get("ip") or ""),
        )
        artifact = _artifact_storage_view(artifact, screenshots_dir, Path("screenshots"))
        evidence.update({
            "screenshotArtifactId": artifact.get("id"), "screenshotOriginalFilename": artifact.get("originalFilename"),
            "screenshotFilename": artifact.get("storedFilename"), "screenshotRelPath": f"screenshots/{artifact.get('storedFilename')}",
            "screenshotAbsPath": str((screenshots_dir / artifact.get("storedFilename")).resolve()),
            "screenshotSha256": artifact.get("sha256"), "screenshotCurrentRevision": artifact.get("currentRevision"),
            "screenshotRevisionCount": artifact.get("revisionCount"), "screenshotRevisions": artifact.get("revisions", []),
            "screenshotImportHistory": artifact.get("importHistory", []), "screenshotImportOutcome": artifact.get("importOutcome"),
        })

    evidence["noteTargets"] = evidence.get("noteTargets") if isinstance(evidence.get("noteTargets"), list) else []
    return evidence


def evidence_matches(ev, allowed):
    return ev.get("type") in allowed


def host_attack_path_text(host):
    host = host if isinstance(host, dict) else {}
    sections = []
    manual = str(host.get("attackPath") or "").strip()
    if manual:
        sections.append(manual)
    steps = host.get("attackPathSteps") or []
    if isinstance(steps, list):
        active_steps = [step for step in steps if isinstance(step, dict) and step.get("active") is not False]
        active_steps.sort(key=lambda step: str(step.get("createdAt") or ""))
        sections.extend(str(step.get("summary") or "").strip() for step in active_steps if str(step.get("summary") or "").strip())
    return "\n\n".join(sections)


def add_evidence_items(doc, host, allowed_types=None):
    items = [ev for ev in (host.get("evidence") or []) if not isinstance(ev, dict) or ev.get("active") is not False]
    if allowed_types:
        items = [ev for ev in items if evidence_matches(ev, allowed_types)]
    if not items:
        return

    for ev in items:
        ev_type = ev.get("type")

        # Initial Access and Privilege Escalation should look like the OffSec sample:
        # Vulnerability Explanation, Vulnerability Fix, Severity, Steps to reproduce.
        # The finding title is already in the section heading, so don't print an extra Evidence line here.
        if ev_type in {"initial_access", "privilege_escalation"}:
            add_finding_detail_fields(doc, host, ev)
        else:
            add_labeled_para_with_gap(doc, "Evidence", ev.get("title") or evidence_type_label(ev.get("type")))

            if ev.get("proofFile"):
                add_labeled_para_with_gap(doc, "Proof File", ev.get("proofFile"))

            if ev.get("proofValue"):
                add_labeled_para_with_gap(doc, "Proof Value", ev.get("proofValue"))

            if ev.get("lootPath"):
                add_labeled_para_with_gap(doc, "Credential/File Path", ev.get("lootPath"))

            if ev.get("credentials"):
                add_labeled_para_with_gap(doc, "Credentials Found", "\n".join(credential_pair_lines(ev.get("credentials"))))

            if ev.get("credentialContext"):
                add_labeled_para_with_gap(doc, "Credential Context / Reuse Notes", ev.get("credentialContext"))

            if ev.get("sourceHost") or ev.get("targetHost"):
                add_labeled_para_with_gap(doc, "Lateral Movement Path", f"{ev.get('sourceHost') or 'source TBD'} -> {ev.get('targetHost') or 'target TBD'}")

            if ev.get("method"):
                add_labeled_para_with_gap(doc, "Method / Protocol", ev.get("method"))

            if ev.get("credentialUsed"):
                add_labeled_para_with_gap(doc, "Credential Used", ev.get("credentialUsed"))

            if ev.get("explanation"):
                add_labeled_para_with_gap(doc, "Vulnerability Explanation", ev.get("explanation"))

            if ev.get("fix"):
                add_labeled_para_with_gap(doc, "Vulnerability Fix", ev.get("fix"))

            if ev.get("severity"):
                add_severity_para_with_gap(doc, ev.get("severity"))

            if ev.get("notes"):
                add_labeled_para_with_gap(doc, "Evidence Notes", ev.get("notes"))

        img = ev.get("screenshotAbsPath")
        if img and Path(img).exists():
            try:
                doc.add_picture(img, width=Inches(5.7))
                add_report_spacer(doc, GAP_BLOCK)
            except Exception as e:
                add_labeled_para_with_gap(doc, "Screenshot", f"Could not embed image: {img} ({e})")
        elif ev.get("screenshotRelPath"):
            add_labeled_para_with_gap(doc, "Screenshot", ev.get("screenshotRelPath"))

        add_report_spacer(doc, GAP_BLOCK)


















def path_is_inside(child, parent):
    try:
        child = Path(child).resolve()
        parent = Path(parent).resolve()
        child.relative_to(parent)
        return True
    except Exception:
        return False


def _nearest_existing_directory(path):
    candidate = Path(path).expanduser()
    while not candidate.exists() and candidate != candidate.parent:
        candidate = candidate.parent
    return candidate if candidate.is_dir() else candidate.parent


def _read_only_oscp_template_status(template_id=""):
    selected_id = str(template_id or "").strip()
    candidates = []
    if REPORT_TEMPLATES_REGISTRY.is_file():
        registry = json.loads(REPORT_TEMPLATES_REGISTRY.read_text(encoding="utf-8"))
        templates = registry.get("templates") if isinstance(registry, dict) else []
        templates = templates if isinstance(templates, list) else []
        if not selected_id:
            selected_id = str(registry.get("defaultTemplateId") or "")
        for item in templates:
            if not isinstance(item, dict):
                continue
            if selected_id and str(item.get("id") or "") != selected_id:
                continue
            filename = safe_name(item.get("filename") or "")
            if filename:
                candidates.append(
                    (
                        REPORT_TEMPLATES_ROOT / filename,
                        str(item.get("displayName") or filename),
                    )
                )
    bundled = ROOT / "report_templates" / "OSCP-Exam-Report.docx"
    if not selected_id or selected_id == "built-in-oscp":
        candidates.append((bundled, "Bundled OSCP Exam Report"))
    for path, label in candidates:
        try:
            validate_docx_bytes(path.read_bytes())
            return {
                "ready": True,
                "currentValue": f"{label} ({path})",
                "reason": "The selected OSCP DOCX template is readable and structurally valid.",
                "consequence": "",
                "correctiveAction": "",
                "actionTarget": "reporting",
            }
        except Exception:
            continue
    return {
        "ready": False,
        "currentValue": selected_id or "No OSCP template selected",
        "reason": "No readable, structurally valid OSCP DOCX template was found.",
        "consequence": "Word report generation may fail until a template is selected.",
        "correctiveAction": "Open Reporting Options and select or import a valid OSCP DOCX template.",
        "actionTarget": "reporting",
    }


def _reference_pack_archive_manifest(archive_path):
    """Validate a Reference Pack archive without extracting or mutating it."""
    archive_path = Path(archive_path)
    if not archive_path.is_file():
        raise ValueError("Reference Pack archive is missing")
    if not 0 < archive_path.stat().st_size <= MAX_REFERENCE_PACK_BYTES:
        raise ValueError("Reference Pack archive is empty or exceeds the size limit")
    try:
        with zipfile.ZipFile(archive_path, "r") as bundle:
            members = bundle.infolist()
            if len(members) > MAX_REFERENCE_PACK_MEMBERS:
                raise ValueError(
                    f"Reference Pack contains too many files (maximum {MAX_REFERENCE_PACK_MEMBERS})"
                )
            normalized_members = {}
            total = 0
            for info in members:
                normalized = safe_recon_zip_member(info)
                if normalized in normalized_members:
                    raise ValueError(f"Reference Pack contains a duplicate path: {normalized}")
                normalized_members[normalized] = info
                if info.is_dir():
                    continue
                suffix = PurePosixPath(normalized).suffix.lower()
                if normalized != REFERENCE_PACK_MANIFEST and suffix != ".md":
                    raise ValueError(f"Reference Pack contains an unsupported file: {normalized}")
                if info.file_size <= 0:
                    raise ValueError(f"Reference Pack contains an empty file: {normalized}")
                if normalized == REFERENCE_PACK_MANIFEST:
                    if info.file_size > 8 * 1024 * 1024:
                        raise ValueError("Reference Pack manifest is too large")
                elif info.file_size > MAX_REFERENCE_NOTE_BYTES:
                    raise ValueError(f"Reference note exceeds the size limit: {normalized}")
                total += info.file_size
                if total > MAX_REFERENCE_PACK_TOTAL_BYTES:
                    raise ValueError("Expanded Reference Pack exceeds the size limit")
                if (
                    info.compress_size == 0
                    or info.file_size / max(info.compress_size, 1)
                    > MAX_REFERENCE_PACK_COMPRESSION_RATIO
                ):
                    raise ValueError(f"Suspicious compression ratio in Reference Pack: {normalized}")
            manifest_info = normalized_members.get(REFERENCE_PACK_MANIFEST)
            if manifest_info is None or manifest_info.is_dir():
                raise ValueError("Reference Pack is missing root index.json")
            bad_member = bundle.testzip()
            if bad_member:
                raise ValueError(f"Reference Pack CRC check failed: {bad_member}")
            manifest = json.loads(bundle.read(REFERENCE_PACK_MANIFEST).decode("utf-8"))
    except (OSError, zipfile.BadZipFile, KeyError, UnicodeDecodeError, json.JSONDecodeError) as error:
        raise ValueError(f"Reference Pack archive is invalid: {error}") from error

    if not isinstance(manifest, dict):
        raise ValueError("Reference Pack manifest must be an object")
    if str(manifest.get("packType") or "") != "aeros-reference-notes":
        raise ValueError("Unsupported Reference Pack type")
    if int(manifest.get("schemaVersion") or 0) != 1:
        raise ValueError("Unsupported Reference Pack schema version")
    raw_pack_id = str(manifest.get("packId") or "").strip()
    if not raw_pack_id:
        raise ValueError("Reference Pack is missing packId")
    pack_id = safe_name(raw_pack_id).lower()
    version = str(manifest.get("version") or "").strip()
    revision = str(manifest.get("contentRevision") or "").strip()
    if not version or not revision:
        raise ValueError("Reference Pack is missing version or contentRevision")
    notes = manifest.get("notes")
    if not isinstance(notes, dict):
        raise ValueError("Reference Pack manifest is missing notes")
    declared_count = manifest.get("count")
    if isinstance(declared_count, bool) or not isinstance(declared_count, int):
        raise ValueError("Reference Pack manifest count must be an integer")
    if declared_count != len(notes):
        raise ValueError(
            f"Reference Pack count mismatch: declared {declared_count}, indexed {len(notes)}"
        )
    seen_ids = set()
    seen_paths = set()
    for manifest_id, raw_entry in notes.items():
        if not isinstance(raw_entry, dict):
            raise ValueError(f"Reference note {manifest_id} must be an object")
        entry = dict(raw_entry)
        runtime_id = str(entry.get("id") or manifest_id).strip()
        if not runtime_id or runtime_id in seen_ids:
            raise ValueError(f"Reference Pack contains duplicate runtime note ID: {runtime_id}")
        seen_ids.add(runtime_id)
        entry["id"] = runtime_id
        _validate_reference_note_metadata(entry)
        relative = _safe_reference_relative_path(entry.get("path") or entry.get("notePath"))
        normalized_path = relative.as_posix()
        if normalized_path in seen_paths:
            raise ValueError(f"Reference Pack contains duplicate indexed path: {normalized_path}")
        seen_paths.add(normalized_path)
        note_info = normalized_members.get(normalized_path)
        if note_info is None or note_info.is_dir():
            raise ValueError(f"Reference note is missing: {normalized_path}")
    indexed_files = {REFERENCE_PACK_MANIFEST, *seen_paths}
    archived_files = {
        name for name, info in normalized_members.items() if not info.is_dir()
    }
    unindexed_files = sorted(archived_files - indexed_files)
    if unindexed_files:
        raise ValueError(
            f"Reference Pack contains an unindexed file: {unindexed_files[0]}"
        )
    return {
        **manifest,
        "packId": pack_id,
        "displayName": str(
            manifest.get("displayName") or manifest.get("library") or pack_id
        ),
        "version": version,
        "contentRevision": revision,
        "count": declared_count,
        "archiveMembers": len([info for info in members if not info.is_dir()]),
        "archivePath": str(archive_path),
    }


def _canonical_reference_pack_expectation():
    """Derive current build truth from the canonical archive manifest."""
    candidates = [
        ROOT
        / "AEROS_PROJECT_CONTROL"
        / "reference-packs"
        / CANONICAL_REFERENCE_PACK_FILENAME,
        ROOT / CANONICAL_REFERENCE_PACK_FILENAME,
    ]
    errors = []
    for archive in candidates:
        if not archive.is_file():
            continue
        try:
            manifest = _reference_pack_archive_manifest(archive)
            if manifest["packId"] != CANONICAL_REFERENCE_PACK_ID:
                raise ValueError(
                    f"Canonical packId mismatch: {manifest['packId'] or '<missing>'}"
                )
            return {
                "valid": True,
                "packId": manifest["packId"],
                "displayName": manifest["displayName"],
                "version": manifest["version"],
                "contentRevision": manifest["contentRevision"],
                "count": manifest["count"],
                "archiveMembers": manifest["archiveMembers"],
                "archivePath": manifest["archivePath"],
            }
        except Exception as error:
            errors.append(f"{archive}: {error}")
    if not errors:
        errors.append(
            f"{CANONICAL_REFERENCE_PACK_FILENAME} was not found in an accepted build location"
        )
    return {
        "valid": False,
        "packId": CANONICAL_REFERENCE_PACK_ID,
        "displayName": "AEROS Pentesting Notes V1",
        "version": "",
        "contentRevision": "",
        "count": 0,
        "archiveMembers": 0,
        "archivePath": "",
        "error": "; ".join(errors),
    }


def _reference_pack_status_value(pack, *, include_name=False):
    parts = []
    if include_name:
        parts.append(str(pack.get("displayName") or pack.get("packId") or "Reference Pack"))
    parts.append(str(pack.get("version") or "Version not declared"))
    revision = str(pack.get("contentRevision") or "").strip()
    parts.append(f"revision {revision}" if revision else "revision not declared")
    parts.append(f"{int(pack.get('count') or 0)} notes")
    return " · ".join(parts)


def _canonical_reference_pack_status():
    expectation = _canonical_reference_pack_expectation()
    if not expectation["valid"]:
        return {
            "ready": False,
            "currentValue": "Canonical Reference Pack expectation unavailable",
            "reason": (
                "This build's canonical Reference Pack expectation could not be verified: "
                f"{expectation['error']}"
            ),
            "consequence": "Installed OSCP methodology notes cannot be compared with build truth.",
            "correctiveAction": (
                f"Restore {CANONICAL_REFERENCE_PACK_FILENAME} in an accepted build location, "
                "then re-run OSCP preflight."
            ),
            "actionTarget": "application",
        }

    required = _reference_pack_status_value(expectation)
    installed_root = REFERENCE_PACKS_ROOT / expectation["packId"]
    import_action = (
        f"Import {Path(expectation['archivePath']).name} through "
        "Options → Application → Optional Reference Packs."
    )
    if not installed_root.is_dir():
        return {
            "ready": False,
            "currentValue": f"Installed: Not installed · Required: {required}",
            "reason": "The current canonical Optional Reference Pack is not installed.",
            "consequence": "OSCP methodology notes may be incomplete or unavailable.",
            "correctiveAction": import_action,
            "actionTarget": "application",
        }
    try:
        installed = _reference_pack_manifest(installed_root)
    except Exception as error:
        return {
            "ready": False,
            "currentValue": f"Installed: Invalid · Required: {required}",
            "reason": f"The installed Optional Reference Pack is invalid: {error}",
            "consequence": "The invalid installed pack is excluded from the active Reference Notes library.",
            "correctiveAction": import_action,
            "actionTarget": "application",
        }

    matches = all(
        (
            installed.get("packId") == expectation.get("packId"),
            str(installed.get("version") or "") == str(expectation.get("version") or ""),
            str(installed.get("contentRevision") or "")
            == str(expectation.get("contentRevision") or ""),
            int(installed.get("count") or 0) == int(expectation.get("count") or 0),
        )
    )
    if matches:
        return {
            "ready": True,
            "currentValue": _reference_pack_status_value(installed, include_name=True),
            "reason": "The installed Optional Reference Pack matches the current canonical pack for this build.",
            "consequence": "",
            "correctiveAction": "",
            "actionTarget": "application",
        }
    return {
        "ready": False,
        "currentValue": (
            f"Installed: {_reference_pack_status_value(installed)} · Required: {required}"
        ),
        "reason": "The installed Optional Reference Pack does not match the current canonical pack for this build.",
        "consequence": "OSCP methodology notes may be older, incomplete, or incompatible with current guidance.",
        "correctiveAction": import_action,
        "actionTarget": "application",
    }


def oscp_preflight_facts(
    template_id="", operator_output_root="", import_staging_root="", expected_build_identity=""
):
    checked_at = datetime.datetime.now(datetime.timezone.utc).isoformat()
    expected_build = str(expected_build_identity or "").strip().lower()
    server_matches = not expected_build or expected_build == BUILD_IDENTITY.lower()
    checks = {
        "server": {
            "ready": server_matches,
            "currentValue": (
                f"{APPLICATION_ID} build {BUILD_IDENTITY}; data root {DATA_ROOT_FINGERPRINT}"
            ),
            "reason": (
                "The local application and server identity match."
                if server_matches
                else "The browser application expected a different local server build."
            ),
            "consequence": (
                ""
                if server_matches
                else "Loading an engagement against this server could use incompatible code or storage."
            ),
            "correctiveAction": (
                ""
                if server_matches
                else "Close the stale AEROS process and restart this application build."
            ),
            "actionTarget": "application",
        }
    }
    try:
        writable = STORE.writable_status()
        checks["persistence"] = {
            "ready": True,
            "currentValue": f"Writable engagement file catalog: {writable['path']}",
            "reason": "An atomic file write and exclusive catalog lock succeeded.",
            "consequence": "",
            "correctiveAction": "",
            "actionTarget": "data",
        }
    except Exception as error:
        checks["persistence"] = {
            "ready": False,
            "currentValue": str(STORAGE_PATH),
            "reason": f"File storage write check failed: {error}",
            "consequence": "AEROS cannot safely save this OSCP engagement.",
            "correctiveAction": "Restore write access to the AEROS data directory, then check again.",
            "actionTarget": "data",
        }
    checks["template"] = _read_only_oscp_template_status(template_id)
    checks["referencePack"] = _canonical_reference_pack_status()
    operator_value = str(operator_output_root or "").strip()
    checks["operatorOutput"] = {
        "ready": bool(operator_value),
        "currentValue": operator_value or "Not configured",
        "reason": (
            "The current Operator Output Root will be inserted into operator-environment commands."
            if operator_value
            else "No Operator Output Root is configured."
        ),
        "consequence": (
            ""
            if operator_value
            else "Generated operator commands cannot target the intended output location."
        ),
        "correctiveAction": (
            ""
            if operator_value
            else "Set Operator Output Root in Options. AEROS does not require filesystem access to it."
        ),
        "actionTarget": "application",
    }

    staging_value = str(import_staging_root or "").strip()
    if staging_value:
        staging_path = Path(staging_value).expanduser()
        staging_parent = _nearest_existing_directory(staging_path)
        staging_ready = staging_parent.is_dir() and os.access(staging_parent, os.W_OK)
        checks["staging"] = {
            "ready": staging_ready,
            "currentValue": str(staging_path),
            "reason": (
                f"Nearest existing directory is writable: {staging_parent}"
                if staging_ready
                else f"No writable parent is available for: {staging_path}"
            ),
            "consequence": (
                ""
                if staging_ready
                else "Local import staging cannot receive operator-selected files."
            ),
            "correctiveAction": (
                ""
                if staging_ready
                else "Choose a local staging root under a writable directory."
            ),
            "actionTarget": "application",
        }
    else:
        checks["staging"] = {
            "ready": False,
            "currentValue": "Not configured",
            "reason": "No AEROS Import Staging Root is configured.",
            "consequence": "Fast local import staging will require manual file selection.",
            "correctiveAction": "Set AEROS Import Staging Root in OSCP setup or Options.",
            "actionTarget": "application",
        }
    checks["renderer"] = {
        "ready": Document is not None,
        "currentValue": "python-docx available" if Document is not None else "python-docx unavailable",
        "reason": (
            "The local Word report renderer is importable."
            if Document is not None
            else "The python-docx report renderer is unavailable."
        ),
        "consequence": "" if Document is not None else "Word report export will fail.",
        "correctiveAction": "" if Document is not None else "Install python-docx in the AEROS runtime.",
        "actionTarget": "reporting",
    }
    backup_parent = _nearest_existing_directory(LAB_BACKUPS_ROOT)
    backup_ready = backup_parent.is_dir() and os.access(backup_parent, os.W_OK)
    checks["backup"] = {
        "ready": backup_ready,
        "currentValue": str(LAB_BACKUPS_ROOT),
        "reason": (
            f"Internal backup parent is writable: {backup_parent}"
            if backup_ready
            else "The internal backup destination has no writable parent."
        ),
        "consequence": "" if backup_ready else "Automatic local recovery checkpoints may fail.",
        "correctiveAction": "" if backup_ready else "Restore write access to the AEROS data directory.",
        "actionTarget": "data",
    }
    return {
        "ok": True,
        "checkedAt": checked_at,
        "networkActions": 0,
        "checks": checks,
        "instance": {
            "applicationId": APPLICATION_ID,
            "buildIdentity": BUILD_IDENTITY,
            "dataRootFingerprint": DATA_ROOT_FINGERPRINT,
            "dataRootOverride": DATA_ROOT_OVERRIDE,
        },
    }


STATIC_ROOT_JSON_FILES = {
    "binary-contexts.json",
    "linux-home-loot-auto-files.json",
    "privileged-known-normal.json",
    "sgid-ignore.json",
    "writable-path-known-normal.json",
}
STATIC_IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".svg", ".ico"}
RUNTIME_ASSET_ATTRIBUTE_RE = re.compile(
    r'(?P<prefix><(?:script|link)\b[^>]*?\b(?:src|href)=")'
    r'(?P<url>[^"]+\.(?:js|css))(?P<suffix>")',
    re.IGNORECASE,
)


def static_path_allowed(candidate):
    """Expose only browser runtime assets, never source, tests, or .git data."""
    candidate = Path(candidate).resolve()
    root = ROOT.resolve()
    if not path_is_inside(candidate, root) or not candidate.is_file():
        return False
    relative = candidate.relative_to(root)
    if len(relative.parts) == 1:
        return relative.suffix.lower() in {".js", ".css"} or relative.name in STATIC_ROOT_JSON_FILES
    top = relative.parts[0].lower()
    if top == "app":
        if relative.suffix.lower() == ".js":
            return True
        return (
            len(relative.parts) == 3
            and relative.parts[1].lower() == "styles"
            and relative.suffix.lower() == ".css"
        )
    if top == "assets":
        return relative.suffix.lower() in STATIC_IMAGE_EXTENSIONS
    if top == "home-loot-baselines":
        return relative.suffix.lower() == ".json"
    return False


def version_runtime_asset_urls(html, build_identity=BUILD_IDENTITY):
    """Give every local JS/CSS URL a build-specific cache key."""
    version = str(build_identity or "").strip().lower()
    if not version:
        return str(html or "")

    def replace(match):
        url = match.group("url")
        if url.startswith(("http://", "https://", "//", "data:")):
            return match.group(0)
        separator = "&" if "?" in url else "?"
        return f'{match.group("prefix")}{url}{separator}v={version}{match.group("suffix")}'

    return RUNTIME_ASSET_ATTRIBUTE_RE.sub(replace, str(html or ""))



class Handler(SimpleHTTPRequestHandler):
    protocol_version = "HTTP/1.1"
    server_version = "AEROS"
    sys_version = ""

    def setup(self):
        super().setup()
        self.connection.settimeout(30)

    def translate_path(self, path):
        # Confine static file serving to the app folder. The stock handler
        # collapses "..", but this override previously did not, reintroducing
        # traversal. Resolve and verify the result stays under ROOT.
        path = path.split("?",1)[0].split("#",1)[0]
        candidate = (ROOT / unquote(path).lstrip("/")).resolve()
        root = ROOT.resolve()
        if candidate != root and root not in candidate.parents:
            return str(root / ".aeros-static-not-found")
        if path not in {"/", "/index.html"} and not static_path_allowed(candidate):
            return str(root / ".aeros-static-not-found")
        return str(candidate)

    def _authorized(self):
        # Reject requests from any other browser origin, and require the
        # per-launch token on API calls. Browser-loaded images use the same
        # token in an HttpOnly, SameSite cookie because <img> cannot add headers.
        host = str(self.headers.get("Host") or "").strip().lower()
        if host not in ALLOWED_HOSTS:
            # Blocks DNS-rebinding: the rebound page reaches the socket but its
            # Host header carries the attacker's domain, not a loopback name.
            return False
        origin = str(self.headers.get("Origin") or "").strip().lower()
        if origin and origin not in ALLOWED_ORIGINS:
            return False
        request_path = self.path.split("?", 1)[0]
        if request_path == "/api/native-shutdown":
            supplied = str(self.headers.get("X-AEROS-Shutdown-Token") or "")
            return bool(NATIVE_SHUTDOWN_TOKEN) and secrets.compare_digest(supplied, NATIVE_SHUTDOWN_TOKEN)
        if request_path == "/api/serve-file":
            cookie_header = str(self.headers.get("Cookie") or "")
            cookies = {}
            for part in cookie_header.split(";"):
                key, separator, value = part.strip().partition("=")
                if separator:
                    cookies[key] = value
            return secrets.compare_digest(cookies.get(SESSION_COOKIE_NAME, ""), APP_TOKEN)
        if request_path.startswith("/api/"):
            if not secrets.compare_digest(str(self.headers.get("X-App-Token") or ""), APP_TOKEN):
                return False
        return True

    def end_headers(self):
        # Defence-in-depth headers on every response.
        self.send_header("X-Content-Type-Options", "nosniff")
        self.send_header("X-Frame-Options", "DENY")
        self.send_header("Referrer-Policy", "no-referrer")
        self.send_header("Cross-Origin-Resource-Policy", "same-origin")
        self.send_header("Permissions-Policy", "camera=(), microphone=(), geolocation=(), payment=(), usb=()")
        self.send_header("Content-Security-Policy",
                         "default-src 'self'; base-uri 'none'; object-src 'none'; frame-ancestors 'none'; "
                         "form-action 'self'; connect-src 'self'; img-src 'self' data: blob:; "
                         f"style-src 'self' 'unsafe-inline'; script-src 'self' 'nonce-{CSP_NONCE}'")
        super().end_headers()

    def _read_body(self):
        length = int(self.headers.get("Content-Length", "0") or "0")
        if length < 0 or length > MAX_BODY_BYTES:
            raise ValueError("Request body too large")
        return self.rfile.read(length)

    def _serve_index(self):
        # Serve index.html with the token + a fetch shim injected into <head>.
        try:
            html = (ROOT / "index.html").read_text(encoding="utf-8")
        except Exception:
            return self._json(404, {"ok": False, "error": "index.html not found"})
        html = version_runtime_asset_urls(html)
        shim = (
            f'<head>\n<script nonce="{CSP_NONCE}">window.__APP_TOKEN__=' + json.dumps(APP_TOKEN) + ";"
            "(function(){var t=window.__APP_TOKEN__,f=window.fetch.bind(window);"
            "window.fetch=function(i,o){o=o||{};var u=(typeof i===\"string\")?i:(i&&i.url)||\"\";"
            "var p=new URL(u,window.location.href);"
            "if(p.origin===window.location.origin&&p.pathname.indexOf(\"/api/\")===0&&p.pathname!==\"/api/serve-file\") {"
            "var h=new Headers(o.headers||((typeof i!==\"string\"&&i&&i.headers)||{}));"
            "h.set(\"X-App-Token\",t);o.headers=h;}"
            "return f(i,o);};})();</script>"
        )
        html = html.replace("<head>", shim, 1)
        html = html.replace("<script>", f'<script nonce="{CSP_NONCE}">')
        body = html.encode("utf-8")
        self.send_response(200)
        self.send_header("Content-Type", "text/html; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store, no-cache, must-revalidate, max-age=0")
        self.send_header("Pragma", "no-cache")
        self.send_header("Set-Cookie", f"{SESSION_COOKIE_NAME}={APP_TOKEN}; HttpOnly; SameSite=Strict; Path=/")
        self.end_headers()
        self.wfile.write(body)

    def _json(self, status, data):
        b=json.dumps(data,separators=(",", ":"),ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type","application/json; charset=utf-8")
        self.send_header("Content-Length",str(len(b)))
        self.send_header("Cache-Control","no-store, no-cache, must-revalidate, max-age=0")
        self.send_header("Pragma","no-cache")
        self.send_header("Expires","0")
        self.end_headers()
        self.wfile.write(b)

    def _send_file(self, path, content_type=None, download_name=None, cache_control="no-store"):
        path = Path(path)
        self.send_response(200)
        self.send_header("Content-Type", content_type or mimetypes.guess_type(path.name)[0] or "application/octet-stream")
        if download_name:
            self.send_header("Content-Disposition", f'attachment; filename="{safe_name(download_name)}"')
        self.send_header("Content-Length", str(path.stat().st_size))
        self.send_header("Cache-Control", cache_control)
        self.end_headers()
        try:
            with path.open("rb") as source:
                shutil.copyfileobj(source, self.wfile, length=1024 * 1024)
        except (BrokenPipeError, ConnectionResetError):
            pass
    def do_GET(self):
        if not self._authorized(): return self._json(403, {"ok": False, "error": "Forbidden"})
        request_path = self.path.split("?", 1)[0]
        if request_path in ("/", "/index.html"): return self._serve_index()
        if request_path == "/api/health":
            return self._json(200, {
                "ok": True,
                "application": "AEROS",
                "apiVersion": 3,
                "instance": {
                    "applicationId": APPLICATION_ID,
                    "buildIdentity": BUILD_IDENTITY,
                    "serverPid": os.getpid(),
                    "dataRootFingerprint": DATA_ROOT_FINGERPRINT,
                    "dataRootOverride": DATA_ROOT_OVERRIDE,
                },
                "capabilities": {"reconImport": True, "unifiedImportIntake": True, "peasImport": True, "commandNotes": True, "referenceNotes": True, "referencePackImport": True, "referenceNoteEditing": True, "oscpPreflight": True},
            })
        if request_path == "/api/oscp-preflight":
            from urllib.parse import urlparse, parse_qs
            qs = parse_qs(urlparse(self.path).query)
            try:
                return self._json(
                    200,
                    oscp_preflight_facts(
                        template_id=qs.get("templateId", [""])[0],
                        operator_output_root=qs.get("operatorOutputRoot", [""])[0],
                        import_staging_root=qs.get("importStagingRoot", [""])[0],
                        expected_build_identity=qs.get("expectedBuildIdentity", [""])[0],
                    ),
                )
            except Exception as error:
                return self._json(400, {"ok": False, "error": str(error)})
        if request_path == "/api/serve-file":
            from urllib.parse import urlparse, parse_qs
            qs = parse_qs(urlparse(self.path).query)
            raw_path = qs.get("path", [""])[0]
            try:
                p = Path(raw_path).resolve()
                # Also confine the authenticated image request to the asset store.
                if not STORE.asset_owner(p):
                    return self._json(403, {"ok": False, "error": "Path not allowed"})
                if not p.exists() or not p.is_file():
                    return self._json(404, {"ok": False, "error": "File not found"})
                if p.suffix.lower() not in IMAGE_EXTS:
                    return self._json(400, {"ok": False, "error": "Unsupported file type"})
                ctype = mimetypes.guess_type(str(p))[0] or "application/octet-stream"
                return self._send_file(p, ctype, cache_control="private, max-age=60")
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})
        if request_path == "/api/reference-notes":
            try:
                manifest = load_reference_notes_manifest()
                return self._json(200, {"ok": True, **manifest})
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})
        if request_path == "/api/reference-note":
            from urllib.parse import urlparse, parse_qs
            qs = parse_qs(urlparse(self.path).query)
            try:
                result = read_reference_note(qs.get("noteId", [""])[0], qs.get("profile", [""])[0])
                return self._json(200, {"ok": True, **result})
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})
        if request_path == "/api/command-notes":
            try:
                manifest = load_command_notes_manifest()
                return self._json(200, {"ok": True, **manifest})
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})
        if request_path == "/api/command-note":
            from urllib.parse import urlparse, parse_qs
            qs = parse_qs(urlparse(self.path).query)
            try:
                result = read_command_note(qs.get("noteId", [""])[0], qs.get("profile", [""])[0])
                return self._json(200, {"ok": True, **result})
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})
        if self.path == "/api/profiles":
            try:
                profiles = read_profile_store()
                return self._json(200, {"ok": True, "profiles": profiles})
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})
        if self.path.startswith("/api/download-engagement-document"):
            from urllib.parse import urlparse, parse_qs
            qs = parse_qs(urlparse(self.path).query)
            try:
                lab = qs.get("labName", [""])[0]; stored = qs.get("storedFilename", [""])[0]
                path = engagement_document_path(lab, stored)
                download_name = safe_name(qs.get("downloadName", [path.name])[0])
                return self._send_file(path, download_name=download_name)
            except Exception as e: return self._json(400, {"ok": False, "error": str(e)})
        if self.path.startswith("/api/download-roe-document"):
            from urllib.parse import urlparse, parse_qs
            qs = parse_qs(urlparse(self.path).query)
            try:
                lab = qs.get("labName", [""])[0]
                stored = qs.get("storedFilename", [""])[0]
                download_name = safe_name(qs.get("downloadName", [stored or "engagement-document"])[0])
                path = roe_document_path(lab, stored)
                return self._send_file(path, download_name=download_name)
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})
        if self.path.startswith("/api/download-peas-artifact"):
            from urllib.parse import urlparse, parse_qs
            qs = parse_qs(urlparse(self.path).query)
            try:
                path = peas_artifact_from_stored_path(qs.get("path", [""])[0])
                return self._send_file(path, "text/plain; charset=utf-8", qs.get("name", [path.name])[0])
            except FileNotFoundError:
                return self._json(404, {"ok": False, "error": "PEAS output artifact not found"})
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})

        if self.path.startswith("/api/download-exploit-artifact"):
            from urllib.parse import urlparse, parse_qs
            qs = parse_qs(urlparse(self.path).query)
            try:
                path = exploit_attempt_artifact_from_stored_path(qs.get("path", [""])[0])
                download_name = safe_name(qs.get("downloadName", [path.name])[0])
                return self._send_file(path, mimetypes.guess_type(download_name)[0], download_name)
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})
        if self.path.startswith("/api/download-scan-artifact"):
            from urllib.parse import urlparse, parse_qs
            qs = parse_qs(urlparse(self.path).query)
            try:
                lab = qs.get("labName", [""])[0]; host_key = qs.get("hostKey", [""])[0]; stored = qs.get("storedFilename", [""])[0]
                path = scan_artifact_path(lab, host_key, stored)
                download_name = safe_name(qs.get("downloadName", [path.name])[0])
                return self._send_file(path, "application/octet-stream", download_name)
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})
        if self.path == "/api/report-templates":
            try:
                registry = load_report_template_registry()
                return self._json(200, {"ok": True, **registry, "storagePath": str(REPORT_TEMPLATES_ROOT)})
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})
        if self.path.startswith("/api/download-report-template"):
            from urllib.parse import urlparse, parse_qs
            qs = parse_qs(urlparse(self.path).query)
            try:
                path, item = report_template_path(qs.get("templateId", [""])[0])
                filename = safe_name(item.get("originalFilename") or item.get("displayName") or "report-template")
                if not filename.lower().endswith(".docx"): filename += ".docx"
                return self._send_file(path, "application/vnd.openxmlformats-officedocument.wordprocessingml.document", filename)
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})
        if self.path.startswith("/api/storage-status"):
            try:
                return self._json(200, {"ok": True, **STORE.status(), "legacyDataMigratedFrom": str(LEGACY_DATA_MIGRATED_FROM or "")})
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})
        if self.path.startswith("/api/list-archives"):
            try:
                return self._json(200, {"ok": True, "archives": list_engagement_archives()})
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})
        if self.path.startswith("/api/download-archive"):
            from urllib.parse import urlparse, parse_qs
            qs = parse_qs(urlparse(self.path).query)
            archive_id = qs.get("archiveId", [""])[0]
            try:
                path = archive_path_for_id(archive_id)
                if not path.exists() or not path.is_file():
                    return self._json(404, {"ok": False, "error": "Archive not found"})
                return self._send_file(path, "application/zip", path.name)
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})
        if self.path.startswith("/api/list-labs"):
            try:
                summaries = STORE.list_project_summaries(include_archived=False)
                labs = [row["name"] for row in summaries]
                return self._json(200, {"ok": True, "labs": labs, "summaries": summaries, "source": "files", "storagePath": str(STORAGE_PATH)})
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})
        if self.path.startswith("/api/load-project"):
            from urllib.parse import urlparse, parse_qs
            qs = parse_qs(urlparse(self.path).query)
            lab = qs.get("labName", [""])[0]
            try:
                record = STORE.load_project_record(lab)
                if record is None:
                    state_path = project_state_path(lab)
                    if state_path.exists():
                        STORE.import_legacy_file(state_path)
                        record = STORE.load_project_record(lab)
                if record is not None:
                    return self._json(200, {
                        "ok": True,
                        "state": record["state"],
                        "engagementId": record["id"],
                        "revision": record["revision"],
                        "source": str(STORAGE_PATH),
                        "storage": "files",
                    })
                return self._json(404, {"ok": False, "error": "Engagement not found"})
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})
        if request_path == "/api/read-import-result-member":
            from urllib.parse import urlparse, parse_qs
            qs = parse_qs(urlparse(self.path).query)
            try:
                path = import_intake_member_path(
                    qs.get("token", [""])[0],
                    qs.get("member", [""])[0],
                )
                return self._send_file(path, mimetypes.guess_type(path.name)[0] or "application/octet-stream")
            except FileNotFoundError as e:
                return self._json(404, {"ok": False, "error": str(e)})
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})
        if request_path.startswith("/api/"):
            return self._json(404, {"ok": False, "error": f"Unknown API endpoint: {request_path}"})
        return super().do_GET()
    def do_POST(self):
        if not self._authorized(): return self._json(403, {"ok": False, "error": "Forbidden"})
        if self.path.split("?", 1)[0] == "/api/native-shutdown":
            self._json(200, {"ok": True, "ownerPid": NATIVE_OWNER_PID})
            import threading
            threading.Thread(target=self.server.shutdown, name="aeros-native-shutdown", daemon=True).start()
            return
        if self.path == "/api/inspect-import-results":
            try:
                body = self._read_body()
                fields, files = parse_multipart_form(self.headers, body)
                result = inspect_import_archive(files.get("input") or files.get("archive"))
                return self._json(200, {"ok": True, **result})
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e), "limits": import_intake_limits()})

        if self.path == "/api/discard-import-results":
            try:
                raw = self._read_body().decode("utf-8")
                payload = json.loads(raw or "{}")
                discarded = discard_import_intake(payload.get("token", ""))
                return self._json(200, {"ok": True, "discarded": discarded})
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})

        if self.path == "/api/import-recon":
            try:
                body = self._read_body()
                fields, files = parse_multipart_form(self.headers, body)
                result = parse_recon_upload(files.get("collection") or files.get("recon"))
                return self._json(200, {"ok": True, **result})
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})

        if self.path == "/api/import-reference-pack":
            try:
                body = self._read_body()
                fields, files = parse_multipart_form(self.headers, body)
                result = import_reference_pack(files.get("referencePack") or files.get("pack"))
                return self._json(200, {"ok": True, **result})
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})

        if self.path == "/api/upload-scan-artifact":
            try:
                body = self._read_body(); fields, files = parse_multipart_form(self.headers, body)
                result = upload_scan_artifact(
                    fields.get("labName", ""),
                    fields.get("hostKey", ""),
                    files.get("artifact"),
                    fields.get("objective", ""),
                    fields.get("source", ""),
                    fields.get("logicalPath", ""),
                    fields.get("knownArtifacts", "[]"),
                    fields.get("physicalSourcePath", ""),
                    fields.get("runId", ""),
                )
                return self._json(200, {"ok": True, **result, "storagePath": str(scan_artifacts_root(fields.get("labName", ""), fields.get("hostKey", "")))})
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})

        if self.path == "/api/upload-engagement-document":
            try:
                body = self._read_body(); fields, files = parse_multipart_form(self.headers, body)
                item = upload_engagement_document(fields.get("labName", ""), files.get("document"), fields.get("displayName", ""), fields.get("category", "supporting"), fields.get("notes", ""), fields.get("knownDocuments", "[]"))
                return self._json(200, {"ok": True, "document": item, "storagePath": str(general_documents_root(fields.get("labName", "")))})
            except Exception as e: return self._json(400, {"ok": False, "error": str(e)})

        if self.path == "/api/upload-roe-document":
            try:
                body = self._read_body(); fields, files = parse_multipart_form(self.headers, body)
                item = upload_roe_document(fields.get("labName", ""), files.get("document"), fields.get("displayName", ""), fields.get("documentKind", "roe"), fields.get("knownDocuments", "[]"))
                return self._json(200, {"ok": True, "document": item, "storagePath": str(roe_documents_root(fields.get("labName", "")))})
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})

        if self.path == "/api/upload-report-template":
            try:
                body = self._read_body(); fields, files = parse_multipart_form(self.headers, body)
                item = upload_report_template(files.get("template"), fields.get("displayName", ""))
                return self._json(200, {"ok": True, "template": item})
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})

        if self.path == "/api/upload-exploitation-image":
            try:
                fields, files = parse_multipart_form(self.headers, self._read_body())
                image = upload_exploitation_image(fields.get('labName', ''), fields.get('hostId', ''), fields.get('pathId', ''), files.get('screenshot'))
                return self._json(200, {'ok': True, 'image': image})
            except Exception as error:
                return self._json(400, {'ok': False, 'error': str(error)})

        if self.path == "/api/upload-scan-image":
            try:
                body = self._read_body()
                fields, files = parse_multipart_form(self.headers, body)
                lab_name = fields.get("labName", "")
                host = json.loads(fields.get("host", "{}"))
                scan_type = fields.get("scanType", "")
                command = fields.get("command", "")
                file_item = files.get("screenshot")
                image = upload_scan_image(lab_name, host, scan_type, command, file_item)
                return self._json(200, {"ok": True, "image": image})
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})

        if self.path == "/api/upload-peas-artifact":
            try:
                body = self._read_body(); fields, files = parse_multipart_form(self.headers, body)
                host = json.loads(fields.get("host", "{}"))
                artifact = upload_peas_artifact(
                    fields.get("labName", ""), host, fields.get("accessContextId", ""),
                    fields.get("tool", "unknown"), files.get("artifact"), fields.get("knownArtifacts", "[]")
                )
                return self._json(200, {"ok": True, "artifact": artifact})
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})

        if self.path == "/api/upload-exploit-artifact":
            try:
                body = self._read_body()
                fields, files = parse_multipart_form(self.headers, body)
                host = json.loads(fields.get("host", "{}"))
                known_artifact = json.loads(fields.get("knownArtifact", "{}") or "{}")
                artifact = upload_exploit_attempt_artifact(
                    fields.get("labName", ""), host, fields.get("attemptId", ""), files.get("artifact"), known_artifact
                )
                return self._json(200, {"ok": True, "artifact": artifact})
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})

        if self.path == "/api/post-evidence":
            try:
                body = self._read_body()
                fields, files = parse_multipart_form(self.headers, body)
                lab_name = fields.get("labName", "")
                host = json.loads(fields.get("host", "{}"))
                evidence = json.loads(fields.get("evidence", "{}"))
                file_item = files.get("screenshot")
                ev = post_evidence(lab_name, host, evidence, file_item)
                return self._json(200, {"ok": True, "evidence": ev})
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})

        try: raw=self._read_body().decode("utf-8")
        except Exception as e: return self._json(413, {"ok": False, "error": str(e)})
        try: payload=json.loads(raw or "{}")
        except Exception: return self._json(400, {"ok": False, "error": "Invalid JSON"})
        if self.path == "/api/set-default-report-template":
            try:
                template_id = payload.get("templateId", "")
                report_template_path(template_id)
                registry = load_report_template_registry(); registry["defaultTemplateId"] = template_id; save_report_template_registry(registry)
                return self._json(200, {"ok": True})
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})
        if self.path == "/api/delete-report-template":
            try:
                if not delete_report_template(payload.get("templateId", "")):
                    return self._json(404, {"ok": False, "error": "Template not found"})
                return self._json(200, {"ok": True})
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})
        if self.path == "/api/delete-peas-artifact":
            try:
                paths = payload.get("storedPaths", [])
                if not isinstance(paths, list): paths = [payload.get("storedPath", "")]
                deleted = 0
                for raw_path in paths:
                    if not raw_path: continue
                    try:
                        deleted += 1 if delete_peas_artifact_by_path(raw_path) else 0
                    except FileNotFoundError:
                        pass
                return self._json(200, {"ok": True, "deleted": deleted})
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})

        if self.path == "/api/delete-exploit-artifact":
            try:
                deleted = delete_exploit_attempt_artifact_by_path(payload.get("storedPath", ""))
                return self._json(200, {"ok": True, "deleted": bool(deleted)})
            except FileNotFoundError:
                return self._json(200, {"ok": True, "deleted": False})
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})
        if self.path == "/api/delete-internal-screenshot":
            try:
                deleted = delete_internal_screenshot_by_path(payload.get("storedPath", ""))
                return self._json(200, {"ok": True, "deleted": bool(deleted)})
            except FileNotFoundError:
                return self._json(200, {"ok": True, "deleted": False})
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})
        if self.path == "/api/delete-reference-pack":
            try:
                deleted = delete_reference_pack(payload.get("packId"))
                if not deleted:
                    return self._json(404, {"ok": False, "error": "Reference pack not found"})
                return self._json(200, {"ok": True})
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})
        if self.path == "/api/save-command-note":
            try:
                result = save_custom_command_note(payload.get("noteId"), payload.get("profile"), payload.get("content"))
                return self._json(200, {"ok": True, **result})
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})
        if self.path == "/api/reset-command-note":
            try:
                result = reset_custom_command_note(payload.get("noteId"), payload.get("profile"))
                return self._json(200, {"ok": True, **result})
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})
        if self.path == "/api/save-reference-note":
            try:
                result = save_custom_reference_note(payload.get("noteId"), payload.get("profile"), payload.get("content"))
                return self._json(200, {"ok": True, **result})
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})
        if self.path == "/api/reset-reference-note":
            try:
                result = reset_custom_reference_note(payload.get("noteId"), payload.get("profile"))
                return self._json(200, {"ok": True, **result})
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})
        if self.path == "/api/save-profiles":
            try:
                profiles = payload.get("profiles", {})
                save_profile_store(profiles)
                return self._json(200, {"ok": True, "path": str(PROFILES_PATH), "count": len(profiles)})
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})
        if self.path == "/api/relink-project":
            try:
                with PROJECT_MUTATION_LOCK:
                    result = STORE.relink_project(payload.get("labName", ""), payload.get("directory", ""))
                return self._json(200, {"ok": True, **result})
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})
        if self.path == "/api/save-project":
            try:
                lab = payload.get("projectName") or payload.get("labName")
                if not lab:
                    raise ValueError("projectName is required")
                guard = payload.get("_saveGuard") or {}
                expected_lab = str(guard.get("expectedLabName") or "").strip()
                if not expected_lab or expected_lab.casefold() != str(lab).strip().casefold():
                    raise ValueError("Save blocked because the loaded engagement does not match the requested engagement.")
                expected_revision = guard.get("expectedRevision")
                if expected_revision is not None:
                    try:
                        expected_revision = int(expected_revision)
                    except (TypeError, ValueError) as exc:
                        raise ValueError("Save guard revision must be an integer") from exc
                clean_reserved_hosts_from_state(payload)
                path = project_state_path(lab)
                project_payload = dict(payload)
                # Application-wide runtime data must never be stored inside an
                # engagement. prepare_state enforces the same boundary.
                for transient_field in (
                    "labs", "profiles", "serverAvailable", "dynamicNotes",
                    "dynamicNoteRoot", "dynamicNoteCount", "dynamicNoteError",
                    "_saveGuard",
                ):
                    project_payload.pop(transient_field, None)
                with PROJECT_MUTATION_LOCK:
                    result = STORE.save_project(
                        project_payload,
                        expected_name=expected_lab,
                        expected_revision=expected_revision,
                        allow_empty_hosts=bool(guard.get("allowEmptyHosts")),
                        reason="application-save",
                        create_only=guard.get("createOnly") is True,
                    )
                return self._json(200, {
                    "ok": True,
                    "path": result["path"],
                    "storagePath": result["storagePath"],
                    "storage": result["state"]["engagementConfig"]["storage"],
                    "pathRewrites": result["pathRewrites"],
                    "engagementId": result["engagementId"],
                    "revision": result["revision"],
                    "changed": result.get("changed", True),
                    "schemaVersion": SCHEMA_VERSION,
                    "applicationVersion": APPLICATION_VERSION,
                    "backupPath": result["backupPath"],
                    "intelFiles": [],
                })
            except EngagementAlreadyExistsError as e:
                return self._json(409, {"ok": False, "error": str(e), "code": "engagement-exists"})
            except RevisionConflictError as e:
                return self._json(409, {"ok": False, "error": str(e), "code": "revision-conflict"})
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})
        if self.path == "/api/archive-project":
            try:
                lab = payload.get("labName") or payload.get("projectName")
                result = archive_engagement(lab, bool(payload.get("removeWorkingCopy")))
                return self._json(200, {"ok": True, **result})
            except RevisionConflictError as e:
                return self._json(409, {"ok": False, "error": str(e), "code": "revision-conflict"})
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})
        if self.path == "/api/restore-archive":
            try:
                result = restore_engagement_archive(payload.get("archiveId"), payload.get("newName", ""))
                return self._json(200, {"ok": True, **result})
            except FileExistsError as e:
                return self._json(409, {"ok": False, "error": str(e)})
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})
        if self.path == "/api/delete-archive":
            try:
                if not delete_engagement_archive(payload.get("archiveId")):
                    return self._json(404, {"ok": False, "error": "Archive not found"})
                return self._json(200, {"ok": True})
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})
        if self.path == "/api/delete-engagement-document":
            try:
                deleted = delete_engagement_document(payload.get("labName", ""), payload.get("storedFilename", ""), payload.get("extractedFilename", ""), payload.get("storedFilenames", []))
                return self._json(200, {"ok": True, "deleted": deleted})
            except Exception as e: return self._json(400, {"ok": False, "error": str(e)})
        if self.path == "/api/delete-roe-document":
            try:
                lab = payload.get("labName", "")
                deleted = delete_roe_document(lab, payload.get("storedFilename", ""), payload.get("extractedFilename", ""), payload.get("storedFilenames", []))
                return self._json(200, {"ok": True, "deleted": deleted})
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})
        if self.path == "/api/delete-project":
            try:
                lab = payload.get("labName") or payload.get("projectName")
                expected_revision = payload.get("expectedRevision")
                if expected_revision is not None:
                    try:
                        expected_revision = int(expected_revision)
                    except (TypeError, ValueError) as exc:
                        raise ValueError("Delete revision must be an integer") from exc
                if not delete_project(lab, expected_revision=expected_revision):
                    return self._json(404, {"ok": False, "error": "Lab not found"})
                return self._json(200, {"ok": True, "labName": lab})
            except RevisionConflictError as e:
                return self._json(409, {"ok": False, "error": str(e), "code": "revision-conflict"})
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})
        if self.path == "/api/report-docx":
            try:
                data=generate_report_docx(payload)
                filename="OSCP-Exam-Report.docx"
                name = str(payload.get("projectName") or "")
                if not STORE.project_exists(name):
                    raise ValueError("Save the engagement before exporting a report")
                report_path = contained(STORE.project_root(name), "reports", filename)
                atomic_write_bytes(report_path, data)
                self.send_response(200)
                self.send_header("Content-Type","application/vnd.openxmlformats-officedocument.wordprocessingml.document")
                self.send_header("Content-Disposition",f'attachment; filename="{filename}"')
                self.send_header("Content-Length",str(len(data)))
                self.end_headers(); self.wfile.write(data); return
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})
        if self.path == "/api/report-html":
            try:
                name = str(payload.get("labName") or "")
                if not STORE.project_exists(name):
                    raise ValueError("Save the engagement before exporting a report")
                html = payload.get("html")
                if not isinstance(html, str) or not html.strip():
                    raise ValueError("Report HTML is required")
                path = contained(STORE.project_root(name), "reports", "report.html")
                atomic_write_text(path, html)
                return self._json(200, {"ok": True, "path": str(path)})
            except Exception as e:
                return self._json(400, {"ok": False, "error": str(e)})
        return self._json(404, {"ok":False,"error":"Not found"})


class LocalThreadingHTTPServer(ThreadingHTTPServer):
    daemon_threads = True
    allow_reuse_address = True
    request_queue_size = 16

    def __init__(self, *args, max_concurrent_requests=8, **kwargs):
        self._request_slots = threading.BoundedSemaphore(max_concurrent_requests)
        super().__init__(*args, **kwargs)

    def process_request(self, request, client_address):
        if not self._request_slots.acquire(blocking=False):
            body = b"AEROS server is busy.\n"
            try:
                request.sendall(
                    b"HTTP/1.1 503 Service Unavailable\r\n"
                    b"Connection: close\r\n"
                    b"Content-Type: text/plain; charset=utf-8\r\n"
                    + f"Content-Length: {len(body)}\r\n\r\n".encode("ascii")
                    + body
                )
            except OSError:
                pass
            self.shutdown_request(request)
            return
        try:
            super().process_request(request, client_address)
        except Exception:
            self._request_slots.release()
            raise

    def process_request_thread(self, request, client_address):
        try:
            super().process_request_thread(request, client_address)
        finally:
            self._request_slots.release()


if __name__ == "__main__":
    os.chdir(ROOT)
    imported = initialize_persistence()
    url = f"http://{HOST}:{PORT}/index.html"
    print(f"AEROS V1 running at {url}")
    print(f"File storage store: {STORAGE_PATH}")
    if LEGACY_DATA_MIGRATED_FROM:
        print(f"Migrated development data from: {LEGACY_DATA_MIGRATED_FROM}")
    if imported:
        print(f"Imported {len(imported)} JSON engagement(s) into File storage.")
    print("Press Ctrl+C to stop.")
    LocalThreadingHTTPServer((HOST, PORT), Handler).serve_forever()
